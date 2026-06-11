import os, sys
WORKDIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, WORKDIR)
from e100_data import T
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "Transactions"
hdr = ["Invoice","Plate","Date","Time","Station","Product","Qty (L/un)","Gross price EUR (incl VAT)",
       "Discount EUR/L","Discount total EUR","Net price EUR/L (doc)","VAT EUR","Amount incl VAT EUR",
       "Chk1: discount=qty x rate","Chk2: amount=qty x gross - disc","Chk3: VAT=21/121 x amount","Chk4: net price"]
ws.append(hdr)
fill = PatternFill("solid", start_color="1B7340")
bw = Font(bold=True, color="FFFFFF", name="Arial", size=9)
norm = Font(name="Arial", size=9)
for c in ws[1]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(horizontal="center", wrap_text=True)

r = 2
for inv, plate, d, t, stn, prod, qty, brut, du, rem, np, tva, gross in T:
    ws.append([inv, plate, "2026-"+d, t, stn, prod, qty, brut, du, rem, np, tva, gross,
        f'=IF(ABS(G{r}*I{r}-J{r})<=0.011,"OK","CHECK")',
        f'=IF(ABS(ROUND(G{r}*H{r}-J{r},2)-M{r})<=0.011,"OK","CHECK")',
        f'=IF(ABS(ROUND(M{r}*21/121,2)-L{r})<=0.011,"OK","CHECK")',
        f'=IF(ABS((M{r}-L{r})/G{r}-K{r})<=0.0006,"OK","CHECK")'])
    r += 1
last = r - 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = norm
for col, f in (("G","#,##0.00"),("H","0.0000"),("I","0.0000"),("J","#,##0.00"),
               ("K","0.0000"),("L","#,##0.00"),("M","#,##0.00")):
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=ord(col)-64, max_col=ord(col)-64):
        row[0].number_format = f
for col, w in zip("ABCDEFGHIJKLMNOPQ",[9,8,10,6,22,7,9,10,9,9,10,9,11,10,12,11,9]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:Q{last}"

# Invoice totals check
ws2 = wb.create_sheet("Invoice totals check")
ws2.append(["Check item","Invoice states","Calculated","Diff","Check"])
for c in ws2[1]:
    c.font = bw; c.fill = fill
def SC(inv, prod, col):
    return (f'=ROUND(SUMPRODUCT((Transactions!$A$2:$A${last}="{inv}")*'
            f'(Transactions!$F$2:$F${last}="{prod}")*Transactions!${col}$2:${col}${last}),2)')
items = [
 ("BE98759 Diesel litres",22336.88,SC("BE98759","Diesel","G"),0.01),
 ("BE98759 Diesel discount EUR",4589.11,SC("BE98759","Diesel","J"),0.02),
 ("BE98759 Diesel net (excl VAT) EUR",36562.02,
  f'=ROUND(SUMPRODUCT((Transactions!$A$2:$A${last}="BE98759")*(Transactions!$F$2:$F${last}="Diesel")*(Transactions!$M$2:$M${last}-Transactions!$L$2:$L${last})),2)',0.05),
 ("BE98759 AdBlue litres",2140.93,SC("BE98759","AdBlue","G"),0.01),
 ("BE98759 AdBlue gross EUR",1750.00,SC("BE98759","AdBlue","M"),0.02),
 ("BE98759 total discount EUR",5037.04,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE98759",Transactions!$J$2:$J${last}),2)',0.02),
 ("BE98759 total VAT EUR",7981.74,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE98759",Transactions!$L$2:$L${last}),2)',0.05),
 ("BE98759 TOTAL incl VAT EUR (payable)",45990.04,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE98759",Transactions!$M$2:$M${last}),2)',0.02),
 ("BE99954 Diesel litres",24454.60,SC("BE99954","Diesel","G"),0.01),
 ("BE99954 Diesel discount EUR",5386.14,SC("BE99954","Diesel","J"),0.02),
 ("BE99954 Diesel net (excl VAT) EUR",38360.05,
  f'=ROUND(SUMPRODUCT((Transactions!$A$2:$A${last}="BE99954")*(Transactions!$F$2:$F${last}="Diesel")*(Transactions!$M$2:$M${last}-Transactions!$L$2:$L${last})),2)',0.05),
 ("BE99954 AdBlue litres",2388.88,SC("BE99954","AdBlue","G"),0.01),
 ("BE99954 AdBlue gross EUR",1994.78,SC("BE99954","AdBlue","M"),0.02),
 ("BE99954 Parking gross EUR",36.66,SC("BE99954","Parking","M"),0.01),
 ("BE99954 total discount EUR",6034.96,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE99954",Transactions!$J$2:$J${last}),2)',0.02),
 ("BE99954 total VAT EUR",8408.17,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE99954",Transactions!$L$2:$L${last}),2)',0.05),
 ("BE99954 TOTAL incl VAT EUR (payable)",48447.10,
  f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"BE99954",Transactions!$M$2:$M${last}),2)',0.02),
 ("BOTH invoices payable EUR",94437.14,
  f'=ROUND(SUM(Transactions!$M$2:$M${last}),2)',0.03),
]
r2 = 2
for name, doc, calc, tol in items:
    ws2.append([name, doc, calc, f"=C{r2}-B{r2}", f'=IF(ABS(D{r2})<={tol},"OK","CHECK")'])
    r2 += 1
for row in ws2.iter_rows(min_row=2, max_row=r2-1):
    for c in row: c.font = norm
    for c in row[1:4]: c.number_format = "#,##0.00"
for col, w in zip("ABCDE",[36,15,15,9,8]):
    ws2.column_dimensions[col].width = w

# Card (vehicle) totals check
ws3 = wb.create_sheet("Card totals check")
ws3.append(["Invoice","Plate","Stated card totals: discount","VAT","Gross","Calc discount","Calc VAT","Calc gross","Check"])
for c in ws3[1]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(wrap_text=True)
ct = {
("BE98759","MTG433"):(324.79,552.12,3181.34),("BE98759","MZD950"):(10.11,6.41,36.92),
("BE98759","MZD970"):(183.69,400.02,2304.85),("BE98759","NGM038"):(359.83,587.10,3382.87),
("BE98759","NGM043"):(175.13,310.77,1790.61),("BE98759","NGM601"):(44.00,208.26,1199.96),
("BE98759","NGM594"):(132.86,243.00,1400.14),("BE98759","NGM884"):(354.02,494.65,2850.14),
("BE98759","NGM883"):(59.19,92.10,530.70),("BE98759","NEP680"):(25.73,52.91,304.90),
("BE98759","MYI227"):(378.73,634.57,3656.33),("BE98759","MYI228"):(6.60,4.18,24.11),
("BE98759","MYI669"):(185.19,255.99,1475.03),("BE98759","MJK906"):(180.86,258.62,1490.17),
("BE98759","NOC558"):(322.10,449.18,2588.09),("BE98759","NOC560"):(169.41,251.15,1447.10),
("BE98759","NOC561"):(369.79,534.62,3080.44),("BE98759","NOC562"):(391.92,661.95,3814.10),
("BE98759","NOC563"):(68.75,128.78,741.98),("BE98759","NOC566"):(17.11,10.85,62.54),
("BE98759","NOC569"):(5.50,3.49,20.09),("BE98759","MMC392"):(328.94,460.16,2651.39),
("BE98759","NGM039"):(242.86,366.51,2111.77),("BE98759","AIB691"):(363.73,527.28,3038.14),
("BE98759","AIB697"):(199.95,282.09,1625.33),("BE98759","AIB689"):(132.77,186.02,1071.81),
("BE98759","MZD958"):(3.48,18.94,109.19),
("BE99954","MTG433"):(452.05,630.96,3635.50),("BE99954","MSF043"):(4.20,13.96,80.44),
("BE99954","MZD954"):(35.00,109.39,630.29),("BE99954","MZD970"):(307.31,529.09,3048.54),
("BE99954","NGM038"):(372.56,469.31,2704.12),("BE99954","NGM043"):(379.25,513.27,2957.47),
("BE99954","NGM601"):(13.33,86.52,498.50),("BE99954","NGM884"):(265.65,400.36,2306.87),
("BE99954","NGM883"):(167.37,232.98,1342.42),("BE99954","NEP680"):(13.32,24.13,139.03),
("BE99954","MYI227"):(344.40,475.80,2741.47),("BE99954","MYI669"):(466.63,617.57,3558.39),
("BE99954","MJK906"):(168.43,226.40,1304.49),("BE99954","NOC558"):(268.45,435.44,2508.90),
("BE99954","NOC560"):(31.22,14.07,81.09),("BE99954","NOC562"):(460.97,589.73,3397.97),
("BE99954","NOC563"):(395.09,507.27,2922.82),("BE99954","NOC569"):(33.37,15.20,87.61),
("BE99954","MMC392"):(511.55,694.13,3999.54),("BE99954","NGM039"):(327.00,403.55,2325.22),
("BE99954","NRE619"):(8.36,27.79,160.14),("BE99954","AIB691"):(7.04,4.46,25.72),
("BE99954","AIB697"):(301.20,407.43,2347.69),("BE99954","AIB692"):(339.85,481.75,2775.74),
("BE99954","AIB689"):(353.87,492.85,2839.75),("BE99954","MZD958"):(7.49,4.75,27.38),
}
r3 = 2
for (inv, plate), (drem, dvat, dgr) in ct.items():
    base = (f'SUMPRODUCT((Transactions!$A$2:$A${last}=A{r3})*'
            f'(Transactions!$B$2:$B${last}=B{r3})*Transactions!')
    ws3.append([inv, plate, drem, dvat, dgr,
        f"=ROUND({base}$J$2:$J${last}),2)",
        f"=ROUND({base}$L$2:$L${last}),2)",
        f"=ROUND({base}$M$2:$M${last}),2)",
        f'=IF(AND(ABS(F{r3}-C{r3})<=0.02,ABS(G{r3}-D{r3})<=0.05,ABS(H{r3}-E{r3})<=0.02),"OK","CHECK")'])
    r3 += 1
last3 = r3 - 1
for row in ws3.iter_rows(min_row=2, max_row=last3):
    for c in row: c.font = norm
    for c in row[2:8]: c.number_format = "#,##0.00"
for col, w in zip("ABCDEFGHI",[9,9,13,10,11,12,10,11,8]):
    ws3.column_dimensions[col].width = w

# NET price summary
ws4 = wb.create_sheet("NET price summary", 0)
ws4.append(["NET PRICE OVERVIEW (excl. VAT, after E100 discount) - UAB Vestroidas, May 2026, Belgium"])
ws4["A1"].font = Font(bold=True, name="Arial", size=11)
ws4.append([])
ws4.append(["Diesel by station","Litres","Net EUR","Wtd avg net EUR/L","Min net EUR/L","Max net EUR/L","Discount EUR/L"])
for c in ws4[3]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(wrap_text=True, horizontal="center")
stations = [("BE94 Poperinge","0.2300"),("BE725 Longlier","0.1800"),("BE428 La Louviere","0.2300"),
            ("BE92 Veurne","0.1050"),("BE9 Hoogstraten","0.0800"),("BE1042 Meer (orange)","0.1930"),
            ("BE768 Meer-Hoogstraten","0.1850")]
r4 = 4
for stn, disc in stations:
    b = (f'SUMPRODUCT((Transactions!$E$2:$E${last}=A{r4})*'
         f'(Transactions!$F$2:$F${last}="Diesel")*')
    ws4.append([stn,
        f"={b}Transactions!$G$2:$G${last})",
        f"=ROUND({b}(Transactions!$M$2:$M${last}-Transactions!$L$2:$L${last})),2)",
        f"=C{r4}/B{r4}",
        f'=ROUND(MIN(IF((Transactions!$E$2:$E${last}=A{r4})*(Transactions!$F$2:$F${last}="Diesel"),Transactions!$K$2:$K${last})),4)',
        f'=ROUND(MAX(IF((Transactions!$E$2:$E${last}=A{r4})*(Transactions!$F$2:$F${last}="Diesel"),Transactions!$K$2:$K${last})),4)',
        disc])
    r4 += 1
tr4 = r4
ws4.append(["TOTAL DIESEL", f"=SUM(B4:B{r4-1})", f"=SUM(C4:C{r4-1})", f"=C{tr4}/B{tr4}",
            f'=ROUND(MIN(IF(Transactions!$F$2:$F${last}="Diesel",Transactions!$K$2:$K${last})),4)',
            f'=ROUND(MAX(IF(Transactions!$F$2:$F${last}="Diesel",Transactions!$K$2:$K${last})),4)',""])
r4 += 1
ws4.append([])
r4 += 1
ws4.append(["AdBlue (all stations)",
    f'=SUMIF(Transactions!$F$2:$F${last},"AdBlue",Transactions!$G$2:$G${last})',
    f'=ROUND(SUMPRODUCT((Transactions!$F$2:$F${last}="AdBlue")*(Transactions!$M$2:$M${last}-Transactions!$L$2:$L${last})),2)',
    f"=C{r4}/B{r4}",
    f'=ROUND(MIN(IF(Transactions!$F$2:$F${last}="AdBlue",Transactions!$K$2:$K${last})),4)',
    f'=ROUND(MAX(IF(Transactions!$F$2:$F${last}="AdBlue",Transactions!$K$2:$K${last})),4)',"varies"])
r4 += 1
ws4.append(["Parking (1 un.)", 1, 30.30, "", "", "", ""])
r4 += 1
ws4.append([])
r4 += 1
ws4.append(["Diesel net price by half-month","Litres","Net EUR","Wtd avg net EUR/L"])
for c in ws4[r4]:
    if c.value: c.font = Font(bold=True, name="Arial", size=10)
r4 += 1
for inv, label in (("BE98759","01-15 May (BE98759)"),("BE99954","16-31 May (BE99954)")):
    b = (f'SUMPRODUCT((Transactions!$A$2:$A${last}="{inv}")*'
         f'(Transactions!$F$2:$F${last}="Diesel")*')
    ws4.append([label, f"={b}Transactions!$G$2:$G${last})",
        f"=ROUND({b}(Transactions!$M$2:$M${last}-Transactions!$L$2:$L${last})),2)",
        f"=C{r4}/B{r4}"])
    r4 += 1
for row in ws4.iter_rows(min_row=4, max_row=r4-1):
    for c in row:
        if c.font is None or not c.font.bold: c.font = norm
    if row[1].value is not None: row[1].number_format = "#,##0.00"
    if row[2].value is not None: row[2].number_format = "#,##0.00"
    for c in row[3:6]:
        c.number_format = "0.0000"
for c in ws4[tr4]: c.font = Font(bold=True, name="Arial", size=10)
for col, w in zip("ABCDEFG",[26,12,12,15,12,12,12]):
    ws4.column_dimensions[col].width = w

wb.save(os.path.join(WORKDIR, "E100_BE98759_BE99954_transactions.xlsx"))
print("saved", last-1, "lines,", last3-1, "cards")
