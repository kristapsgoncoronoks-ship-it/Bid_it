"""
EXCEL REPORTING - polished workbooks from fuel_history.db.

Reusable styling helpers (styled headers, number formats, banded rows, totals,
conditional colour scales, charts, freeze panes) plus a one-click executive
SUMMARY report. All prices are NET EUR/L, final (VAT excluded, rebates applied).

    python3 reports.py [period]      -> Fleet_Fuel_Summary_<period>.xlsx
"""
import os
import sqlite3
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter

import money
import queries

WORKDIR = os.path.dirname(os.path.abspath(__file__))
DB = f"{WORKDIR}/fuel_history.db"

# palette
INK = "1A2733"; ACC = "0E5FA8"; OKG = "1B7340"; BADR = "C8102E"
HEADER_FILL = PatternFill("solid", fgColor=ACC)
TITLE_FILL = PatternFill("solid", fgColor=INK)
BAND_FILL = PatternFill("solid", fgColor="F2F6FA")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="DDE4EA")
TOPB = Border(top=Side(style="thin", color="9FB3C4"))

FMT_INT = "#,##0"
FMT_EUR = "#,##0.00"
FMT_PRICE = "#,##0.0000"
FMT_PCT = "0.0%"


def connect():
    import db_tuning
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    db_tuning.tune(con)  # WAL + busy_timeout for safe multi-process access
    return con


# ---------------------------------------------------------------- styling
def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = WHITE_BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26


def band_rows(ws, first, last, ncols):
    for r in range(first, last + 1):
        if (r - first) % 2:
            for c in range(1, ncols + 1):
                ws.cell(r, c).fill = BAND_FILL


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def number_format(ws, col, first, last, fmt):
    for r in range(first, last + 1):
        ws.cell(r, col).number_format = fmt


def totals_row(ws, row, ncols, label_col=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = Font(bold=True)
        cell.border = TOPB
    ws.cell(row, label_col).font = Font(bold=True)


# ---------------------------------------------------------------- data
def _periods(con):
    return [r[0] for r in con.execute("SELECT DISTINCT period FROM transactions ORDER BY period DESC")]

def _kpis(con, period):
    r = con.execute("""
        SELECT ROUND(SUM(net_eur),2) net, ROUND(SUM(vat_eur),2) vat,
               ROUND(SUM(net_eur)+SUM(vat_eur),2) gross, COUNT(*) lines,
               (SELECT ROUND(SUM(qty),0) FROM transactions WHERE period=? AND product_group='Diesel') dl,
               (SELECT ROUND(SUM(net_eur_eff)/NULLIF(SUM(qty),0),4) FROM transactions
                  WHERE period=? AND product_group='Diesel') eurl
        FROM transactions WHERE period=?""", (period, period, period)).fetchone()
    return r

def _by(con, period, dim):
    return con.execute(f"""
        SELECT {dim} k, ROUND(SUM(qty),0) litres, ROUND(SUM(net_eur),2) net,
               ROUND(SUM(vat_eur),2) vat,
               ROUND(SUM(net_eur_eff)/NULLIF(SUM(qty),0),4) eurl
        FROM transactions WHERE period=? GROUP BY {dim} ORDER BY net DESC""", (period,)).fetchall()

def _by_entity(con, period):
    return con.execute("""
        SELECT entity k, country, ROUND(SUM(net_eur),2) net, ROUND(SUM(vat_eur),2) vat,
               ROUND(SUM(net_eur)+SUM(vat_eur),2) gross
        FROM transactions WHERE period=? GROUP BY entity, country ORDER BY net DESC""",
        (period,)).fetchall()

def _trend(con):
    return con.execute("""
        SELECT period, ROUND(SUM(qty),0) litres,
               ROUND(SUM(net_eur),2) net,
               ROUND(SUM(net_eur_eff)/NULLIF(SUM(qty),0),4) eurl
        FROM transactions WHERE product_group='Diesel' GROUP BY period ORDER BY period""").fetchall()

def _savings(con, period):
    # Single source of truth for the avoidable-overpay loop lives in
    # queries.q_savings; here we only remap to this module's key names.
    r = queries.q_savings(con, period)
    return {"total": r["total"],
            "by_sup": r["by_supplier"],
            "by_ctry": r["by_country"]}


# ---------------------------------------------------------------- summary report
def summary_workbook(period=None, path=None):
    con = connect()
    ps = _periods(con)
    period = period or (ps[0] if ps else None)
    if period is None:
        con.close()
        raise ValueError("no data loaded — nothing to report")
    k = _kpis(con, period)
    sv = _savings(con, period)
    wb = Workbook()
    _sheet_summary(wb, period, k, sv,
                   _by(con, period, "supplier"), _by(con, period, "country"))
    _sheet_table(wb, "By supplier", _by(con, period, "supplier"), period, chart=True)
    _sheet_table(wb, "By country", _by(con, period, "country"), period)
    _sheet_entity(wb, _by_entity(con, period), period)
    _sheet_trend(wb, _trend(con))
    _sheet_savings(wb, sv, period)
    con.close()
    path = path or os.path.join(WORKDIR, f"Fleet_Fuel_Summary_{period}.xlsx")
    wb.save(path)
    return path


def _title(ws, period, subtitle, span="A1:F1"):
    ws.merge_cells(span)
    t = ws.cell(1, 1, f"Fleet Fuel Report — {period}")
    t.font = Font(bold=True, color="FFFFFF", size=16)
    t.fill = TITLE_FILL
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    last_col = span.split(":")[1][0]
    ws.merge_cells(f"A2:{last_col}2")
    s = ws.cell(2, 1, subtitle)
    s.font = Font(italic=True, color="5B6B7A", size=9)
    s.alignment = Alignment(horizontal="left", indent=1)


def _kpi_card(ws, row, col, value, label, fmt=FMT_EUR, accent=ACC):
    v = ws.cell(row, col, value)
    v.font = Font(bold=True, size=15, color=accent)
    v.number_format = fmt
    v.alignment = Alignment(horizontal="left")
    lab = ws.cell(row + 1, col, label)
    lab.font = Font(size=9, color="5B6B7A")
    for r in (row, row + 1):
        ws.cell(r, col).border = Border(left=Side(style="thick", color=accent))


def _sheet_summary(wb, period, k, sv, by_sup, by_ctry):
    ws = wb.active; ws.title = "Summary"
    gen = datetime.date.today().isoformat()
    _title(ws, period, f"NET EUR/L, final (VAT excluded, rebates applied) · generated {gen}", "A1:F1")
    set_widths(ws, [24, 16, 16, 16, 16, 16])

    # KPI cards (row 4 value / row 5 label)
    cards = [
        (f"{(k['dl'] or 0):,.0f}", "Diesel litres", FMT_INT, ACC),
        (k["eurl"] or 0, "Fleet eff. €/L", FMT_PRICE, ACC),
        (k["net"] or 0, "Net spend (EUR)", FMT_EUR, INK),
        (k["vat"] or 0, "Reclaimable VAT (EUR)", FMT_EUR, OKG),
        (sv["total"], "Avoidable overpay (EUR)", FMT_EUR, BADR),
    ]
    for i, (val, lab, fmt, acc) in enumerate(cards):
        # litres value is a string already formatted; keep numeric where possible
        if isinstance(val, str):
            c = ws.cell(4, i + 1, k["dl"] or 0); c.number_format = FMT_INT
            c.font = Font(bold=True, size=15, color=acc)
            ws.cell(5, i + 1, lab).font = Font(size=9, color="5B6B7A")
            for r in (4, 5):
                ws.cell(r, i + 1).border = Border(left=Side(style="thick", color=acc))
        else:
            _kpi_card(ws, 4, i + 1, val, lab, fmt, acc)

    # mini table: top suppliers by net spend
    r0 = 8
    ws.cell(r0, 1, "Spend by supplier").font = Font(bold=True, size=11)
    hdr = ["Supplier", "Litres", "Net EUR", "VAT EUR", "€/L eff", "Share"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r0 + 1, j, h)
    style_header(ws, r0 + 1, len(hdr))
    tot_net = sum((r["net"] or 0) for r in by_sup) or 1
    rr = r0 + 2
    for r in by_sup:
        ws.cell(rr, 1, r["k"])
        ws.cell(rr, 2, r["litres"] or 0).number_format = FMT_INT
        ws.cell(rr, 3, r["net"] or 0).number_format = FMT_EUR
        ws.cell(rr, 4, r["vat"] or 0).number_format = FMT_EUR
        ws.cell(rr, 5, r["eurl"] or 0).number_format = FMT_PRICE
        ws.cell(rr, 6, (r["net"] or 0) / tot_net).number_format = FMT_PCT
        rr += 1
    band_rows(ws, r0 + 2, rr - 1, len(hdr))
    # totals
    ws.cell(rr, 1, "TOTAL")
    ws.cell(rr, 2, f"=SUM(B{r0+2}:B{rr-1})").number_format = FMT_INT
    ws.cell(rr, 3, f"=SUM(C{r0+2}:C{rr-1})").number_format = FMT_EUR
    ws.cell(rr, 4, f"=SUM(D{r0+2}:D{rr-1})").number_format = FMT_EUR
    ws.cell(rr, 6, 1).number_format = FMT_PCT
    totals_row(ws, rr, len(hdr))
    # conditional colour scale on €/L (green cheap -> red dear) + data bars on litres
    ws.conditional_formatting.add(f"E{r0+2}:E{rr-1}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"))
    ws.conditional_formatting.add(f"B{r0+2}:B{rr-1}",
        DataBarRule(start_type="min", end_type="max", color="9CC3E6"))

    # bar chart: net spend by supplier
    ch = BarChart(); ch.type = "col"; ch.title = "Net spend by supplier (EUR)"
    ch.height = 7.5; ch.width = 16; ch.legend = None
    data = Reference(ws, min_col=3, min_row=r0 + 1, max_row=rr - 1)
    cats = Reference(ws, min_col=1, min_row=r0 + 2, max_row=rr - 1)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ws.add_chart(ch, f"H{r0+1}")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"


def _sheet_table(wb, title, rows, period, chart=False):
    ws = wb.create_sheet(title)
    hdr = [title.replace("By ", "").capitalize(), "Litres", "Net EUR", "VAT EUR", "€/L eff"]
    for j, h in enumerate(hdr, 1):
        ws.cell(1, j, h)
    style_header(ws, 1, len(hdr))
    rr = 2
    for r in rows:
        ws.cell(rr, 1, r["k"])
        ws.cell(rr, 2, r["litres"] or 0).number_format = FMT_INT
        ws.cell(rr, 3, r["net"] or 0).number_format = FMT_EUR
        ws.cell(rr, 4, r["vat"] or 0).number_format = FMT_EUR
        ws.cell(rr, 5, r["eurl"] or 0).number_format = FMT_PRICE
        rr += 1
    band_rows(ws, 2, rr - 1, len(hdr))
    ws.cell(rr, 1, "TOTAL")
    ws.cell(rr, 2, f"=SUM(B2:B{rr-1})").number_format = FMT_INT
    ws.cell(rr, 3, f"=SUM(C2:C{rr-1})").number_format = FMT_EUR
    ws.cell(rr, 4, f"=SUM(D2:D{rr-1})").number_format = FMT_EUR
    totals_row(ws, rr, len(hdr))
    ws.conditional_formatting.add(f"E2:E{rr-1}",
        ColorScaleRule(start_type="min", start_color="63BE7B",
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color="F8696B"))
    set_widths(ws, [22, 14, 16, 14, 12])
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    if chart:
        ch = BarChart(); ch.type = "bar"; ch.title = f"{title} — net spend"
        ch.height = 8; ch.width = 14; ch.legend = None
        data = Reference(ws, min_col=3, min_row=1, max_row=rr - 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "G2")
        # second chart: effective €/L (col E, NET EUR/L basis) — directly comparable bars
        # so the dearest supplier is obvious. References the SAME cells already written.
        ch2 = BarChart(); ch2.type = "col"; ch2.title = "Effective €/L by supplier (NET, VAT-excl)"
        ch2.height = 8; ch2.width = 14; ch2.legend = None
        d2 = Reference(ws, min_col=5, min_row=1, max_row=rr - 1)
        ch2.add_data(d2, titles_from_data=True); ch2.set_categories(cats)
        ws.add_chart(ch2, "G18")


def _sheet_entity(wb, rows, period):
    ws = wb.create_sheet("By entity (VAT)")
    hdr = ["Entity", "Country", "Net EUR", "VAT reclaimable", "Gross EUR"]
    for j, h in enumerate(hdr, 1):
        ws.cell(1, j, h)
    style_header(ws, 1, len(hdr))
    rr = 2
    for r in rows:
        ws.cell(rr, 1, r["k"]); ws.cell(rr, 2, r["country"])
        ws.cell(rr, 3, r["net"] or 0).number_format = FMT_EUR
        ws.cell(rr, 4, r["vat"] or 0).number_format = FMT_EUR
        ws.cell(rr, 5, r["gross"] or 0).number_format = FMT_EUR
        rr += 1
    band_rows(ws, 2, rr - 1, len(hdr))
    ws.cell(rr, 1, "TOTAL")
    for col, L in ((3, "C"), (4, "D"), (5, "E")):
        ws.cell(rr, col, f"=SUM({L}2:{L}{rr-1})").number_format = FMT_EUR
    totals_row(ws, rr, len(hdr))
    set_widths(ws, [26, 14, 16, 16, 16])
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    # chart: net spend by entity x country (col C = Net EUR), referencing the cells above.
    if rr > 3:
        ch = BarChart(); ch.type = "bar"; ch.title = "Net spend by entity (EUR)"
        ch.height = 8; ch.width = 14; ch.legend = None
        data = Reference(ws, min_col=3, min_row=1, max_row=rr - 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "G2")


def _sheet_trend(wb, rows):
    ws = wb.create_sheet("Trend")
    hdr = ["Period", "Diesel litres", "Net EUR", "Fleet eff. €/L"]
    for j, h in enumerate(hdr, 1):
        ws.cell(1, j, h)
    style_header(ws, 1, len(hdr))
    rr = 2
    for r in rows:
        ws.cell(rr, 1, r["period"])
        ws.cell(rr, 2, r["litres"] or 0).number_format = FMT_INT
        ws.cell(rr, 3, r["net"] or 0).number_format = FMT_EUR
        ws.cell(rr, 4, r["eurl"] or 0).number_format = FMT_PRICE
        rr += 1
    set_widths(ws, [12, 14, 16, 14])
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    if rr > 3:
        ch = LineChart(); ch.title = "Fleet effective €/L trend"; ch.height = 8; ch.width = 16
        data = Reference(ws, min_col=4, min_row=1, max_row=rr - 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "F2")


def _sheet_savings(wb, sv, period):
    ws = wb.create_sheet("Savings")
    _title(ws, period, "Avoidable overpay vs the cheapest same-day, same-country rival (diesel)", "A1:D1")
    set_widths(ws, [22, 16, 4, 0])
    ws.cell(4, 1, "Total avoidable overpay (EUR)").font = Font(bold=True)
    tc = ws.cell(4, 2, sv["total"]); tc.number_format = FMT_EUR
    tc.font = Font(bold=True, size=13, color=BADR)
    ctry_block = None  # (header_row, first_data_row, last_data_row) for the chart
    for start_row, title, items in ((7, "Overpay by supplier", sv["by_sup"]),
                                    (7 + len(sv["by_sup"]) + 4, "Overpay by country", sv["by_ctry"])):
        ws.cell(start_row, 1, title).font = Font(bold=True, size=11)
        ws.cell(start_row + 1, 1, "Name"); ws.cell(start_row + 1, 2, "Overpay EUR")
        style_header(ws, start_row + 1, 2)
        r = start_row + 2
        for name, val in items:
            ws.cell(r, 1, name)
            ws.cell(r, 2, money.f2(val)).number_format = FMT_EUR
            r += 1
        band_rows(ws, start_row + 2, r - 1, 2)
        if r > start_row + 2:
            ws.conditional_formatting.add(f"B{start_row+2}:B{r-1}",
                DataBarRule(start_type="min", end_type="max", color="F4A6A6"))
            if title == "Overpay by country":
                ctry_block = (start_row + 1, start_row + 2, r - 1)
    ws.freeze_panes = "A3"; ws.sheet_view.showGridLines = False
    # chart: avoidable overpay (EUR) by country, referencing the country block cells.
    if ctry_block:
        hdr_row, first, last = ctry_block
        ch = BarChart(); ch.type = "col"; ch.title = "Avoidable overpay (EUR) by country"
        ch.height = 7.5; ch.width = 14; ch.legend = None
        data = Reference(ws, min_col=2, min_row=hdr_row, max_row=last)
        cats = Reference(ws, min_col=1, min_row=first, max_row=last)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "E7")


def savings_intel_workbook(s, path=None):
    """Consolidated "Savings & Intelligence" workbook from a savings_intel.summary()
    dict `s`. A KPI/summary sheet (headline addressable EUR + by-country breakdown)
    plus a top-actions detail sheet. Fuel-cost intelligence ONLY — NO VAT/claim/
    recovery figures. All prices NET EUR/L, final (VAT excluded, rebates applied);
    EUR totals are HALF_UP (money.f2), produced upstream by the detectors."""
    period = s.get("period") or "all"
    wb = Workbook()

    # --- Summary / KPI sheet -------------------------------------------------
    ws = wb.active; ws.title = "Summary"
    _title(ws, period,
           "Fuel-cost intelligence — NET EUR/L, final (VAT excluded, rebates applied). "
           "No VAT/claim figures.", "A1:E1")
    set_widths(ws, [26, 18, 18, 18, 18])
    cards = [
        (s["avoidable_overpay_eur"], "Avoidable overpay (EUR)", FMT_EUR, BADR),
        (s["recoverable_contract_eur"], "Recoverable contract (EUR)", FMT_EUR, OKG),
        (s["anomaly_count"], "Anomalies flagged", FMT_INT, ACC),
        (s["total_addressable_eur"], "TOTAL addressable (EUR)", FMT_EUR, INK),
    ]
    for i, (val, lab, fmt, acc) in enumerate(cards):
        _kpi_card(ws, 4, i + 1, val, lab, fmt, acc)

    # by-country breakdown
    r0 = 8
    ws.cell(r0, 1, "Addressable by country").font = Font(bold=True, size=11)
    hdr = ["Country", "Avoidable overpay EUR", "Recoverable contract EUR", "Addressable EUR"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r0 + 1, j, h)
    style_header(ws, r0 + 1, len(hdr))
    rr = r0 + 2
    for c in s["by_country"]:
        ws.cell(rr, 1, c["country"])
        ws.cell(rr, 2, c["overpay_eur"]).number_format = FMT_EUR
        ws.cell(rr, 3, c["recover_eur"]).number_format = FMT_EUR
        ws.cell(rr, 4, c["addressable_eur"]).number_format = FMT_EUR
        rr += 1
    band_rows(ws, r0 + 2, rr - 1, len(hdr))
    if rr > r0 + 2:
        ws.cell(rr, 1, "TOTAL")
        for col, L in ((2, "B"), (3, "C"), (4, "D")):
            ws.cell(rr, col, f"=SUM({L}{r0+2}:{L}{rr-1})").number_format = FMT_EUR
        totals_row(ws, rr, len(hdr))
        ws.conditional_formatting.add(f"D{r0+2}:D{rr-1}",
            DataBarRule(start_type="min", end_type="max", color="F4A6A6"))
        # chart: avoidable overpay (EUR) by country
        ch = BarChart(); ch.type = "col"; ch.title = "Avoidable overpay (EUR) by country"
        ch.height = 7.5; ch.width = 15; ch.legend = None
        data = Reference(ws, min_col=2, min_row=r0 + 1, max_row=rr - 1)
        cats = Reference(ws, min_col=1, min_row=r0 + 2, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "F8")
    ws.freeze_panes = "A3"; ws.sheet_view.showGridLines = False

    # --- Top actions detail sheet -------------------------------------------
    wa = wb.create_sheet("Top actions")
    hdr = ["Opportunity", "Country", "Detail", "Addressable EUR"]
    for j, h in enumerate(hdr, 1):
        wa.cell(1, j, h)
    style_header(wa, 1, len(hdr))
    rr = 2
    for a in s["top_actions"]:
        wa.cell(rr, 1, a["kind"]); wa.cell(rr, 2, a["country"])
        wa.cell(rr, 3, a["detail"])
        wa.cell(rr, 4, a["eur"]).number_format = FMT_EUR
        rr += 1
    band_rows(wa, 2, rr - 1, len(hdr))
    if rr > 2:
        wa.cell(rr, 1, "TOTAL addressable")
        wa.cell(rr, 4, f"=SUM(D2:D{rr-1})").number_format = FMT_EUR
        totals_row(wa, rr, len(hdr))
    set_widths(wa, [34, 12, 70, 16])
    wa.freeze_panes = "A2"; wa.sheet_view.showGridLines = False

    path = path or os.path.join(WORKDIR, f"Savings_Intelligence_{period}.xlsx")
    wb.save(path)
    return path


def overpay_review_workbook(period=None, supplier=None, path=None):
    """Supplier price-competitiveness review packet for `period` (optionally one
    `supplier`). Turns queries.q_savings_lines (the per-fuelling-day overpay detail)
    into an actionable Excel: a per-supplier SUMMARY sheet and a DETAIL sheet of the
    individual fuelling days that drove the overpay.

    Framing (printed on every sheet): this is a price-COMPETITIVENESS / negotiation
    review — "supplier X charged €Y more than the cheapest same-day, same-country
    rival across N fuellings" — NOT a contractual claim-back / debt. Evidence for
    renegotiation or steering volume, not money the supplier owes. NET EUR/L basis,
    final (VAT excluded, rebates applied). EUR totals HALF_UP (money.f2)."""
    con = connect()
    ps = _periods(con)
    period = period or (ps[0] if ps else None)
    if period is None:
        con.close()
        raise ValueError("no data loaded — nothing to report")
    lines = queries.q_savings_lines(con, period, supplier)
    con.close()

    gen = datetime.date.today().isoformat()
    caveat = ("Competitiveness review, NOT a contractual claim — supplier charged more "
              "than the cheapest same-day, same-country diesel rival. NET EUR/L, final "
              "(VAT excluded, rebates applied). "
              + (f"Supplier: {supplier}. " if supplier else "")
              + f"Generated {gen}.")
    wb = Workbook()

    # --- Summary sheet: per-supplier rollup ---------------------------------
    ws = wb.active; ws.title = "Summary"
    _title(ws, period, caveat, "A1:D1")
    set_widths(ws, [24, 14, 16, 18])
    # roll the detail lines up per supplier
    roll = {}
    for ln in lines:
        a = roll.setdefault(ln["supplier"], {"fuellings": 0, "litres": 0.0, "overpay": 0.0})
        a["fuellings"] += 1
        a["litres"] += ln["litres"] or 0
        a["overpay"] += ln["overpay_eur"] or 0
    ordered = sorted(roll.items(), key=lambda x: -x[1]["overpay"])
    r0 = 4
    ws.cell(r0, 1, "Most-overpaid supplier first").font = Font(bold=True, size=11)
    hdr = ["Supplier", "Fuellings", "Litres", "Total overpay EUR"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r0 + 1, j, h)
    style_header(ws, r0 + 1, len(hdr))
    rr = r0 + 2
    for sup, a in ordered:
        ws.cell(rr, 1, sup)
        ws.cell(rr, 2, a["fuellings"]).number_format = FMT_INT
        ws.cell(rr, 3, a["litres"]).number_format = FMT_INT
        ws.cell(rr, 4, money.f2(a["overpay"])).number_format = FMT_EUR
        rr += 1
    band_rows(ws, r0 + 2, rr - 1, len(hdr))
    if rr > r0 + 2:
        ws.cell(rr, 1, "TOTAL")
        ws.cell(rr, 2, f"=SUM(B{r0+2}:B{rr-1})").number_format = FMT_INT
        ws.cell(rr, 3, f"=SUM(C{r0+2}:C{rr-1})").number_format = FMT_INT
        ws.cell(rr, 4, f"=SUM(D{r0+2}:D{rr-1})").number_format = FMT_EUR
        totals_row(ws, rr, len(hdr))
        ws.conditional_formatting.add(f"D{r0+2}:D{rr-1}",
            DataBarRule(start_type="min", end_type="max", color="F4A6A6"))
        ch = BarChart(); ch.type = "col"; ch.title = "Total overpay (EUR) by supplier"
        ch.height = 7.5; ch.width = 15; ch.legend = None
        data = Reference(ws, min_col=4, min_row=r0 + 1, max_row=rr - 1)
        cats = Reference(ws, min_col=1, min_row=r0 + 2, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
        ws.add_chart(ch, "F4")
    ws.freeze_panes = "A3"; ws.sheet_view.showGridLines = False

    # --- Detail sheet: per fuelling-day -------------------------------------
    wd = wb.create_sheet("Detail")
    _title(wd, period, caveat, "A1:I1")
    set_widths(wd, [22, 13, 14, 12, 12, 12, 22, 12, 16])
    hdr = ["Supplier", "Date", "Country", "Litres", "This €/L", "Cheapest €/L",
           "Cheapest supplier", "Delta €/L", "Overpay EUR"]
    for j, h in enumerate(hdr, 1):
        wd.cell(4, j, h)
    style_header(wd, 4, len(hdr))
    rr = 5
    for ln in lines:
        wd.cell(rr, 1, ln["supplier"])
        wd.cell(rr, 2, ln["date"])
        wd.cell(rr, 3, ln["country"])
        wd.cell(rr, 4, ln["litres"] or 0).number_format = FMT_INT
        wd.cell(rr, 5, ln["eur_l"] or 0).number_format = FMT_PRICE
        wd.cell(rr, 6, ln["cheapest_eur_l"] or 0).number_format = FMT_PRICE
        wd.cell(rr, 7, ln["cheapest_supplier"])
        wd.cell(rr, 8, ln["delta_eur_l"] or 0).number_format = FMT_PRICE
        wd.cell(rr, 9, ln["overpay_eur"] or 0).number_format = FMT_EUR
        rr += 1
    band_rows(wd, 5, rr - 1, len(hdr))
    if rr > 5:
        wd.cell(rr, 1, "TOTAL")
        wd.cell(rr, 4, f"=SUM(D5:D{rr-1})").number_format = FMT_INT
        wd.cell(rr, 9, f"=SUM(I5:I{rr-1})").number_format = FMT_EUR
        totals_row(wd, rr, len(hdr))
    wd.freeze_panes = "A5"; wd.sheet_view.showGridLines = False

    if path is None:
        suffix = f"_{supplier}" if supplier else ""
        # mirror q_savings_lines: keep filename filesystem-safe
        safe = "".join(ch if ch.isalnum() else "_" for ch in suffix)
        path = os.path.join(WORKDIR, f"Overpay_Review_{period}{safe}.xlsx")
    wb.save(path)
    return path


def expense_report_workbook(period=None, entity=None, path=None):
    """Finance-facing company expense / cost-allocation report for `period` (optionally
    one `entity`). Turns the validated transactions into a NET / VAT / gross spend
    breakdown a finance team can export: a per-ENTITY summary (cost centre) with a
    per-product-group split, and a per-VEHICLE detail sheet with the effective NET €/L.

    BASIS (printed on the report): NET EUR, final (rebates applied); VAT shown
    separately; gross = net + VAT. EUR totals HALF_UP (money.f2)."""
    con = connect()
    ps = _periods(con)
    period = period or (ps[0] if ps else None)
    if period is None:
        con.close()
        raise ValueError("no data loaded — nothing to report")
    data = queries.q_expense(con, period, entity)
    con.close()

    gen = datetime.date.today().isoformat()
    caveat = ("Company fuel & toll expense report. NET EUR, final (rebates applied); "
              "VAT shown separately; gross = net + VAT. "
              + (f"Entity: {entity}. " if entity else "")
              + f"Generated {gen}.")
    wb = Workbook()

    # --- Summary sheet: per-entity table + per-product-group split -----------
    ws = wb.active; ws.title = "Summary"
    _title(ws, period, caveat, "A1:G1")
    set_widths(ws, [26, 12, 11, 14, 16, 16, 16])
    r0 = 4
    ws.cell(r0, 1, "Expense by entity (cost centre)").font = Font(bold=True, size=11)
    hdr = ["Entity", "Fuellings", "Vehicles", "Litres", "Net EUR", "VAT EUR", "Gross EUR"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r0 + 1, j, h)
    style_header(ws, r0 + 1, len(hdr))
    rr = r0 + 2
    for e in data["by_entity"]:
        ws.cell(rr, 1, e["entity"])
        ws.cell(rr, 2, e["fuellings"]).number_format = FMT_INT
        ws.cell(rr, 3, e["n_vehicles"]).number_format = FMT_INT
        ws.cell(rr, 4, e["litres"]).number_format = FMT_INT
        ws.cell(rr, 5, e["net_eur"]).number_format = FMT_EUR
        ws.cell(rr, 6, e["vat_eur"]).number_format = FMT_EUR
        ws.cell(rr, 7, e["gross_eur"]).number_format = FMT_EUR
        rr += 1
    band_rows(ws, r0 + 2, rr - 1, len(hdr))
    if rr > r0 + 2:
        ws.cell(rr, 1, "TOTAL")
        for col, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G")):
            ws.cell(rr, col, f"=SUM({L}{r0+2}:{L}{rr-1})").number_format = (
                FMT_INT if col in (2, 3, 4) else FMT_EUR)
        totals_row(ws, rr, len(hdr))
        ws.conditional_formatting.add(f"E{r0+2}:E{rr-1}",
            DataBarRule(start_type="min", end_type="max", color="9CC3E6"))

    # per-product-group split below the entity table
    p0 = rr + 3
    ws.cell(p0, 1, "Expense by product group").font = Font(bold=True, size=11)
    phdr = ["Product group", "Litres", "Net EUR", "VAT EUR", "Gross EUR"]
    for j, h in enumerate(phdr, 1):
        ws.cell(p0 + 1, j, h)
    style_header(ws, p0 + 1, len(phdr))
    pr = p0 + 2
    for g in data["by_product"]:
        ws.cell(pr, 1, g["product_group"])
        ws.cell(pr, 2, g["litres"]).number_format = FMT_INT
        ws.cell(pr, 3, g["net_eur"]).number_format = FMT_EUR
        ws.cell(pr, 4, g["vat_eur"]).number_format = FMT_EUR
        ws.cell(pr, 5, g["gross_eur"]).number_format = FMT_EUR
        pr += 1
    band_rows(ws, p0 + 2, pr - 1, len(phdr))
    if pr > p0 + 2:
        ws.cell(pr, 1, "TOTAL")
        for col, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
            ws.cell(pr, col, f"=SUM({L}{p0+2}:{L}{pr-1})").number_format = (
                FMT_INT if col == 2 else FMT_EUR)
        totals_row(ws, pr, len(phdr))
    ws.freeze_panes = "A3"; ws.sheet_view.showGridLines = False

    # --- Vehicles sheet: per-vehicle detail ---------------------------------
    wv = wb.create_sheet("Vehicles")
    _title(wv, period, caveat, "A1:I1")
    set_widths(wv, [22, 18, 12, 14, 16, 16, 16, 12, 12])
    hdr = ["Entity", "Vehicle", "Fuellings", "Litres", "Net EUR", "VAT EUR",
           "Gross EUR", "NET €/L", "Countries"]
    for j, h in enumerate(hdr, 1):
        wv.cell(4, j, h)
    style_header(wv, 4, len(hdr))
    rr = 5
    for v in data["by_vehicle"]:
        wv.cell(rr, 1, v["entity"])
        wv.cell(rr, 2, v["vehicle"])
        wv.cell(rr, 3, v["fuellings"]).number_format = FMT_INT
        wv.cell(rr, 4, v["litres"]).number_format = FMT_INT
        wv.cell(rr, 5, v["net_eur"]).number_format = FMT_EUR
        wv.cell(rr, 6, v["vat_eur"]).number_format = FMT_EUR
        wv.cell(rr, 7, v["gross_eur"]).number_format = FMT_EUR
        wv.cell(rr, 8, v["net_eur_l"]).number_format = FMT_PRICE
        wv.cell(rr, 9, v["n_countries"]).number_format = FMT_INT
        rr += 1
    band_rows(wv, 5, rr - 1, len(hdr))
    if rr > 5:
        wv.cell(rr, 1, "TOTAL")
        wv.cell(rr, 3, f"=SUM(C5:C{rr-1})").number_format = FMT_INT
        wv.cell(rr, 4, f"=SUM(D5:D{rr-1})").number_format = FMT_INT
        wv.cell(rr, 5, f"=SUM(E5:E{rr-1})").number_format = FMT_EUR
        wv.cell(rr, 6, f"=SUM(F5:F{rr-1})").number_format = FMT_EUR
        wv.cell(rr, 7, f"=SUM(G5:G{rr-1})").number_format = FMT_EUR
        totals_row(wv, rr, len(hdr))
    wv.freeze_panes = "A5"; wv.sheet_view.showGridLines = False

    if path is None:
        suffix = f"_{entity}" if entity else ""
        # mirror overpay_review_workbook: keep the filename filesystem-safe
        safe = "".join(ch if ch.isalnum() else "_" for ch in suffix)
        path = os.path.join(WORKDIR, f"Expense_Report_{period}{safe}.xlsx")
    wb.save(path)
    return path


# ---- ledger columns: (q_ledger key, human header) in export order --------------
_LEDGER_COLS = [
    ("date", "Date"), ("period", "Period"), ("entity", "Entity"),
    ("supplier", "Supplier"), ("country", "Country"), ("vehicle", "Vehicle"),
    ("station", "Station"), ("product", "Product"), ("product_group", "Product group"),
    ("qty", "Qty"), ("currency", "Currency"), ("net_local", "Net local"),
    ("vat_local", "VAT local"), ("gross_local", "Gross local"), ("net_eur", "Net EUR"),
    ("vat_eur", "VAT EUR"), ("gross_eur", "Gross EUR"), ("vat_rate_pct", "VAT rate %"),
    ("note", "Note"),
]


def accounting_ledger_csv(period=None, entity=None):
    """Accounting / ERP ledger export (CSV) — the decision-free first cut of the
    SAF-T/ERP-export capability. A clean, universally-importable transaction-level CSV
    (one row per transaction, no client chart-of-accounts, no country-specific SAF-T XML)
    that finance can derive any journal from or import into Xero/QuickBooks/DATEV/a
    spreadsheet. Read-only over the engine-owned product DB.

    BASIS: NET EUR, final (rebates applied); VAT shown separately; gross = net + VAT.

    The CSV is CLEAN — a single human-readable header row then one data row per
    `queries.q_ledger` record, with NO comment/preamble lines (many ERP importers choke
    on them) and numeric cells as plain numbers (no currency symbols/separators). Encoded
    utf-8-sig so Excel opens it cleanly.

    Returns (download_name, csv_bytes). Raises ValueError when no data/period."""
    import csv, io
    con = connect()
    ps = _periods(con)
    period = period or (ps[0] if ps else None)
    if period is None:
        con.close()
        raise ValueError("no data loaded — nothing to export")
    rows = queries.q_ledger(con, period, entity)
    con.close()

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in _LEDGER_COLS])
    for r in rows:
        w.writerow([r.get(k, "") for k, _ in _LEDGER_COLS])
    data = buf.getvalue().encode("utf-8-sig")

    suffix = f"_{entity}" if entity else ""
    safe = "".join(ch if ch.isalnum() else "_" for ch in suffix)
    download_name = f"Accounting_Ledger_{period}{safe}.csv"
    return download_name, data


def fee_report_workbook(claim, path=None):
    """One-sheet service-fee invoice / calculation for a single VAT claim. `claim`
    is a vat_applications row (dict). Fee = % of the refunded amount, or the minimum
    per claim when the % is lower. The settlement depends on where the refund was
    paid: to the customer (we invoice the fee) or to us (we deduct and remit net)."""
    from openpyxl.styles import Font
    ent = claim.get("entity") or ""
    ctry = claim.get("refund_country") or claim.get("country") or ""
    per = claim.get("ref_period") or claim.get("period") or ""
    amount = (claim.get("paid_amount") or claim.get("vat_eur") or 0) or 0
    pct = claim.get("fee_pct") or 0
    mn = money.f2(claim.get("fee_min") or 0)
    pct_fee = money.f2(pct / 100.0 * amount)
    fee = claim.get("fee_eur")
    fee = money.f2(fee) if fee is not None else max(pct_fee, mn)
    basis = "percent (% of refund)" if pct_fee >= mn else "minimum per claim"
    inv_no = claim.get("fee_invoice_no")
    payout = claim.get("payout_to") or "customer"
    net_to_customer = money.f2(amount - fee)
    wb = Workbook(); ws = wb.active; ws.title = "Fee invoice" if inv_no else "Fee calculation"
    heading = (f"Service fee invoice {inv_no}" if inv_no else "Service fee calculation")
    _title(ws, per, f"{heading} — {ent} / {ctry}", "A1:C1")
    set_widths(ws, [38, 18, 4])
    lines = [
        ("Customer", ent, None), ("Refund country", ctry, None),
        ("Claim period", per, None), ("Status", claim.get("status") or "", None),
    ]
    if inv_no:
        lines += [("Fee invoice no.", inv_no, None),
                  ("Invoice date", claim.get("fee_invoice_date") or "", None)]
    lines += [
        ("", "", None),
        ("Refunded VAT amount (EUR)", amount, FMT_EUR),
        (f"Fee at contract rate ({pct:g}%)", pct_fee, FMT_EUR),
        ("Minimum fee per claim (EUR)", mn, FMT_EUR),
        ("FEE CHARGED (EUR)", fee, FMT_EUR),
        ("Basis", basis, None),
        ("", "", None),
        ("Refund paid to", "OUR account" if payout == "us" else "CUSTOMER account", None),
    ]
    if payout == "us":
        lines += [("Less: our service fee (EUR)", fee, FMT_EUR),
                  ("NET REMITTED TO CUSTOMER (EUR)", net_to_customer, FMT_EUR)]
    else:
        lines += [("FEE PAYABLE BY CUSTOMER (EUR)", fee, FMT_EUR),
                  ("(remit to our account)", "", None)]
    lines += [
        ("", "", None),
        ("Submitted", claim.get("submitted_date") or claim.get("submitted") or "", None),
        ("Refund paid", claim.get("paid_date") or claim.get("paid") or "", None),
        ("Fee billed", claim.get("fee_billed_date") or "(charged when refund is paid)", None),
    ]
    r0 = 4
    for i, (lbl, val, fmt) in enumerate(lines):
        rr = r0 + i
        c1 = ws.cell(rr, 1, lbl)
        c2 = ws.cell(rr, 2, val if val is not None else "")
        if fmt:
            c2.number_format = fmt
        if str(lbl).startswith(("FEE CHARGED", "NET REMITTED", "FEE PAYABLE")):
            c1.font = Font(bold=True, size=12)
            c2.font = Font(bold=True, size=12, color=ACC)
    ws.sheet_view.showGridLines = False
    safe = f"Fee_{ent}_{ctry}_{per}".replace(" ", "_")
    safe = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in safe)
    path = path or os.path.join(WORKDIR, safe + ".xlsx")
    wb.save(path)
    return path


def claims_overview_workbook(overview, year, path=None):
    """VAT-refund submission readiness: 'Ready to submit' + 'Open claims' sheets."""
    wb = Workbook(); first = True
    # ready-to-submit
    ws = wb.active; ws.title = "Ready to submit"
    hdr = ["Entity", "Country", "Period", "VAT EUR", "Can submit?", "Blocking reasons"]
    for j, h in enumerate(hdr, 1):
        ws.cell(1, j, h)
    style_header(ws, 1, len(hdr))
    rr = 2
    for c in overview["to_submit"]:
        ws.cell(rr, 1, c["entity"]); ws.cell(rr, 2, c["country"]); ws.cell(rr, 3, c["period"])
        ws.cell(rr, 4, c["vat_eur"] or 0).number_format = FMT_EUR
        cell = ws.cell(rr, 5, "READY" if c["ready"] else "BLOCKED")
        cell.font = Font(bold=True, color=OKG if c["ready"] else BADR)
        ws.cell(rr, 6, "; ".join(c["issues"]))
        rr += 1
    band_rows(ws, 2, rr - 1, len(hdr))
    set_widths(ws, [24, 14, 12, 14, 12, 50]); ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    # open claims
    ws2 = wb.create_sheet("Open claims")
    hdr2 = ["Entity", "Country", "Period", "VAT EUR", "Status", "Submitted", "Age (days)"]
    for j, h in enumerate(hdr2, 1):
        ws2.cell(1, j, h)
    style_header(ws2, 1, len(hdr2))
    rr = 2
    for c in overview["open"]:
        ws2.cell(rr, 1, c["entity"]); ws2.cell(rr, 2, c["country"]); ws2.cell(rr, 3, c["period"])
        ws2.cell(rr, 4, c["vat_eur"] or 0).number_format = FMT_EUR
        ws2.cell(rr, 5, c["status"]); ws2.cell(rr, 6, c["submitted"] or "")
        ws2.cell(rr, 7, c["age_days"] if c["age_days"] != "" else "")
        rr += 1
    band_rows(ws2, 2, rr - 1, len(hdr2))
    set_widths(ws2, [24, 14, 12, 14, 12, 13, 11]); ws2.freeze_panes = "A2"; ws2.sheet_view.showGridLines = False
    path = path or os.path.join(WORKDIR, f"VAT_Claim_Readiness_{year}.xlsx")
    wb.save(path)
    return path


def fees_statement_workbook(rows, year, path=None):
    """Monthly fees statement: charged fees and net remittances per customer.
    `rows` is recovery_report() output. Only claims whose fee has been billed
    (the refund was paid) count. Sheet 1 aggregates by customer; sheet 2 lists
    every charged line. All EUR, VAT-excluded."""
    from openpyxl.styles import Font
    billed = [r for r in rows if r.get("fee_billed_date")]
    # aggregate per entity
    agg = {}
    detail = []
    for r in billed:
        ent = r.get("entity") or ""
        vat = r.get("vat_eur") or 0
        fee = r.get("fee_eur") or 0
        refund = r.get("paid_amount") or vat
        payout = r.get("payout_to") or "customer"
        receivable = fee if payout == "customer" else 0.0
        net = (refund - fee) if payout == "us" else refund
        pct_fee = (r.get("fee_pct") or 0) / 100.0 * vat
        basis = "percent" if pct_fee >= (r.get("fee_min") or 0) else "minimum"
        a = agg.setdefault(ent, {"n": 0, "vat": 0.0, "fee": 0.0, "recv": 0.0, "net": 0.0})
        a["n"] += 1; a["vat"] += vat; a["fee"] += fee
        a["recv"] += receivable; a["net"] += (net if payout == "us" else 0.0)
        detail.append((ent, r.get("country") or "", r.get("period") or "", vat, fee, basis,
                       payout, receivable, net, r.get("fee_invoice_no") or "",
                       r.get("fee_billed_date") or ""))

    wb = Workbook()
    ws = wb.active; ws.title = "Fees by customer"
    _title(ws, str(year), "Service fees statement — charged fees & net remittances", "A1:F1")
    hdr = ["Customer", "Claims", "VAT refunded", "Fees charged",
           "Receivable (invoiced)", "Net remitted (deducted)"]
    for j, h in enumerate(hdr, 1):
        ws.cell(3, j, h)
    style_header(ws, 3, len(hdr))
    rr = 4
    for ent in sorted(agg):
        a = agg[ent]
        ws.cell(rr, 1, ent)
        ws.cell(rr, 2, a["n"]).number_format = FMT_INT
        ws.cell(rr, 3, money.f2(a["vat"])).number_format = FMT_EUR
        ws.cell(rr, 4, money.f2(a["fee"])).number_format = FMT_EUR
        ws.cell(rr, 5, money.f2(a["recv"])).number_format = FMT_EUR
        ws.cell(rr, 6, money.f2(a["net"])).number_format = FMT_EUR
        rr += 1
    band_rows(ws, 4, rr - 1, len(hdr))
    if rr > 4:
        totals_row(ws, rr, len(hdr))
        ws.cell(rr, 1, "TOTAL").font = Font(bold=True)
        for col in (2, 3, 4, 5, 6):
            L = get_column_letter(col)
            ws.cell(rr, col).value = f"=SUM({L}4:{L}{rr-1})"
            ws.cell(rr, col).number_format = FMT_INT if col == 2 else FMT_EUR
    set_widths(ws, [26, 9, 16, 14, 18, 20]); ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    ws2 = wb.create_sheet("Fee detail")
    hdr2 = ["Customer", "Country", "Period", "VAT EUR", "Fee EUR", "Basis",
            "Refund paid to", "Receivable", "Net remitted", "Fee invoice", "Billed"]
    for j, h in enumerate(hdr2, 1):
        ws2.cell(1, j, h)
    style_header(ws2, 1, len(hdr2))
    rr = 2
    for d in detail:
        ent, ctry, per, vat, fee, basis, payout, recv, net, inv, billed_dt = d
        ws2.cell(rr, 1, ent); ws2.cell(rr, 2, ctry); ws2.cell(rr, 3, per)
        ws2.cell(rr, 4, money.f2(vat)).number_format = FMT_EUR
        ws2.cell(rr, 5, money.f2(fee)).number_format = FMT_EUR
        ws2.cell(rr, 6, basis)
        ws2.cell(rr, 7, "us" if payout == "us" else "customer")
        ws2.cell(rr, 8, money.f2(recv)).number_format = FMT_EUR
        ws2.cell(rr, 9, money.f2(net) if payout == "us" else "").number_format = FMT_EUR
        ws2.cell(rr, 10, inv); ws2.cell(rr, 11, billed_dt)
        rr += 1
    band_rows(ws2, 2, rr - 1, len(hdr2))
    set_widths(ws2, [24, 10, 10, 13, 12, 10, 13, 13, 14, 14, 12])
    ws2.freeze_panes = "A2"; ws2.sheet_view.showGridLines = False

    path = path or os.path.join(WORKDIR, f"VAT_Fees_Statement_{year}.xlsx")
    wb.save(path)
    return path


def receivables_forecast_workbook(fc, year, path=None):
    """VAT-receivable / financing-ready view (INTERNAL, data-only). `fc` is the dict
    from vat_refund.receivables_forecast(): per-claim receivable rows with the frozen
    fee and the ROUTE-AWARE settlement figures (refund receivable owed by the state,
    agency fee, customer net), the open-receivable aging, cycle-time medians, and the
    realization rate per refund country. All EUR, NET basis (VAT-excluded prices; the
    EUR columns are VAT amounts/refunds). The refund receivable and the agency fee are
    two SEPARATE cash flows — never summed across the 'customer'/'us' routes. No outward
    send — an export for the admin."""
    from openpyxl.styles import Font
    rows = fc.get("rows", [])
    wb = Workbook()
    ws = wb.active; ws.title = "Receivables"
    _title(ws, str(year), "VAT receivables & payout forecast — internal, NET basis", "A1:H1")
    hdr = ["Entity", "Country", "Period", "Status", "Payout route",
           "Refund receivable EUR (from state)", "Agency fee EUR (frozen)",
           "Customer net EUR", "Paid amount", "Submitted", "Approved", "Paid",
           "Age (days)", "Aging band"]
    for j, h in enumerate(hdr, 1):
        ws.cell(3, j, h)
    style_header(ws, 3, len(hdr))
    rr = 4
    for r in rows:
        route = r.get("route") or "customer"
        ws.cell(rr, 1, r.get("entity") or "")
        ws.cell(rr, 2, r.get("country") or "")
        ws.cell(rr, 3, r.get("period") or "")
        ws.cell(rr, 4, f"{r.get('status_code') or ''} {r.get('status_label') or ''}".strip())
        ws.cell(rr, 5, "deduct (to us)" if route == "us" else "direct (to customer)")
        ws.cell(rr, 6, money.f2(r.get("refund_receivable_eur") or 0)).number_format = FMT_EUR
        ws.cell(rr, 7, money.f2(r.get("fee_eur") or 0)).number_format = FMT_EUR
        ws.cell(rr, 8, money.f2(r.get("net_to_customer_eur") or 0)).number_format = FMT_EUR
        ws.cell(rr, 9, money.f2(r["paid_amount"]) if r.get("paid_amount") is not None else "").number_format = FMT_EUR
        ws.cell(rr, 10, r.get("submitted") or "")
        ws.cell(rr, 11, r.get("approved") or "")
        ws.cell(rr, 12, r.get("paid") or "")
        ws.cell(rr, 13, r["age_days"] if isinstance(r.get("age_days"), int) else "")
        ws.cell(rr, 14, r.get("aging_band") or "")
        rr += 1
    band_rows(ws, 4, rr - 1, len(hdr))
    set_widths(ws, [24, 11, 10, 18, 20, 22, 18, 16, 13, 11, 11, 11, 10, 11])
    ws.freeze_panes = "A4"; ws.sheet_view.showGridLines = False

    # Forecast / aging summary
    ws2 = wb.create_sheet("Forecast")
    forecast = fc.get("forecast", {})
    aging = fc.get("aging", {})
    ws2.cell(1, 1, "Open-receivable cash forecast (unpaid submitted/approved) — "
                   "two SEPARATE flows, never summed across routes").font = Font(bold=True)
    ws2.cell(3, 1, "Open claims"); ws2.cell(3, 2, forecast.get("open_count") or 0).number_format = FMT_INT
    ws2.cell(4, 1, "Open refund receivable EUR (from state)")
    ws2.cell(4, 2, money.f2(forecast.get("open_refund_receivable_eur") or 0)).number_format = FMT_EUR
    ws2.cell(5, 1, "Realization-weighted refund EUR")
    ws2.cell(5, 2, money.f2(forecast.get("open_weighted_refund_eur") or 0)).number_format = FMT_EUR
    ws2.cell(6, 1, "Open agency fee receivable EUR")
    ws2.cell(6, 2, money.f2(forecast.get("open_fee_receivable_eur") or 0)).number_format = FMT_EUR
    ws2.cell(8, 1, "Aging band").font = Font(bold=True)
    ws2.cell(8, 2, "Refund receivable EUR").font = Font(bold=True)
    ws2.cell(8, 3, "Count").font = Font(bold=True)
    by_band = aging.get("by_band", {})
    r = 9
    for band in ("0-30", "30-60", "60-90", "90+"):
        b = by_band.get(band, {})
        ws2.cell(r, 1, band)
        ws2.cell(r, 2, money.f2(b.get("eur") or 0)).number_format = FMT_EUR
        ws2.cell(r, 3, b.get("count") or 0).number_format = FMT_INT
        r += 1
    ws2.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws2.cell(r, 2, money.f2(aging.get("total_eur") or 0)).number_format = FMT_EUR
    ws2.cell(r, 3, aging.get("total_count") or 0).number_format = FMT_INT
    set_widths(ws2, [34, 20, 10])
    ws2.sheet_view.showGridLines = False

    # Cycle-time & realization by country
    ws3 = wb.create_sheet("Cycle time & realization")
    ct = fc.get("cycle_time", {})
    realization = fc.get("realization", {})
    ws3.cell(1, 1, "Median submitted→paid days, and realization (paid/claimed) by country").font = Font(bold=True)
    hdr3 = ["Country", "Median days (paid)", "Claimed EUR", "Paid EUR", "Realization %"]
    for j, h in enumerate(hdr3, 1):
        ws3.cell(3, j, h)
    style_header(ws3, 3, len(hdr3))
    by_ctry = ct.get("by_country", {})
    countries = sorted(set(by_ctry) | {c for c in realization if c != "overall"})
    rr = 4
    for c in countries:
        rz = realization.get(c, {})
        ws3.cell(rr, 1, c)
        med = by_ctry.get(c)
        ws3.cell(rr, 2, med if med is not None else "")
        ws3.cell(rr, 3, money.f2(rz.get("claimed") or 0)).number_format = FMT_EUR
        ws3.cell(rr, 4, money.f2(rz.get("paid") or 0)).number_format = FMT_EUR
        ws3.cell(rr, 5, (rz.get("rate") if rz.get("rate") is not None else "")).number_format = FMT_PCT
        rr += 1
    band_rows(ws3, 4, rr - 1, len(hdr3))
    ovr = realization.get("overall", {})
    ws3.cell(rr, 1, "OVERALL").font = Font(bold=True)
    ws3.cell(rr, 2, ct.get("overall") if ct.get("overall") is not None else "")
    ws3.cell(rr, 3, money.f2(ovr.get("claimed") or 0)).number_format = FMT_EUR
    ws3.cell(rr, 4, money.f2(ovr.get("paid") or 0)).number_format = FMT_EUR
    ws3.cell(rr, 5, (ovr.get("rate") if ovr.get("rate") is not None else "")).number_format = FMT_PCT
    set_widths(ws3, [16, 18, 14, 14, 14])
    ws3.freeze_panes = "A4"; ws3.sheet_view.showGridLines = False

    path = path or os.path.join(WORKDIR, f"VAT_Receivables_Forecast_{year}.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys
    per = sys.argv[1] if len(sys.argv) > 1 else None
    p = summary_workbook(per)
    print("wrote", p)
