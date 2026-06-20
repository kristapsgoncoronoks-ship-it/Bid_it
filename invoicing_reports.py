"""
INVOICING REPORTS — Phase 5 of the sales-invoicing module: a READ-ONLY reporting suite
over the invoicing data (issued invoices, lines, payments, credit notes from Phases 1-4).

This module NEVER mutates the invoicing data — it only READS it (via invoicing.connect()
and the invoicing.py read helpers) and reuses the Phase 1-4 money/derivation logic
(`compute_*`, `outstanding`, `credited_total`, `display_status`, `_age_bucket`). No figure
is forked: the AR aging reuses invoicing.accounts_receivable()'s exact math, and the
VAT-output / revenue / statement reports sum the SAME stored line/payment/credit figures.

FOUR reports, each with on-screen data + Excel (.xlsx, openpyxl) + PDF (HTML→wkhtmltopdf,
Latvian-capable) export:

  1. VAT OUTPUT (PVN) — output VAT owed for a period, grouped by VAT rate (21/12/5/0%) +
     reverse-charge + exempt SEPARATE lines, with CREDIT NOTES SUBTRACTED. Tax point =
     the issue date. Drafts excluded. The per-rate VAT sums to the total output VAT.
  2. SALES / REVENUE — NET revenue over a period, by month / by customer / by service,
     issued invoices minus credit notes.
  3. CUSTOMER STATEMENT — a per-customer account statement (invoices as debits; credit
     notes + payments as credits) with a running + opening/closing balance.
  4. AR / AGING — outstanding by customer with aging buckets, reusing
     invoicing.accounts_receivable() (NOT a fork).

BASIS: all amounts are EUR; revenue/VAT figures are on a NET (VAT-excluded) basis except
where a GROSS column is explicitly labelled. Money is summed via money.fsum (cents-exact,
ROUND_HALF_UP) — never a bare round(). The tax point for the VAT-output and revenue
reports is the ISSUE date (the invoice/credit-note becomes reportable when issued).

PERIOD MODEL: a report covers a calendar period — a month (YYYY-MM), a quarter
(YYYY-Qn), or a year (YYYY) — resolved to an inclusive [start, end] ISO-date window by
`period_window()`. Filtering is by ISSUE date, consistently across the reports.
"""
import datetime
import io

import applog
import money

log = applog.get("invoicing_reports")


# ============================================================ shared period helper
def period_window(kind, year, sub=None):
    """Resolve a (kind, year, sub) period selection to an inclusive ISO date window
    (start, end, label).

      kind='month'   sub=1..12       -> that calendar month
      kind='quarter' sub=1..4        -> that calendar quarter
      kind='year'    sub ignored     -> the whole year

    Returns (start_iso, end_iso, label). Raises ValueError on a bad selection so the route
    can surface a clear error (the caller validates the inputs)."""
    try:
        y = int(year)
    except (TypeError, ValueError):
        raise ValueError("a valid year is required")
    if kind == "year":
        return f"{y:04d}-01-01", f"{y:04d}-12-31", f"{y:04d}"
    if kind == "quarter":
        try:
            q = int(sub)
        except (TypeError, ValueError):
            raise ValueError("a valid quarter (1-4) is required")
        if q < 1 or q > 4:
            raise ValueError("quarter must be 1-4")
        sm = (q - 1) * 3 + 1
        em = sm + 2
        start = datetime.date(y, sm, 1)
        end = _month_end(y, em)
        return start.isoformat(), end.isoformat(), f"{y:04d}-Q{q}"
    if kind == "month":
        try:
            m = int(sub)
        except (TypeError, ValueError):
            raise ValueError("a valid month (1-12) is required")
        if m < 1 or m > 12:
            raise ValueError("month must be 1-12")
        start = datetime.date(y, m, 1)
        end = _month_end(y, m)
        return start.isoformat(), end.isoformat(), f"{y:04d}-{m:02d}"
    raise ValueError("period kind must be month, quarter or year")


def _month_end(year, month):
    """Last calendar day of (year, month) as a date."""
    if month == 12:
        nxt = datetime.date(year + 1, 1, 1)
    else:
        nxt = datetime.date(year, month + 1, 1)
    return nxt - datetime.timedelta(days=1)


# ============================================================ data access (read-only)
def _issued_docs(start, end, doc_type=None):
    """Every ISSUED (non-draft) invoice/credit-note whose ISSUE DATE is in [start, end],
    tenant-scoped, with the customer name joined. doc_type=None returns both. Read-only;
    never raises -> []. Reuses invoicing.connect()/scope_clause so this is tenant-safe.

    PROFORMA / QUOTE are EXCLUDED — they are not tax invoices and must never be counted in
    the VAT-output / revenue reports (they carry no legal amount / no output VAT)."""
    import invoicing
    import tenancy
    frag, tp = tenancy.scope_clause(column="i.tenant_id")
    where = ["i.status<>'draft'", "i.issue_date IS NOT NULL",
             "i.issue_date>=?", "i.issue_date<=?",
             "i.doc_type IN ('invoice','credit_note')"]
    params = [start, end]
    if doc_type is not None:
        where.append("i.doc_type=?")
        params.append(doc_type)
    sql = ("SELECT i.*, c.name AS customer_name FROM invoices i "
           "LEFT JOIN bill_customers c ON c.id=i.customer_id "
           "WHERE " + " AND ".join(where) + frag
           + " ORDER BY i.issue_date, i.id")
    try:
        con = invoicing.connect()
        try:
            rows = con.execute(sql, [*params, *tp]).fetchall()
        finally:
            con.close()
        return [invoicing._inv_dict(r) for r in rows]
    except Exception as e:
        log.warning("_issued_docs(%s,%s) failed: %s", start, end, e)
        return []


# ============================================================ 1. VAT OUTPUT (PVN)
# OUTPUT VAT the company OWES for a period. Tax point = the ISSUE date. An ISSUED invoice
# ADDS taxable supplies + output VAT; an ISSUED credit note SUBTRACTS (it reduces output
# VAT for the period in which the credit note is issued). Drafts carry no legal amount and
# are excluded. Reverse-charge (AE) and 0%/exempt supplies are SHOWN SEPARATELY (no output
# VAT but reportable). The per-rate VAT sums to the grand-total output VAT (the tie-out).
#
# RATE GROUPING toward the Latvian PVN deklarācija: a line is classified into ONE bucket:
#   - reverse_charge invoice  -> 'reverse_charge' (AE): net reportable, VAT 0
#   - vat_rate == 0           -> 'zero' (0% / exempt): net reportable, VAT 0
#   - else                    -> the standard/reduced rate bucket keyed by the rate
# (21% / 12% / 5% are the Latvian presets; a custom rate gets its own bucket.)

# Standard LV PVN-rate buckets, in display order (most reports group to these).
LV_RATE_ORDER = (0.21, 0.12, 0.05)


def _rate_bucket_key(line, reverse_charge):
    """Classify a stored line into a VAT-output bucket KEY:
      'reverse_charge' | 'zero' | a float rate (0.21/0.12/0.05/custom). Pure."""
    if reverse_charge:
        return "reverse_charge"
    rate = money.D(line.get("vat_rate") or 0)
    if rate == 0:
        return "zero"
    return float(rate)


def vat_output_report(start, end, label=""):
    """The OUTPUT-VAT (PVN) report for a period [start, end]. Read-only; never raises.

    Returns:
      {"label", "start", "end",
       "rows": [{"key", "rate"(float|None), "kind"('standard'|'zero'|'reverse_charge'),
                 "label", "net"(float), "vat"(float)}],   # one per rate-bucket, ordered
       "net_total", "vat_total",        # grand totals (NET supplies, output VAT)
       "invoices_net", "invoices_vat",  # the positive (invoice) contribution
       "credits_net", "credits_vat"}    # the credit-note contribution (already subtracted)

    TIE-OUT: sum(rows[i].vat) == vat_total (cents-exact); a credit note REDUCES both the
    bucket net and its VAT for the period. Reverse-charge + zero rows carry vat==0."""
    import invoicing
    docs = _issued_docs(start, end)
    # bucket -> {"net":[...signed...], "vat":[...signed...]}
    buckets = {}
    inv_net, inv_vat, cr_net, cr_vat = [], [], [], []
    for d in docs:
        is_credit = d.get("doc_type") == invoicing.DOC_CREDIT_NOTE
        sign = -1 if is_credit else 1
        rc = bool(d.get("reverse_charge"))
        for ln in invoicing.get_lines(d["id"]):
            key = _rate_bucket_key(ln, rc)
            net = sign * money.q2(ln.get("line_net") or 0)
            vat = sign * money.q2(ln.get("line_vat") or 0)
            b = buckets.setdefault(key, {"net": [], "vat": []})
            b["net"].append(float(net))
            b["vat"].append(float(vat))
            if is_credit:
                cr_net.append(float(-net))   # store the positive credit magnitude
                cr_vat.append(float(-vat))
            else:
                inv_net.append(float(net))
                inv_vat.append(float(vat))

    rows = []
    # ordered: the standard LV rates first, then any custom rate (desc), then zero, then AE.
    def _order(key):
        if key == "reverse_charge":
            return (3, 0.0)
        if key == "zero":
            return (2, 0.0)
        if key in LV_RATE_ORDER:
            return (0, -float(key))
        return (1, -float(key))

    for key in sorted(buckets, key=_order):
        b = buckets[key]
        net = money.fsum(b["net"])
        vat = money.fsum(b["vat"])
        if key == "reverse_charge":
            kind, rate, lbl = "reverse_charge", None, "Reverse charge (AE)"
        elif key == "zero":
            kind, rate, lbl = "zero", 0.0, "0% / exempt"
        else:
            kind, rate = "standard", float(key)
            lbl = f"{rate * 100:g}%"
        rows.append({"key": key, "rate": rate, "kind": kind, "label": lbl,
                     "net": net, "vat": vat})

    net_total = money.fsum([r["net"] for r in rows])
    vat_total = money.fsum([r["vat"] for r in rows])
    return {"label": label, "start": start, "end": end, "rows": rows,
            "net_total": net_total, "vat_total": vat_total,
            "invoices_net": money.fsum(inv_net), "invoices_vat": money.fsum(inv_vat),
            "credits_net": money.fsum(cr_net), "credits_vat": money.fsum(cr_vat)}


# ============================================================ 2. SALES / REVENUE
# NET revenue (VAT excluded) over a period: issued invoices MINUS issued credit notes. Tax
# point = the ISSUE date. Broken down BY month (trend), BY customer, and BY service (the
# line description). Counts + totals. Reverse-charge / zero-rated supplies still count as
# revenue (they are NET sales, just with no output VAT).

def revenue_report(start, end, label=""):
    """The sales/revenue report for [start, end]. Read-only; never raises.

    Returns:
      {"label","start","end",
       "by_month":   [{"month":'YYYY-MM',"net":..,"invoices":n,"credits":n}],
       "by_customer":[{"customer":name,"net":..,"invoices":n,"credits":n}],
       "by_service": [{"service":desc,"net":..,"lines":n}],
       "net_total", "invoice_count", "credit_count"}

    NET is the signed sum (invoice net positive, credit-note net negative). 'invoices' /
    'credits' are document counts; 'lines' is a line count for the service breakdown."""
    import invoicing
    docs = _issued_docs(start, end)
    months, custs, svcs = {}, {}, {}
    inv_count = cr_count = 0
    for d in docs:
        is_credit = d.get("doc_type") == invoicing.DOC_CREDIT_NOTE
        sign = -1 if is_credit else 1
        net = float(sign * money.q2(d.get("net_total") or 0))
        if is_credit:
            cr_count += 1
        else:
            inv_count += 1
        mo = str(d.get("issue_date"))[:7]
        m = months.setdefault(mo, {"net": [], "invoices": 0, "credits": 0})
        m["net"].append(net)
        m["credits" if is_credit else "invoices"] += 1
        cname = d.get("customer_name") or "—"
        cb = custs.setdefault(cname, {"net": [], "invoices": 0, "credits": 0})
        cb["net"].append(net)
        cb["credits" if is_credit else "invoices"] += 1
        for ln in invoicing.get_lines(d["id"]):
            desc = (ln.get("description") or "").strip() or "(no description)"
            lnet = float(sign * money.q2(ln.get("line_net") or 0))
            sb = svcs.setdefault(desc, {"net": [], "lines": 0})
            sb["net"].append(lnet)
            sb["lines"] += 1

    by_month = [{"month": mo, "net": money.fsum(v["net"]),
                 "invoices": v["invoices"], "credits": v["credits"]}
                for mo, v in sorted(months.items())]
    by_customer = sorted(
        [{"customer": c, "net": money.fsum(v["net"]),
          "invoices": v["invoices"], "credits": v["credits"]}
         for c, v in custs.items()],
        key=lambda r: r["net"], reverse=True)
    by_service = sorted(
        [{"service": s, "net": money.fsum(v["net"]), "lines": v["lines"]}
         for s, v in svcs.items()],
        key=lambda r: r["net"], reverse=True)
    net_total = money.fsum([r["net"] for r in by_month])
    return {"label": label, "start": start, "end": end,
            "by_month": by_month, "by_customer": by_customer, "by_service": by_service,
            "net_total": net_total, "invoice_count": inv_count, "credit_count": cr_count}


# ============================================================ 3. CUSTOMER STATEMENT
# A per-customer account statement: a chronological ledger of invoices (DEBITS, gross),
# credit notes (CREDITS, gross) and payments (CREDITS) with a RUNNING balance, plus the
# opening balance (the customer's net position from documents/payments BEFORE the window)
# and the closing balance. The closing balance == opening + sum(debits) − sum(credits).
#
# A DEBIT increases what the customer owes (an invoice); a CREDIT decreases it (a credit
# note or a payment). GROSS basis (a statement is what the customer actually owes/pays).

def _customer_payments(customer_id, start=None, end=None):
    """All payments against the customer's ISSUED invoices (optionally date-bounded by
    paid_date), tenant-scoped, oldest first. Read-only -> []. Each row:
    {paid_date, amount, invoice_id, number, reference}."""
    import invoicing
    import tenancy
    frag, tp = tenancy.scope_clause(column="p.tenant_id")
    where = ["i.customer_id=?", "i.status<>'draft'",
             "i.doc_type IN ('invoice','credit_note')"]
    params = [customer_id]
    if start:
        where.append("p.paid_date>=?")
        params.append(start)
    if end:
        where.append("p.paid_date<=?")
        params.append(end)
    sql = ("SELECT p.paid_date, p.amount, p.invoice_id, p.reference, i.number "
           "FROM invoice_payments p JOIN invoices i ON i.id=p.invoice_id "
           "WHERE " + " AND ".join(where) + frag
           + " ORDER BY p.paid_date, p.id")
    try:
        con = invoicing.connect()
        try:
            rows = con.execute(sql, [*params, *tp]).fetchall()
        finally:
            con.close()
        return [dict(r) for r in rows]
    except Exception as e:
        log.warning("_customer_payments(%s) failed: %s", customer_id, e)
        return []


def _customer_docs(customer_id, start=None, end=None):
    """ISSUED invoices + credit notes for a customer (optionally issue-date bounded),
    tenant-scoped, by issue date. Read-only -> []."""
    import invoicing
    import tenancy
    frag, tp = tenancy.scope_clause(column="i.tenant_id")
    where = ["i.customer_id=?", "i.status<>'draft'", "i.issue_date IS NOT NULL",
             "i.doc_type IN ('invoice','credit_note')"]
    params = [customer_id]
    if start:
        where.append("i.issue_date>=?")
        params.append(start)
    if end:
        where.append("i.issue_date<=?")
        params.append(end)
    sql = ("SELECT i.* FROM invoices i WHERE " + " AND ".join(where) + frag
           + " ORDER BY i.issue_date, i.id")
    try:
        con = invoicing.connect()
        try:
            rows = con.execute(sql, [*params, *tp]).fetchall()
        finally:
            con.close()
        return [invoicing._inv_dict(r) for r in rows]
    except Exception as e:
        log.warning("_customer_docs(%s) failed: %s", customer_id, e)
        return []


def _opening_balance(customer_id, before):
    """The customer's balance owed STRICTLY BEFORE the ISO date `before`: sum of invoice
    gross − credit-note gross (by issue date) − payments (by paid date), cents-exact. A
    blank `before` yields 0 (no opening period). Pure-ish read."""
    import invoicing
    if not before:
        return 0.0
    # documents issued before the window
    debits, credits = [], []
    for d in _customer_docs(customer_id, end=None):
        if str(d.get("issue_date"))[:10] >= before:
            continue
        gross = float(money.q2(d.get("gross_total") or 0))
        if d.get("doc_type") == invoicing.DOC_CREDIT_NOTE:
            credits.append(gross)
        else:
            debits.append(gross)
    for p in _customer_payments(customer_id, end=None):
        if str(p.get("paid_date"))[:10] >= before:
            continue
        credits.append(float(money.q2(p.get("amount") or 0)))
    return float(money.D(money.fsum(debits)) - money.D(money.fsum(credits)))


def customer_statement(customer_id, start, end, label=""):
    """An account statement for one customer over [start, end]. Read-only; never raises.

    Returns:
      {"label","start","end","customer":{...},
       "opening_balance", "closing_balance", "outstanding_total",
       "lines": [{"date","type"('invoice'|'credit_note'|'payment'),
                  "ref","debit","credit","balance"}]}

    A DEBIT is an invoice gross; a CREDIT is a credit-note gross or a payment. The running
    `balance` starts at opening_balance and applies each entry in date order; the closing
    balance == opening + Σdebits − Σcredits. outstanding_total is the customer's CURRENT
    total receivable (across all time, via invoicing.outstanding) — what they still owe."""
    import invoicing
    cust = invoicing.get_customer(customer_id) or {"id": customer_id, "name": "—"}
    opening = _opening_balance(customer_id, start)

    entries = []
    for d in _customer_docs(customer_id, start=start, end=end):
        gross = float(money.q2(d.get("gross_total") or 0))
        if d.get("doc_type") == invoicing.DOC_CREDIT_NOTE:
            entries.append({"date": str(d.get("issue_date"))[:10], "type": "credit_note",
                            "ref": d.get("number") or "—", "debit": 0.0, "credit": gross,
                            "_seq": 0})
        else:
            entries.append({"date": str(d.get("issue_date"))[:10], "type": "invoice",
                            "ref": d.get("number") or "—", "debit": gross, "credit": 0.0,
                            "_seq": 0})
    for p in _customer_payments(customer_id, start=start, end=end):
        amt = float(money.q2(p.get("amount") or 0))
        ref = p.get("number") or p.get("reference") or "—"
        entries.append({"date": str(p.get("paid_date"))[:10], "type": "payment",
                        "ref": ref, "debit": 0.0, "credit": amt, "_seq": 1})

    # date order; within a date, documents (seq 0) before payments (seq 1) for a readable
    # running balance.
    entries.sort(key=lambda e: (e["date"], e["_seq"]))
    bal = money.D(opening)
    lines = []
    for e in entries:
        bal = bal + money.D(e["debit"]) - money.D(e["credit"])
        lines.append({"date": e["date"], "type": e["type"], "ref": e["ref"],
                      "debit": e["debit"], "credit": e["credit"],
                      "balance": float(money.q2(bal))})
    closing = float(money.q2(bal))

    # the customer's CURRENT total outstanding (all open invoices), for the header.
    outstanding_total = money.fsum(
        [invoicing.outstanding(d) for d in _customer_docs(customer_id)
         if d.get("doc_type") != invoicing.DOC_CREDIT_NOTE])
    return {"label": label, "start": start, "end": end, "customer": cust,
            "opening_balance": float(money.q2(opening)), "closing_balance": closing,
            "outstanding_total": outstanding_total, "lines": lines}


# ============================================================ 4. AR / AGING (by customer)
# Outstanding by CUSTOMER with aging buckets. This does NOT fork the math: it consumes
# invoicing.accounts_receivable() (which already derives outstanding/display_status/the
# aging bucket per invoice via the Phase-3 logic) and re-groups its per-invoice rows BY
# CUSTOMER. The grand total == sum of the per-invoice outstanding == AR total_outstanding.

def ar_aging_report(today=None, label=""):
    """The AR aging report grouped BY CUSTOMER for the buckets
    current / 1-30 / 31-60 / 60+. Read-only; never raises.

    Returns:
      {"label","as_of",
       "rows": [{"customer", "current","b1_30","b31_60","b60p","total","overdue"}],
       "buckets": {bucket: outstanding}, "total_outstanding", "total_overdue"}

    Each customer row's bucket amounts are the sum of that customer's open-invoice
    outstanding in each aging bucket; OVERDUE = the past-due buckets (1-30 + 31-60 + 60+).
    Reuses invoicing.accounts_receivable so the totals tie to the AR view exactly."""
    import invoicing
    ref = today or datetime.date.today().isoformat()
    ar = invoicing.accounts_receivable(today=ref)
    # bucket key -> our column name
    colmap = {"current": "current", "1-30": "b1_30", "31-60": "b31_60", "60+": "b60p"}
    by_cust = {}
    for inv in ar.get("rows", []):
        cname = inv.get("customer_name") or "—"
        col = colmap.get(inv.get("bucket"), "current")
        owed = float(inv.get("outstanding") or 0)
        c = by_cust.setdefault(cname, {b: [] for b in
                                       ("current", "b1_30", "b31_60", "b60p")})
        c[col].append(owed)
    rows = []
    for cname in sorted(by_cust):
        c = by_cust[cname]
        cur = money.fsum(c["current"])
        b1 = money.fsum(c["b1_30"])
        b2 = money.fsum(c["b31_60"])
        b3 = money.fsum(c["b60p"])
        total = money.fsum([cur, b1, b2, b3])
        overdue = money.fsum([b1, b2, b3])
        rows.append({"customer": cname, "current": cur, "b1_30": b1, "b31_60": b2,
                     "b60p": b3, "total": total, "overdue": overdue})
    rows.sort(key=lambda r: r["total"], reverse=True)
    buckets = {
        "current": money.fsum([r["current"] for r in rows]),
        "1-30": money.fsum([r["b1_30"] for r in rows]),
        "31-60": money.fsum([r["b31_60"] for r in rows]),
        "60+": money.fsum([r["b60p"] for r in rows]),
    }
    total_outstanding = money.fsum([r["total"] for r in rows])
    total_overdue = money.fsum([r["overdue"] for r in rows])
    return {"label": label, "as_of": ref, "rows": rows, "buckets": buckets,
            "total_outstanding": total_outstanding, "total_overdue": total_overdue}


# ============================================================ EXCEL EXPORT (openpyxl)
# Real .xlsx workbooks built with openpyxl, matching the app's existing reports.py style
# (a dark title band, a styled header row, banded rows, EUR number formats, a totals row).
# Returned as BYTES (the route streams them via send_file) — no temp file churn. EVERY
# free-text DB value is run through reports._formula_safe (CSV/spreadsheet formula-injection
# neutraliser) before it is written. NET-EUR basis is stated in each report's subtitle.

FMT_EUR = "#,##0.00"
FMT_INT = "#,##0"
FMT_PCT = "0%"


def _xlsx_bytes(wb):
    """Serialize an openpyxl workbook to bytes (no temp file)."""
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _safe(v):
    import reports
    return reports._formula_safe(v)


def vat_output_workbook(rep):
    """Excel workbook (bytes) for the VAT-output report `rep` (vat_output_report output)."""
    from openpyxl import Workbook
    import reports
    wb = Workbook()
    caveat = ("Output VAT (PVN) by rate. Tax point = issue date; drafts excluded; "
              "credit notes subtracted. Amounts in EUR, NET (taxable) basis; VAT shown "
              "separately. Reverse-charge and 0%/exempt supplies carry no output VAT.")
    ws = wb.active
    ws.title = "VAT output"
    reports._title(ws, rep.get("label", ""), caveat, "A1:C1")
    reports.set_widths(ws, [34, 20, 20])
    hdr = ["VAT rate", "Taxable net (EUR)", "Output VAT (EUR)"]
    for j, h in enumerate(hdr, 1):
        ws.cell(4, j, h)
    reports.style_header(ws, 4, len(hdr))
    r = 5
    for row in rep["rows"]:
        ws.cell(r, 1, _safe(row["label"]))
        ws.cell(r, 2, row["net"]).number_format = FMT_EUR
        ws.cell(r, 3, row["vat"]).number_format = FMT_EUR
        r += 1
    reports.band_rows(ws, 5, r - 1, len(hdr))
    if r > 5:
        ws.cell(r, 1, "TOTAL output VAT")
        ws.cell(r, 2, f"=SUM(B5:B{r-1})").number_format = FMT_EUR
        ws.cell(r, 3, f"=SUM(C5:C{r-1})").number_format = FMT_EUR
        reports.totals_row(ws, r, len(hdr))
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    return _xlsx_bytes(wb)


def revenue_workbook(rep):
    """Excel workbook (bytes) for the revenue report (by month / customer / service)."""
    from openpyxl import Workbook
    import reports
    wb = Workbook()
    caveat = ("Sales / revenue, NET (VAT-excluded) EUR. Tax point = issue date; drafts "
              "excluded; credit notes subtracted. Counts are document/line counts.")

    def _section(ws, r0, title, headers, widths, rows, total_col):
        ws.cell(r0, 1, title).font = reports.Font(bold=True, size=11)
        for j, h in enumerate(headers, 1):
            ws.cell(r0 + 1, j, h)
        reports.style_header(ws, r0 + 1, len(headers))
        rr = r0 + 2
        for row in rows:
            for j, (val, fmt) in enumerate(row, 1):
                c = ws.cell(rr, j, _safe(val) if fmt is None else val)
                if fmt:
                    c.number_format = fmt
            rr += 1
        reports.band_rows(ws, r0 + 2, rr - 1, len(headers))
        if rr > r0 + 2:
            ws.cell(rr, 1, "TOTAL")
            L = reports.get_column_letter(total_col)
            ws.cell(rr, total_col, f"=SUM({L}{r0+2}:{L}{rr-1})").number_format = FMT_EUR
            reports.totals_row(ws, rr, len(headers))
        return rr + 3

    ws = wb.active
    ws.title = "Revenue"
    reports._title(ws, rep.get("label", ""), caveat, "A1:D1")
    reports.set_widths(ws, [34, 18, 12, 12])
    r0 = _section(ws, 4, "Revenue by month",
                  ["Month", "Net EUR", "Invoices", "Credit notes"],
                  None,
                  [[(m["month"], None), (m["net"], FMT_EUR),
                    (m["invoices"], FMT_INT), (m["credits"], FMT_INT)]
                   for m in rep["by_month"]], 2)
    r0 = _section(ws, r0, "Revenue by customer",
                  ["Customer", "Net EUR", "Invoices", "Credit notes"],
                  None,
                  [[(c["customer"], None), (c["net"], FMT_EUR),
                    (c["invoices"], FMT_INT), (c["credits"], FMT_INT)]
                   for c in rep["by_customer"]], 2)
    _section(ws, r0, "Revenue by service",
             ["Service / description", "Net EUR", "Lines", ""],
             None,
             [[(s["service"], None), (s["net"], FMT_EUR), (s["lines"], FMT_INT), ("", None)]
              for s in rep["by_service"]], 2)
    ws.sheet_view.showGridLines = False
    return _xlsx_bytes(wb)


def customer_statement_workbook(rep):
    """Excel workbook (bytes) for a customer statement (the ledger + balances)."""
    from openpyxl import Workbook
    import reports
    wb = Workbook()
    cust = rep.get("customer") or {}
    caveat = ("Statement of account — " + str(cust.get("name") or "—")
              + ". GROSS EUR. Debit = invoice; credit = credit note / payment.")
    ws = wb.active
    ws.title = "Statement"
    reports._title(ws, rep.get("label", ""), caveat, "A1:E1")
    reports.set_widths(ws, [14, 16, 22, 16, 16])
    ws.cell(4, 1, "Opening balance")
    ws.cell(4, 5, rep["opening_balance"]).number_format = FMT_EUR
    reports.totals_row(ws, 4, 5)
    hdr = ["Date", "Type", "Reference", "Debit", "Credit", "Balance"]
    reports.set_widths(ws, [14, 16, 22, 16, 16, 16])
    for j, h in enumerate(hdr, 1):
        ws.cell(6, j, h)
    reports.style_header(ws, 6, len(hdr))
    r = 7
    for ln in rep["lines"]:
        ws.cell(r, 1, _safe(ln["date"]))
        ws.cell(r, 2, _safe(ln["type"]))
        ws.cell(r, 3, _safe(ln["ref"]))
        ws.cell(r, 4, ln["debit"] or None).number_format = FMT_EUR
        ws.cell(r, 5, ln["credit"] or None).number_format = FMT_EUR
        ws.cell(r, 6, ln["balance"]).number_format = FMT_EUR
        r += 1
    reports.band_rows(ws, 7, r - 1, len(hdr))
    ws.cell(r, 1, "Closing balance")
    ws.cell(r, 6, rep["closing_balance"]).number_format = FMT_EUR
    reports.totals_row(ws, r, len(hdr))
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False
    return _xlsx_bytes(wb)


def ar_aging_workbook(rep):
    """Excel workbook (bytes) for the AR aging-by-customer report."""
    from openpyxl import Workbook
    import reports
    wb = Workbook()
    caveat = ("Accounts receivable aging by customer, as of " + str(rep.get("as_of", ""))
              + ". Outstanding EUR (gross − payments − credits). Overdue = past-due "
              "buckets (1-30 + 31-60 + 60+).")
    ws = wb.active
    ws.title = "AR aging"
    reports._title(ws, rep.get("as_of", ""), caveat, "A1:G1")
    reports.set_widths(ws, [30, 15, 15, 15, 15, 16, 16])
    hdr = ["Customer", "Current", "1-30 days", "31-60 days", "60+ days",
           "Total outstanding", "Overdue"]
    for j, h in enumerate(hdr, 1):
        ws.cell(4, j, h)
    reports.style_header(ws, 4, len(hdr))
    r = 5
    for row in rep["rows"]:
        ws.cell(r, 1, _safe(row["customer"]))
        ws.cell(r, 2, row["current"]).number_format = FMT_EUR
        ws.cell(r, 3, row["b1_30"]).number_format = FMT_EUR
        ws.cell(r, 4, row["b31_60"]).number_format = FMT_EUR
        ws.cell(r, 5, row["b60p"]).number_format = FMT_EUR
        ws.cell(r, 6, row["total"]).number_format = FMT_EUR
        ws.cell(r, 7, row["overdue"]).number_format = FMT_EUR
        r += 1
    reports.band_rows(ws, 5, r - 1, len(hdr))
    if r > 5:
        ws.cell(r, 1, "TOTAL")
        for col in range(2, 8):
            L = reports.get_column_letter(col)
            ws.cell(r, col, f"=SUM({L}5:{L}{r-1})").number_format = FMT_EUR
        reports.totals_row(ws, r, len(hdr))
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    return _xlsx_bytes(wb)


# ============================================================ PDF EXPORT (HTML→wkhtmltopdf)
# PDF via the SAME Latvian-capable HTML→PDF path as the invoice PDF (invoicing._render_pdf_
# wkhtmltopdf with the dependency-free Unicode fallback). Each report builds an escaped
# HTML document (every DB value through markupsafe.escape) which is rendered to PDF bytes.
# The fixed labels are localized via i18n.t in the requested language (EN default).

_REPORT_CSS = """
@page { size: A4; margin: 18mm 14mm; }
body { font-family: 'DejaVu Sans', Arial, sans-serif; color:#1A2733; font-size:10pt; }
h1 { font-size:16pt; margin:0 0 2px; }
.sub { color:#5B6B7A; font-size:8.5pt; margin:0 0 14px; }
table { border-collapse:collapse; width:100%; margin:8px 0 16px; }
th, td { border:1px solid #DDE4EA; padding:5px 8px; text-align:left; }
th { background:#0E5FA8; color:#fff; font-weight:bold; }
td.num, th.num { text-align:right; }
tr.total td { font-weight:bold; border-top:2px solid #9FB3C4; background:#F2F6FA; }
h2 { font-size:12pt; margin:14px 0 4px; }
.kv { margin:2px 0; }
.kv b { display:inline-block; min-width:160px; }
.foot { color:#888; font-size:8pt; margin-top:18px; }
"""


def _fmt(x):
    """Display EUR (2 dp, thousands) for the PDF/HTML reports."""
    try:
        return f"{float(x):,.2f}"
    except (TypeError, ValueError):
        return "0.00"


def _render_report_pdf(html):
    """Render a report HTML document to PDF bytes via the invoice HTML→PDF path
    (wkhtmltopdf primary, dependency-free Unicode fallback). Returns bytes or None."""
    import invoicing
    try:
        data = invoicing._render_pdf_wkhtmltopdf(html)
        if data:
            return data
    except Exception as e:
        log.warning("report pdf wkhtmltopdf path failed: %s", e)
    # dependency-free Unicode fallback: render the visible text lines.
    try:
        return _fallback_pdf(html)
    except Exception as e:
        log.warning("report pdf fallback failed: %s", e)
        return None


def _fallback_pdf(html):
    """A minimal dependency-free PDF (Unicode-correct) when wkhtmltopdf is absent: strip the
    HTML to text lines and emit them via invoicing's Type0/Identity-H Unicode PDF writer."""
    import re
    import invoicing
    # crude HTML→text: drop the <style>…</style>, turn block tags into newlines, strip tags.
    txt = re.sub(r"(?is)<style.*?</style>", "", html)
    txt = re.sub(r"(?i)</(tr|h1|h2|div|p)>", "\n", txt)
    txt = re.sub(r"(?i)</t[dh]>", "  ", txt)
    txt = re.sub(r"(?s)<[^>]+>", "", txt)
    import html as _htmlmod
    txt = _htmlmod.unescape(txt)
    lines = [ln.strip() for ln in txt.splitlines() if ln.strip()]
    return invoicing.text_to_unicode_pdf("\n".join(lines))


def _doc_head(title_html, sub_html):
    return ("<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<style>{_REPORT_CSS}</style></head><body>"
            f"<h1>{title_html}</h1><div class='sub'>{sub_html}</div>")


def vat_output_html(rep, lang=None):
    """The VAT-output report as an escaped, print-ready HTML document (labels localized)."""
    from markupsafe import escape as esc
    import i18n
    L = i18n.normalize(lang) if lang is not None else i18n.current_lang()

    def _t(s):
        return i18n.t(s, L)

    P = [_doc_head(esc(_t("Output VAT report (PVN)")),
                   esc(_t("Period")) + ": " + esc(rep.get("label", "")) + " · "
                   + esc(rep.get("start", "")) + " … " + esc(rep.get("end", "")) + " · "
                   + esc(_t("EUR, NET (taxable) basis; tax point = issue date; credit "
                           "notes subtracted")))]
    P.append(f"<table><thead><tr><th>{esc(_t('VAT rate'))}</th>"
             f"<th class='num'>{esc(_t('Taxable net'))}</th>"
             f"<th class='num'>{esc(_t('Output VAT'))}</th></tr></thead><tbody>")
    for row in rep["rows"]:
        lbl = _vat_row_label(row, _t)
        P.append(f"<tr><td>{esc(lbl)}</td>"
                 f"<td class='num'>{esc(_fmt(row['net']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['vat']))}</td></tr>")
    P.append(f"<tr class='total'><td>{esc(_t('Total output VAT'))}</td>"
             f"<td class='num'>{esc(_fmt(rep['net_total']))}</td>"
             f"<td class='num'>{esc(_fmt(rep['vat_total']))}</td></tr>")
    P.append("</tbody></table>")
    _note = _t("Reverse-charge and 0%/exempt supplies are reportable but carry no "
               "output VAT.")
    P.append(f"<div class='foot'>{esc(_note)}</div>")
    P.append("</body></html>")
    return "".join(str(p) for p in P)


def _vat_row_label(row, _t):
    if row["kind"] == "reverse_charge":
        return _t("Reverse charge (AE)")
    if row["kind"] == "zero":
        return _t("0% / exempt")
    return row["label"]


def revenue_html(rep, lang=None):
    from markupsafe import escape as esc
    import i18n
    L = i18n.normalize(lang) if lang is not None else i18n.current_lang()

    def _t(s):
        return i18n.t(s, L)

    P = [_doc_head(esc(_t("Sales / revenue report")),
                   esc(_t("Period")) + ": " + esc(rep.get("label", "")) + " · "
                   + esc(rep.get("start", "")) + " … " + esc(rep.get("end", "")) + " · "
                   + esc(_t("NET (VAT-excluded) EUR; credit notes subtracted")))]
    # by month
    P.append(f"<h2>{esc(_t('Revenue by month'))}</h2>")
    P.append(f"<table><thead><tr><th>{esc(_t('Month'))}</th>"
             f"<th class='num'>{esc(_t('Net'))}</th>"
             f"<th class='num'>{esc(_t('Invoices'))}</th>"
             f"<th class='num'>{esc(_t('Credit notes'))}</th></tr></thead><tbody>")
    for m in rep["by_month"]:
        P.append(f"<tr><td>{esc(m['month'])}</td><td class='num'>{esc(_fmt(m['net']))}</td>"
                 f"<td class='num'>{m['invoices']}</td><td class='num'>{m['credits']}</td></tr>")
    P.append(f"<tr class='total'><td>{esc(_t('Total'))}</td>"
             f"<td class='num'>{esc(_fmt(rep['net_total']))}</td>"
             f"<td class='num'>{rep['invoice_count']}</td>"
             f"<td class='num'>{rep['credit_count']}</td></tr></tbody></table>")
    # by customer
    P.append(f"<h2>{esc(_t('Revenue by customer'))}</h2>")
    P.append(f"<table><thead><tr><th>{esc(_t('Customer'))}</th>"
             f"<th class='num'>{esc(_t('Net'))}</th>"
             f"<th class='num'>{esc(_t('Invoices'))}</th>"
             f"<th class='num'>{esc(_t('Credit notes'))}</th></tr></thead><tbody>")
    for c in rep["by_customer"]:
        P.append(f"<tr><td>{esc(c['customer'])}</td><td class='num'>{esc(_fmt(c['net']))}</td>"
                 f"<td class='num'>{c['invoices']}</td><td class='num'>{c['credits']}</td></tr>")
    P.append("</tbody></table>")
    # by service
    P.append(f"<h2>{esc(_t('Revenue by service'))}</h2>")
    P.append(f"<table><thead><tr><th>{esc(_t('Service / description'))}</th>"
             f"<th class='num'>{esc(_t('Net'))}</th>"
             f"<th class='num'>{esc(_t('Lines'))}</th></tr></thead><tbody>")
    for s in rep["by_service"]:
        P.append(f"<tr><td>{esc(s['service'])}</td><td class='num'>{esc(_fmt(s['net']))}</td>"
                 f"<td class='num'>{s['lines']}</td></tr>")
    P.append("</tbody></table></body></html>")
    return "".join(str(p) for p in P)


def customer_statement_html(rep, lang=None):
    from markupsafe import escape as esc
    import i18n
    L = i18n.normalize(lang) if lang is not None else i18n.current_lang()

    def _t(s):
        return i18n.t(s, L)

    cust = rep.get("customer") or {}
    P = [_doc_head(esc(_t("Statement of account")),
                   esc(cust.get("name") or "—") + " · " + esc(rep.get("start", ""))
                   + " … " + esc(rep.get("end", "")) + " · "
                   + esc(_t("GROSS EUR; debit = invoice, credit = credit note / payment")))]
    if cust.get("address"):
        P.append(f"<div class='kv'>{esc(cust.get('address'))}</div>")
    if cust.get("vat_number"):
        P.append(f"<div class='kv'><b>{esc(_t('VAT'))}:</b> {esc(cust.get('vat_number'))}</div>")
    P.append(f"<div class='kv'><b>{esc(_t('Opening balance'))}:</b> {esc(_fmt(rep['opening_balance']))} EUR</div>")
    P.append(f"<table><thead><tr><th>{esc(_t('Date'))}</th><th>{esc(_t('Type'))}</th>"
             f"<th>{esc(_t('Reference'))}</th>"
             f"<th class='num'>{esc(_t('Debit'))}</th>"
             f"<th class='num'>{esc(_t('Credit'))}</th>"
             f"<th class='num'>{esc(_t('Balance'))}</th></tr></thead><tbody>")
    for ln in rep["lines"]:
        P.append(f"<tr><td>{esc(ln['date'])}</td><td>{esc(_t(ln['type']))}</td>"
                 f"<td>{esc(ln['ref'])}</td>"
                 f"<td class='num'>{esc(_fmt(ln['debit']) if ln['debit'] else '')}</td>"
                 f"<td class='num'>{esc(_fmt(ln['credit']) if ln['credit'] else '')}</td>"
                 f"<td class='num'>{esc(_fmt(ln['balance']))}</td></tr>")
    P.append(f"<tr class='total'><td colspan='5'>{esc(_t('Closing balance'))}</td>"
             f"<td class='num'>{esc(_fmt(rep['closing_balance']))}</td></tr>")
    P.append("</tbody></table>")
    P.append(f"<div class='kv'><b>{esc(_t('Total outstanding'))}:</b> "
             f"{esc(_fmt(rep['outstanding_total']))} EUR</div>")
    P.append("</body></html>")
    return "".join(str(p) for p in P)


def ar_aging_html(rep, lang=None):
    from markupsafe import escape as esc
    import i18n
    L = i18n.normalize(lang) if lang is not None else i18n.current_lang()

    def _t(s):
        return i18n.t(s, L)

    P = [_doc_head(esc(_t("Accounts receivable aging")),
                   esc(_t("As of")) + " " + esc(rep.get("as_of", "")) + " · "
                   + esc(_t("Outstanding EUR (gross − payments − credits)")))]
    P.append(f"<table><thead><tr><th>{esc(_t('Customer'))}</th>"
             f"<th class='num'>{esc(_t('Current (not due)'))}</th>"
             f"<th class='num'>{esc(_t('1–30 days'))}</th>"
             f"<th class='num'>{esc(_t('31–60 days'))}</th>"
             f"<th class='num'>{esc(_t('60+ days'))}</th>"
             f"<th class='num'>{esc(_t('Total outstanding'))}</th>"
             f"<th class='num'>{esc(_t('Overdue'))}</th></tr></thead><tbody>")
    for row in rep["rows"]:
        P.append(f"<tr><td>{esc(row['customer'])}</td>"
                 f"<td class='num'>{esc(_fmt(row['current']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['b1_30']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['b31_60']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['b60p']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['total']))}</td>"
                 f"<td class='num'>{esc(_fmt(row['overdue']))}</td></tr>")
    b = rep["buckets"]
    P.append(f"<tr class='total'><td>{esc(_t('Total'))}</td>"
             f"<td class='num'>{esc(_fmt(b['current']))}</td>"
             f"<td class='num'>{esc(_fmt(b['1-30']))}</td>"
             f"<td class='num'>{esc(_fmt(b['31-60']))}</td>"
             f"<td class='num'>{esc(_fmt(b['60+']))}</td>"
             f"<td class='num'>{esc(_fmt(rep['total_outstanding']))}</td>"
             f"<td class='num'>{esc(_fmt(rep['total_overdue']))}</td></tr>")
    P.append("</tbody></table></body></html>")
    return "".join(str(p) for p in P)


# PDF byte builders (HTML → wkhtmltopdf / fallback) -------------------------------------
def vat_output_pdf(rep, lang=None):
    return _render_report_pdf(vat_output_html(rep, lang=lang))


def revenue_pdf(rep, lang=None):
    return _render_report_pdf(revenue_html(rep, lang=lang))


def customer_statement_pdf(rep, lang=None):
    return _render_report_pdf(customer_statement_html(rep, lang=lang))


def ar_aging_pdf(rep, lang=None):
    return _render_report_pdf(ar_aging_html(rep, lang=lang))
