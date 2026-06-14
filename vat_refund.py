"""
VAT REFUND MODULE (EU 2008/9/EC)
Builds refund applications per ENTITY x REFUND COUNTRY x PERIOD (Q1-Q4, YEAR)
from fuel_history.db, applies minimum thresholds, tracks application status,
and exports an Excel claim workbook with invoice-level detail per stream.

Monthly/quarterly use:
    python3 vat_refund.py                       -> matrix + VAT_Refund_Claims_<year>.xlsx
    python3 vat_refund.py --set-status "Jupiter Plus AS" Sweden 2026-Q2 submitted
Statuses: draft -> ready -> submitted -> approved -> paid (free text allowed)
"""
import sqlite3, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import supplier_master, customer_master, audit, money
import db, db_tuning, db_migrate, applog
import vat_config
from vat_config import (GOODS_CODE,
                        MIN_QUARTER, MIN_ANNUAL, DEADLINE_FMT,
                        LOCAL_CCY_INPUT, COMPLIANCE_NOTES)

import os
WORKDIR = os.path.dirname(os.path.abspath(__file__))
# The VAT-refund CLAIM records (applications, one-invoice locks, the document vault
# index) live in their OWN database, isolated from the analytics/transactions store
# that history.py rebuilds every month — so a monthly reload can never corrupt the
# legal/financial claim data. Transactions are read from the analytics DB on demand.
DB = f"{WORKDIR}/vat_claims.db"            # claim records (this module owns it)
ANALYTICS_DB = f"{WORKDIR}/fuel_history.db"  # transactions (read-only here)

log = applog.get("vat_refund")

def quarter(period):           # '2026-05' -> '2026-Q2'
    y, m = period.split("-")
    return f"{y}-Q{(int(m)-1)//3+1}"

def q_months(per):             # '2026-Q2' -> Apr-Jun; '2026-YEAR' -> all 12 months
    if per.endswith("-YEAR"):
        y = per.split("-")[0]
        return [f"{y}-{m:02d}" for m in range(1, 13)]
    y, q = per.split("-Q")
    return [f"{y}-{m:02d}" for m in range((int(q)-1)*3+1, (int(q)-1)*3+4)]

def period_end_date(period):
    """Last calendar day of a claim period. '2026-Q2' -> 2026-06-30;
    '2026-YEAR' -> 2026-12-31. Returns a datetime.date."""
    import datetime, calendar
    p = str(period); y = int(p[:4])
    if p.endswith("-YEAR"):
        return datetime.date(y, 12, 31)
    q = int(p.split("-Q")[1]); m = q * 3
    return datetime.date(y, m, calendar.monthrange(y, m)[1])

def period_ended(period, today=None):
    """True once the claim period has fully closed (a period can't be filed early)."""
    import datetime
    return (today or datetime.date.today()) > period_end_date(period)

_SCHEMA_READY = set()   # DB files whose schema is set up this process

def connect():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    db_tuning.tune(con)  # WAL + busy_timeout for safe multi-process access
    audit.bind(con)   # audit triggers call ffs_actor(); register it every connect
    if DB != ":memory:" and DB in _SCHEMA_READY:
        return con
    con.execute("""CREATE TABLE IF NOT EXISTS vat_applications (
        entity TEXT, refund_country TEXT, ref_period TEXT,
        vat_eur REAL, vat_local REAL, currency TEXT,
        status TEXT DEFAULT 'draft', updated TEXT DEFAULT CURRENT_TIMESTAMP,
        submitted_date TEXT, approved_date TEXT, paid_date TEXT, paid_amount REAL,
        PRIMARY KEY (entity, refund_country, ref_period))""")
    con.execute("""CREATE TABLE IF NOT EXISTS vat_claimed_invoices (
        entity TEXT, refund_country TEXT, supplier TEXT, invoice_ref TEXT,
        ref_period TEXT, locked_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE (entity, refund_country, supplier, invoice_ref))""")
    con.execute("""CREATE TABLE IF NOT EXISTS invoice_documents (
        id INTEGER PRIMARY KEY,
        entity TEXT, supplier TEXT, invoice_ref TEXT,
        filename TEXT, stored_path TEXT, sha256 TEXT, size INTEGER,
        kind TEXT DEFAULT 'original_pdf', uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE (entity, supplier, invoice_ref, sha256))""")
    # Receipt-control WAIVER: an admin marks a supplier as "no invoice coming" for one
    # (entity, refund_country, period) so its UNMATCHED transactions (no registered
    # invoice exists for that country) do NOT block submission and are EXCLUDED from the
    # claim. Only a genuinely-uninvoiced supplier is waivable (see add_waiver / R5).
    con.execute("""CREATE TABLE IF NOT EXISTS vat_invoice_waivers (
        entity TEXT, refund_country TEXT, ref_period TEXT, supplier TEXT,
        reason TEXT, waived_by TEXT, waived_at TEXT DEFAULT CURRENT_TIMESTAMP,
        PRIMARY KEY (entity, refund_country, ref_period, supplier))""")
    audit.install_audit(con, ["vat_applications", "vat_claimed_invoices", "invoice_documents",
                              "vat_invoice_waivers"])
    # versioned migrations: each runs ONCE per database (see db_migrate). Append only.
    db_migrate.apply(con, "vat_refund", [
        "ALTER TABLE invoice_documents ADD COLUMN backend TEXT DEFAULT 'local'",
        "ALTER TABLE invoice_documents ADD COLUMN web_url TEXT",
        # our fee: rate FROZEN at submission, fee CHARGED when refund is paid
        "ALTER TABLE vat_applications ADD COLUMN fee_eur REAL",
        "ALTER TABLE vat_applications ADD COLUMN fee_pct REAL",
        "ALTER TABLE vat_applications ADD COLUMN fee_min REAL",
        "ALTER TABLE vat_applications ADD COLUMN fee_billed_date TEXT",
        # settlement: where the refund landed + the fee invoice once issued
        "ALTER TABLE vat_applications ADD COLUMN payout_to TEXT",
        "ALTER TABLE vat_applications ADD COLUMN fee_invoice_no TEXT",
        "ALTER TABLE vat_applications ADD COLUMN fee_invoice_date TEXT",
        # workflow status CODE (1A..5); the legacy `status` column stays the
        # coarse engine state (draft/submitted/approved/paid) that drives locks.
        "ALTER TABLE vat_applications ADD COLUMN status_code TEXT",
        # per-status data: when the decision arrived, why (note), and the
        # deadline of an open action (document request 2B / appeal 3D)
        "ALTER TABLE vat_applications ADD COLUMN decision_date TEXT",
        "ALTER TABLE vat_applications ADD COLUMN status_note TEXT",
        "ALTER TABLE vat_applications ADD COLUMN action_deadline TEXT",
    ])
    _migrate_from_analytics(con)
    if DB != ":memory:":
        _SCHEMA_READY.add(DB)
    return con

def analytics_connect():
    """Read-only connection to the analytics DB (fuel_history) for reading
    `transactions`. fuel_history.db is OWNED/written by the data-processing engine
    (history.py); claim records are NOT here — they live in DB (vat_claims.db).
    Delegates to the dataproduct accessor so all app product READS share one
    read-only window; the name/signature are unchanged for claim-side callers.
    Passes this module's ANALYTICS_DB so the location-independent / test-monkeypatch
    seam is preserved (the file may differ; the read-only window is the same)."""
    import dataproduct
    return dataproduct.connect("fuel_history", path=ANALYTICS_DB)

def _migrate_from_analytics(con):
    """One-time upgrade path: if the claim tables are empty in the (new) claims DB but
    populated in the old shared fuel_history.db, copy them across. Idempotent — skips
    once the claims DB has data. The originals are left in place (now unused)."""
    if DB == ":memory:" or DB in _SCHEMA_READY:
        return
    try:
        if con.execute("SELECT COUNT(*) FROM vat_applications").fetchone()[0] > 0:
            return
        if con.execute("SELECT COUNT(*) FROM invoice_documents").fetchone()[0] > 0:
            return
    except Exception:
        return
    if not os.path.exists(ANALYTICS_DB) or os.path.abspath(ANALYTICS_DB) == os.path.abspath(DB):
        return
    src = sqlite3.connect(ANALYTICS_DB)
    try:
        for t in ("vat_applications", "vat_claimed_invoices", "invoice_documents"):
            try:
                cols = [r[1] for r in src.execute(f"PRAGMA table_info({t})")]
            except db.DBError:
                cols = []
            if not cols:
                continue
            rows = src.execute(f"SELECT {','.join(cols)} FROM {t}").fetchall()
            if rows:
                ph = ",".join("?" * len(cols))
                con.executemany(f"INSERT OR IGNORE INTO {t} ({','.join(cols)}) VALUES ({ph})", rows)
        con.commit()
    except Exception:
        # degrade gracefully (the claims DB simply starts empty), but never
        # silently: an unnoticed failure here looks identical to "no legacy data"
        log.exception("legacy claim-DB import failed — vat_claims.db may be "
                      "missing migrated rows (source: %s)", ANALYTICS_DB)
    finally:
        src.close()

DOCDIR = f"{WORKDIR}/documents"

def attach_document(con, ent, sup, ref, src_path=None, file_bytes=None,
                    filename=None, kind="original_pdf", country=None, period=None):
    """Attach a physical document (original PDF or scan) to an invoice.
    Returns (ok, message). Hash-verified; duplicate file on the SAME invoice is
    skipped; the SAME file on a DIFFERENT invoice raises a warning (likely a
    wrong attachment) but is allowed with the warning recorded in the message.
    The file is archived under the logical vault tree
    <Customer> <RegNo>/<Year>/<Country>/<Claim period>/<file>; country/period are
    looked up from the invoice registry when not passed in."""
    import hashlib, os
    import document_vault
    if src_path:
        file_bytes = open(src_path, "rb").read()
        filename = filename or os.path.basename(src_path)
    sha = hashlib.sha256(file_bytes).hexdigest()
    dup_same = con.execute("""SELECT id FROM invoice_documents WHERE entity=? AND supplier=?
                              AND invoice_ref=? AND sha256=?""", (ent, sup, ref, sha)).fetchone()
    if dup_same:
        return True, f"already attached (identical file, sha {sha[:8]}) - skipped"
    elsewhere = con.execute("""SELECT invoice_ref FROM invoice_documents WHERE sha256=?
                               AND NOT (entity=? AND supplier=? AND invoice_ref=?)""",
                            (sha, ent, sup, ref)).fetchone()
    warn = (f" | WARNING: identical file already attached to invoice "
            f"{elsewhere['invoice_ref']} - verify correct document" if elsewhere else "")
    # resolve the metadata that organises the archive (country/period of the
    # invoice, and the customer's registration number) — all best-effort.
    if country is None or period is None:
        try:
            scon = supplier_master.connect()
            inv = scon.execute("""SELECT country, period FROM supplier_invoices
                                  WHERE supplier=? AND invoice_no=?""", (sup, ref)).fetchone()
            scon.close()
            if inv:
                country = country if country is not None else inv["country"]
                period = period if period is not None else inv["period"]
        except Exception as e:
            log.debug("vat_refund: supplier-invoice enrichment failed for %s/%s: %s", sup, ref, e)
    cust_name, reg = ent, None
    try:
        c = customer_master.get_customer(ent) or {}
        cust_name = c.get("company_name") or ent
        reg = c.get("reg_number")
    except Exception as e:
        log.debug("vat_refund: customer lookup failed for %s: %s", ent, e)
    safe = document_vault.invoice_vault_path(cust_name, reg, country, period, filename)
    be = document_vault.backend(DOCDIR)
    stored, web_url = be.put(safe, file_bytes)
    con.execute("""INSERT INTO invoice_documents (entity, supplier, invoice_ref, filename,
                   stored_path, sha256, size, kind, backend, web_url)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (ent, sup, ref, filename, stored, sha, len(file_bytes), kind, be.name, web_url))
    con.commit()
    return True, (f"attached {filename} ({len(file_bytes):,} B, sha {sha[:8]}, {kind}, "
                  f"storage: {be.name}" + (f", {web_url}" if web_url else "")) + warn

def attach_existing(con, ent, sup, ref, *, source, source_id, kind="scan"):
    """Attach a file that is ALREADY stored (data lake or this customer's vault tree)
    to invoice (ent, sup, ref). Reads the bytes from the chosen source and routes them
    through attach_document(), so SHA dedup and the cross-invoice WARNING are identical
    to a fresh upload — legal integrity is the same. Returns (ok, message)."""
    import data_lake, document_vault
    if source == "lake":
        meta, data = data_lake.get_file(source_id)
        if data is None:
            return False, "source file not found"
        return attach_document(con, ent, sup, ref, file_bytes=data,
                               filename=meta["filename"], kind=kind)
    if source == "vault":
        # source_id is a stored_path of an invoice_documents row for THIS customer.
        row = con.execute("""SELECT filename, stored_path FROM invoice_documents
                             WHERE entity=? AND stored_path=?""",
                          (ent, source_id)).fetchone()
        if not row:
            return False, "source file not found"
        try:
            data = document_vault.get_bytes(row["stored_path"], DOCDIR)
        except Exception as e:
            log.exception("attach_existing: vault read failed for %s", source_id)
            return False, f"source file not found ({str(e)[:80]})"
        return attach_document(con, ent, sup, ref, file_bytes=data,
                               filename=row["filename"], kind=kind)
    return False, f"unknown source '{source}'"

def docs_for(con, ent, sup, ref):
    return con.execute("""SELECT * FROM invoice_documents WHERE entity=? AND supplier=?
                          AND invoice_ref=? ORDER BY uploaded_at""", (ent, sup, ref)).fetchall()

def file_documents_for_claim(con, ent, ctry, period):
    """Re-file the documents of the invoices ACTUALLY locked into this claim
    (vat_claimed_invoices for this exact ref_period) under the claim's period
    folder, so the vault mirrors the real claim composition — dynamically.

    This is what makes merging dynamic: claims are filed per invoice, not per
    calendar quarter. If Q1/Q2 are low and get pulled into the yearly claim while
    Q3 and Q4 are claimed quarterly, only the Q1/Q2 (and any late) invoices in the
    YEAR claim move to 'Annual'; the Q3/Q4 documents stay under their own quarter
    because they are locked to '<year>-Q3'/'<year>-Q4'. A quarterly claim is
    typically a no-op (its documents are already in the matching quarter folder).

    Re-filed per document in a DB-safe order — write the new copy, point the row at
    it, THEN delete the old copy — so a crash never leaves the database referencing
    a missing file. Idempotent. Returns the number of documents moved."""
    import document_vault
    try:
        c = customer_master.get_customer(ent) or {}
        cust_name = c.get("company_name") or ent
        reg = c.get("reg_number")
    except Exception:
        cust_name, reg = ent, None
    locked = con.execute("""SELECT supplier, invoice_ref FROM vat_claimed_invoices
                            WHERE entity=? AND refund_country=? AND ref_period=?""",
                         (ent, ctry, period)).fetchall()
    moved = 0
    for lk in locked:
        for d in docs_for(con, ent, lk["supplier"], lk["invoice_ref"]):
            old = d["stored_path"]
            # filing folder follows the CLAIM's period (Annual for a yearly claim),
            # not the invoice's calendar quarter.
            new_name = document_vault.invoice_vault_path(cust_name, reg, ctry, period, d["filename"])
            data = document_vault.get_bytes(old, DOCDIR)
            new_loc, web_url = document_vault.copy_to(new_name, data, DOCDIR)
            if str(new_loc) == str(old):
                continue                                # already in the right folder
            con.execute("UPDATE invoice_documents SET stored_path=?, web_url=? WHERE id=?",
                        (new_loc, web_url, d["id"]))
            con.commit()                                # row now points at the new copy
            try:
                document_vault.delete(old, DOCDIR)         # safe to drop the old copy
            except Exception as e:
                # orphan at worst, never lost — but record the orphan so it's visible
                log.warning("file_documents_for_claim: could not delete old vault copy %s "
                            "(orphaned, not lost): %s", old, e)
            moved += 1
    return moved

def _verify_one(stored_path, recorded_sha):
    """Re-read ONE stored document and compare its live SHA-256 to the hash recorded
    when it was attached. Returns (status, detail, data) where status is OK /
    CORRUPT / MISSING and `data` is the raw bytes (None on MISSING). This is the
    canonical per-document integrity check reused by both verify_documents() (the
    whole-store sweep) and evidence_pack() (per-claim export) so they apply IDENTICAL
    hashing/verification logic — never re-implement hashing in a caller."""
    import hashlib
    import document_vault
    try:
        data = document_vault.get_bytes(stored_path, DOCDIR)
    except Exception as e:
        return "MISSING", str(e)[:140], None
    actual = hashlib.sha256(data).hexdigest()
    if actual != recorded_sha:
        return "CORRUPT", f"hash {actual[:8]} != recorded {str(recorded_sha)[:8]}", data
    return "OK", "", data

def verify_documents(con=None):
    """Integrity check for the physical documents (PDF/ZIP files): re-read each
    stored file and compare its SHA-256 to the hash recorded when it was attached.
    Detects corrupted or missing/jeopardised files. Returns (rows, summary)."""
    close = False
    if con is None:
        con = connect(); close = True
    rows, ok, corrupt, missing = [], 0, 0, 0
    for r in con.execute("""SELECT entity, supplier, invoice_ref, filename, stored_path,
                            sha256, size, backend FROM invoice_documents ORDER BY id"""):
        status, detail, _data = _verify_one(r["stored_path"], r["sha256"])
        if status == "OK":
            ok += 1
        elif status == "CORRUPT":
            corrupt += 1
        else:
            missing += 1
        rows.append({"entity": r["entity"], "supplier": r["supplier"],
                     "invoice_ref": r["invoice_ref"], "filename": r["filename"],
                     "sha256": r["sha256"], "backend": r["backend"],
                     "status": status, "detail": detail})
    if close:
        con.close()
    return rows, {"total": len(rows), "ok": ok, "corrupt": corrupt, "missing": missing}

def evidence_pack(entity, refund_country, period, con=None):
    """Audit-ready EVIDENCE-EXPORT pack (monetization M6).

    Bundle a VAT claim's SHA-256-verified original documents into a single ZIP so a
    claim's supporting evidence can be handed over provably-intact. The pack contains:
      * the ORIGINAL PDF/ZIP files (named by invoice_ref) read via
        document_vault.get_bytes(...) — the same vault accessor the rest of the module
        uses, so any storage backend (local/SharePoint/FTPS) works unchanged;
      * a MANIFEST.sha256.csv listing invoice_ref, supplier, filename, recorded sha256
        and the live verify STATUS (OK / MISMATCH / MISSING) — mirroring backup.py's
        SHA-256 MANIFEST so the recipient can re-prove integrity offline;
      * a COVER.txt summary (entity / country / period, claimed VAT, # invoices,
        # docs, and any MISSING/MISMATCH flagged loudly).

    Integrity is re-verified at export time via _verify_one() (the SAME logic as
    verify_documents()) — a corrupted or missing file is SURFACED in the cover +
    manifest, never silently dropped. Returns (zip_bytes, summary) where summary is a
    dict with the counts and the list of integrity failures.
    """
    import io, csv, zipfile, datetime
    close = False
    if con is None:
        con = connect(); close = True
    try:
        app = con.execute("""SELECT vat_eur, vat_local, currency, status, status_code
                             FROM vat_applications
                             WHERE entity=? AND refund_country=? AND ref_period=?""",
                          (entity, refund_country, period)).fetchone()
        # the invoices LOCKED into this exact claim (period-stamped registration)
        invoices = con.execute("""SELECT supplier, invoice_ref FROM vat_claimed_invoices
                                  WHERE entity=? AND refund_country=? AND ref_period=?
                                  ORDER BY supplier, invoice_ref""",
                               (entity, refund_country, period)).fetchall()
        manifest_rows = []          # (invoice_ref, supplier, filename, sha256, status, detail)
        failures = []               # human-readable integrity failures
        n_docs = ok = mismatch = missing = 0
        used_names = set()

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for inv in invoices:
                sup, ref = inv["supplier"], inv["invoice_ref"]
                docs = docs_for(con, entity, sup, ref)
                if not docs:
                    manifest_rows.append((ref, sup, "", "", "MISSING",
                                          "no document attached to this invoice"))
                    missing += 1
                    failures.append(f"{ref} ({sup}): MISSING — no document attached")
                    continue
                for d in docs:
                    n_docs += 1
                    # MISMATCH in the manifest == verify_documents' CORRUPT status:
                    # the live bytes no longer hash to the recorded sha256.
                    status, detail, data = _verify_one(d["stored_path"], d["sha256"])
                    arc_status = "MISMATCH" if status == "CORRUPT" else status
                    if status == "OK":
                        ok += 1
                    elif status == "CORRUPT":
                        mismatch += 1
                        failures.append(f"{ref} ({sup}): MISMATCH — {detail}")
                    else:
                        missing += 1
                        failures.append(f"{ref} ({sup}): MISSING — {detail}")
                    # safe, unique archive name: <invoice_ref>__<filename>
                    base = f"{_safe_name(ref)}__{_safe_name(d['filename'] or 'document')}"
                    name = base; i = 2
                    while name in used_names:
                        name = f"{base}.{i}"; i += 1
                    used_names.add(name)
                    manifest_rows.append((ref, sup, name, d["sha256"] or "",
                                          arc_status, detail))
                    # store the original bytes (even a MISMATCH: the recipient sees
                    # exactly what is on disk, with the manifest flagging it).
                    if data is not None:
                        z.writestr(f"documents/{name}", data)

            # ---- MANIFEST (mirrors backup.py's SHA-256 manifest, CSV form) ----
            mbuf = io.StringIO()
            w = csv.writer(mbuf)
            w.writerow(["invoice_ref", "supplier", "filename", "sha256",
                        "verify_status", "detail"])
            for row in manifest_rows:
                w.writerow(row)
            z.writestr("MANIFEST.sha256.csv", mbuf.getvalue())

            # ---- COVER summary ----
            vat_eur = (app["vat_eur"] if app else None)
            cover = [
                "VAT REFUND — EVIDENCE PACK",
                "=" * 60,
                f"Entity         : {entity}",
                f"Refund country : {refund_country}",
                f"Claim period   : {period}",
                f"Claimed VAT    : {money.f2(vat_eur):,.2f} EUR" if vat_eur is not None
                    else "Claimed VAT    : (no application record)",
                f"Status code    : {(app['status_code'] if app and app['status_code'] else (app['status'] if app else 'n/a'))}",
                f"Invoices       : {len(invoices)}",
                f"Documents      : {n_docs}",
                f"Generated (UTC): {datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')}",
                "",
                "INTEGRITY",
                "-" * 60,
                f"Verified OK    : {ok}",
                f"MISMATCH       : {mismatch}",
                f"MISSING        : {missing}",
            ]
            if failures:
                cover.append("")
                cover.append("!! INTEGRITY FAILURES — evidence NOT fully intact:")
                for f in failures:
                    cover.append(f"   - {f}")
            else:
                cover.append("")
                cover.append("All documents verified intact against their recorded SHA-256.")
            cover.append("")
            cover.append("Each document's live SHA-256 was re-verified against the hash recorded")
            cover.append("when it was attached. Re-hash any file and compare to MANIFEST.sha256.csv")
            cover.append("to independently confirm the evidence is unaltered.")
            z.writestr("COVER.txt", "\n".join(cover) + "\n")

        summary = {"entity": entity, "refund_country": refund_country, "period": period,
                   "vat_eur": money.f2(app["vat_eur"]) if app and app["vat_eur"] is not None else 0.0,
                   "invoices": len(invoices), "documents": n_docs,
                   "ok": ok, "mismatch": mismatch, "missing": missing,
                   "failures": failures, "intact": not failures}
        return buf.getvalue(), summary
    finally:
        if close:
            con.close()

def _safe_name(s):
    """Filesystem/zip-safe component of an invoice ref or filename."""
    import re
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(s)).strip("_") or "x"

LOCKING = ("submitted", "approved", "paid")

def _synthetic(ref, vat_id=None):
    """True if a claim line is not tied to ONE registered invoice: an INPUT
    placeholder (no registered invoice / no real VAT ID), an ALL: aggregate, or
    an UNMATCHED transaction (no registered invoice resolves it). Centralizes the
    predicate used by the lock gate, the readiness/checklist gates and the
    workbook so they all block the same set of synthetic refs."""
    ref = str(ref)
    return (("INPUT" in ref) or ref.startswith("ALL:") or (ref == "UNMATCHED")
            or ("INPUT" in str(vat_id)))

def stream_invoices(con, ent, ctry, period, cache=None):
    """Distinct (supplier, invoice_ref) used by a claim stream."""
    return sorted({(L["supplier"], L["invoice"]) for L in invoice_lines(con, ent, ctry, period, cache)})

def docs_index(con):
    """One-query set of (entity, supplier, invoice_ref) that have >=1 document.
    Lets callers check document coverage without an N+1 of docs_for()."""
    return {(r["entity"], r["supplier"], r["invoice_ref"])
            for r in con.execute(
                "SELECT DISTINCT entity, supplier, invoice_ref FROM invoice_documents")}

def _no_registered_invoices(sup, ctry, scon=None):
    """True iff supplier `sup` has NO registered invoice for refund country `ctry`
    (R5 case (a): a genuinely-uninvoiced supplier — its transactions are UNMATCHED
    because the invoice simply isn't coming). supplier_master.get_invoices never
    returns an empty list: with zero rows it yields a single 'INPUT: … invoice'
    placeholder, so 'no real invoices' == every returned ref is an INPUT: stub."""
    regs = supplier_master.get_invoices(sup, ctry, con=scon)
    return all(str(no).startswith("INPUT:") for no, _date in regs)

def _waivable_missing(ref, sup, ctry, scon=None):
    """True iff (sup, ref) is the R5 case-(a) genuinely-uninvoiced situation: a
    synthetic ref that is the INPUT-invoice stub AND the supplier has NO registered
    invoice for `ctry`. This is the ONLY waivable case. It excludes by construction:
      - a real (matched) ref            -> not _synthetic;
      - case-(b) UNMATCHED (>=2 regs)   -> _no_registered_invoices False;
      - an ALL:/vat-id placeholder      -> not the INPUT-invoice stub (startswith guard).
    The lock gate, the checklist gate and the /vat waive UI all key on THIS predicate
    so they never drift apart (mirror of _synthetic centralizing the block set)."""
    ref = str(ref)
    return (_synthetic(ref) and ref.startswith("INPUT")
            and _no_registered_invoices(sup, ctry, scon))

# --------------------------------------------------------------- receipt-control waivers
def list_waivers(con, ent, ctry, period):
    """The set of suppliers WAIVED ('invoice not coming') for one claim stream."""
    return {r["supplier"] for r in con.execute(
        """SELECT supplier FROM vat_invoice_waivers
           WHERE entity=? AND refund_country=? AND ref_period=?""",
        (ent, ctry, period))}

def add_waiver(con, ent, ctry, period, supplier, reason=None):
    """Waive a supplier's missing-invoice receipt-control item for one claim stream.
    Returns (ok, msg). REFUSES (R5) a supplier that HAS registered invoices for `ctry`:
    those transactions are UNMATCHED because of an ambiguous note-match (a matching fix),
    NOT a genuinely-uninvoiced supplier — waiving them would drop claimable VAT. Upsert
    so it is idempotent; the waive actor is stamped and the row is audited by its trigger."""
    if not _no_registered_invoices(supplier, ctry):
        return False, (f"can't waive '{supplier}' — it HAS registered invoice(s) for "
                       f"{ctry}; an UNMATCHED transaction there is a note-matching fix, "
                       "not a missing invoice (register/match the invoice instead)")
    actor = audit._current_actor() or "admin"
    con.execute("""INSERT INTO vat_invoice_waivers
                   (entity, refund_country, ref_period, supplier, reason, waived_by)
                   VALUES (?,?,?,?,?,?)
                   ON CONFLICT(entity, refund_country, ref_period, supplier)
                   DO UPDATE SET reason=excluded.reason, waived_by=excluded.waived_by,
                                 waived_at=CURRENT_TIMESTAMP""",
                (ent, ctry, period, supplier, (reason or None), actor))
    con.commit()
    return True, f"waived '{supplier}' — invoice not coming ({ctry} {period})"

def remove_waiver(con, ent, ctry, period, supplier):
    """Undo a receipt-control waiver. Returns (ok, msg)."""
    cur = con.execute("""DELETE FROM vat_invoice_waivers WHERE entity=? AND refund_country=?
                         AND ref_period=? AND supplier=?""", (ent, ctry, period, supplier))
    con.commit()
    if cur.rowcount:
        return True, f"removed waiver for '{supplier}'"
    return False, f"no waiver for '{supplier}'"

def lock_state(con, ent, ctry, sup, ref):
    r = con.execute("""SELECT ref_period FROM vat_claimed_invoices WHERE entity=? AND
                       refund_country=? AND supplier=? AND invoice_ref=?""",
                    (ent, ctry, sup, ref)).fetchone()
    return r["ref_period"] if r else None

def set_status(con, ent, ctry, period, new, gate_activation=True):
    """Guarded status transition enforcing one-invoice-one-submission.
    Returns (ok, message).

    `gate_activation` (default True) enforces the customer/country activation flags
    when entering a locked state. The workflow layer (`set_status_code`) sets it False
    because it has already enforced the adjustable system checklist, which supersedes
    the activation flags.

    The whole transition (duplicate checks + invoice-lock acquisition + the
    application upsert) runs as ONE transaction. Lock acquisition uses a plain
    INSERT (not INSERT OR IGNORE) so a lost race surfaces as an IntegrityError on
    the UNIQUE(entity, refund_country, supplier, invoice_ref) constraint; on that
    we roll back and abort the status change rather than silently proceeding as if
    we had won the lock."""
    # A tracked customer must be ACTIVATED (onboarding documents complete) before a
    # claim can be submitted on their behalf. Untracked entities are not gated.
    if gate_activation and new in LOCKING and customer_master.is_active(ent) is False:
        return False, (f"customer '{ent}' is not activated — complete the trade registry, "
                       f"bank account and signed contract on the Customers page first")
    # Each refund country is activated separately (request + receive its documents).
    # Once activation has been started for a country it must reach 'active' to submit.
    if gate_activation and new in LOCKING and customer_master.country_active(ent, ctry) is False:
        return False, (f"refund country '{ctry}' is not activated for '{ent}' — request and "
                       f"receive the country documents (power of attorney) on the Customers page")
    try:
        # Open an immediate transaction so concurrent claimants serialize on write.
        con.execute("BEGIN IMMEDIATE")
    except db.DBError:
        # Already inside a transaction (e.g. autocommit off / nested caller) - fine.
        pass
    try:
        cur = con.execute("""SELECT status FROM vat_applications WHERE entity=? AND
                             refund_country=? AND ref_period=?""", (ent, ctry, period)).fetchone()
        cur = cur["status"] if cur else "draft"
        is_annual = str(period).endswith("-YEAR")
        claim_set = None
        if new in LOCKING:
            if cur not in LOCKING:  # entering locked state -> validate & lock invoices
                invs = stream_invoices(con, ent, ctry, period)
                # Receipt-control WAIVERS: drop ONLY a waived, genuinely-uninvoiced
                # supplier's synthetic ref (R5 case (a): NO registered invoice exists for
                # this country, so its ref is the INPUT stub). A real (matched) ref
                # survives (_synthetic False); a supplier that HAS invoices is never
                # dropped (_no_registered_invoices False) so its UNMATCHED ref stays in
                # `bad`. Everything downstream (the `bad` gate, claim_set, locks, nodoc,
                # the frozen vat_eur/vat_local over claim_set keys) excludes the waived
                # transactions by construction.
                waived = list_waivers(con, ent, ctry, period)
                if waived:
                    scon = supplier_master.connect()
                    try:
                        invs = [(s, r) for s, r in invs
                                if not (s in waived and _waivable_missing(r, s, ctry, scon))]
                    finally:
                        scon.close()
                bad = [f"{s}:{r}" for s, r in invs if _synthetic(r)]
                if bad:
                    con.rollback()
                    return False, "BLOCKED - unresolved invoice refs (fill INPUTs first): " + "; ".join(bad)
                # A YEARLY claim is the mop-up for periods NOT already claimed
                # quarterly: invoices already locked to a quarter are EXCLUDED from
                # the annual claim (not a conflict). A quarterly claim still treats
                # any overlap as a duplicate and blocks.
                conflicts, claim_set = [], []
                for s, r in invs:
                    other = lock_state(con, ent, ctry, s, r)
                    if other and other != period:
                        if is_annual:
                            continue                    # claimed in its own quarter
                        conflicts.append(f"{s} invoice {r} already claimed in {other}")
                    else:
                        claim_set.append((s, r))
                if conflicts:
                    con.rollback()
                    return False, "BLOCKED - duplicate submission: " + "; ".join(conflicts)
                if is_annual and not claim_set:
                    con.rollback()
                    return False, ("BLOCKED - nothing to claim annually: every invoice for this "
                                   "year is already claimed in a quarterly filing")
                nodoc = [f"{s} {r}" for s, r in claim_set if not docs_for(con, ent, s, r)]
                if nodoc:
                    con.rollback()
                    return False, ("BLOCKED - physical document missing (attach original PDF "
                                   "or scan first): " + "; ".join(nodoc))
                for s, r in claim_set:
                    try:
                        con.execute("""INSERT INTO vat_claimed_invoices
                                       (entity, refund_country, supplier, invoice_ref, ref_period)
                                       VALUES (?,?,?,?,?)""", (ent, ctry, s, r, period))
                    except db.IntegrityError:
                        # Another claim acquired this invoice lock between our check
                        # and our insert. Abort the entire transition.
                        con.rollback()
                        other = lock_state(con, ent, ctry, s, r)
                        return False, ("BLOCKED - duplicate submission (concurrent claim won the "
                                       f"lock): {s} invoice {r} already claimed"
                                       + (f" in {other}" if other else "") + " - retry not needed")
        elif new == "withdrawn":
            # The ONLY path that frees invoices (used by withdraw_claim). Rejection /
            # appeal / confiscation deliberately KEEP the locks so the invoices can't
            # be re-claimed elsewhere while the claim is contested.
            con.execute("""DELETE FROM vat_claimed_invoices WHERE entity=? AND refund_country=?
                           AND ref_period=?""", (ent, ctry, period))
        elif new == "rejected":
            # A rejected claim KEEPS its invoice locks (mirrors 3B's 'approved' engine
            # state, whose lock-acquisition branch above is a no-op once cur is already
            # locked). Locks are released ONLY by withdraw_claim — a rejection must not
            # free invoices for re-claiming and risk a duplicate submission.
            pass
        elif cur in LOCKING:
            con.rollback()
            return False, (f"BLOCKED - application is '{cur}' and holds invoice locks; "
                           "use 'withdrawn' to release the locks before reverting.")
        stamp = {"submitted": "submitted_date", "approved": "approved_date", "paid": "paid_date"}.get(new)
        con.execute("""INSERT INTO vat_applications (entity, refund_country, ref_period, status)
                       VALUES (?,?,?,?) ON CONFLICT(entity, refund_country, ref_period)
                       DO UPDATE SET status=excluded.status, updated=CURRENT_TIMESTAMP""",
                    (ent, ctry, period, new))
        if stamp:
            con.execute(f"UPDATE vat_applications SET {stamp}=CURRENT_DATE WHERE entity=? "
                        "AND refund_country=? AND ref_period=?", (ent, ctry, period))
        # Freeze the fee RATE onto the claim the moment it is first submitted; once
        # locked the rate can no longer be adjusted (% / minimum changes only affect
        # un-submitted declarations).
        if new in LOCKING and cur not in LOCKING:
            # Freeze the VAT base from EXACTLY the invoices locked into THIS claim
            # (`claim_set`), via invoice_lines — the same one-row-per-(invoice,code)
            # basis used everywhere else. This is the canonical base for BOTH a yearly
            # claim (the deferred/late mop-up, never the whole calendar year) AND a
            # quarterly claim: a raw SUM(vat_eur) over all period transactions would
            # wrongly include period invoices NOT in this claim (e.g. already locked to
            # another claim), so the frozen vat_eur/fee would sit on a larger base than
            # the invoices actually filed. `claim_set` is always populated here (it is
            # built in the same `new in LOCKING and cur not in LOCKING` branch above).
            keys = set(claim_set or ())
            # One pass over the locked claim_set lines: accumulate BOTH the EUR base
            # (vat_eur) and the national-currency base (vat_local) so the frozen row
            # carries them on the SAME claim_set basis. vat_local feeds the local-
            # currency minimum-threshold gate for SE/DK locking claims; vat_eur/fee
            # freezing is byte-identical to before (same builtin-sum-then-f2).
            sve = svl = 0.0
            for L in invoice_lines(con, ent, ctry, period):
                if (L["supplier"], L["invoice"]) in keys:
                    sve += L["vat_eur"]
                    svl += L["vat_local"]
            ve = money.f2(sve)
            vl = money.f2(svl)
            fpct, fmin = customer_master.fee_for(ent, ctry)
            fee, _basis = customer_master.compute_fee(ve, fpct, fmin)
            con.execute("""UPDATE vat_applications SET vat_eur=?, vat_local=?, fee_eur=?, fee_pct=?, fee_min=?
                           WHERE entity=? AND refund_country=? AND ref_period=?""",
                        (ve, vl, fee, fpct, fmin, ent, ctry, period))
        # CHARGE the fee for services only when the money is refunded (status=paid):
        # recompute on the refunded amount (paid_amount, else the claimed VAT) at the
        # frozen rate and stamp the billing date.
        if new == "paid":
            r = con.execute("""SELECT vat_eur, paid_amount, fee_pct, fee_min FROM vat_applications
                               WHERE entity=? AND refund_country=? AND ref_period=?""",
                            (ent, ctry, period)).fetchone()
            base = (r["paid_amount"] if r and r["paid_amount"] else (r["vat_eur"] if r else 0)) or 0
            fee, _b = customer_master.compute_fee(base, (r["fee_pct"] if r else 0) or 0,
                                              (r["fee_min"] if r else 0) or 0)
            con.execute("""UPDATE vat_applications SET fee_eur=?, fee_billed_date=CURRENT_DATE,
                           payout_to=COALESCE(payout_to, ?)
                           WHERE entity=? AND refund_country=? AND ref_period=?""",
                        (fee, customer_master.payout_route(ent), ent, ctry, period))
        con.commit()
    except Exception:
        con.rollback()
        raise
    # Keep the vault in step with the claim: file the documents of the invoices
    # actually locked into THIS claim under the claim's period folder. A quarterly
    # claim is a no-op (docs already in that quarter); a yearly claim pulls its
    # deferred/late invoices into 'Annual'. Done AFTER commit and outside the
    # transaction (file moves aren't transactional); best-effort so it never blocks
    # the already-recorded status change.
    if new in LOCKING and cur not in LOCKING:
        try:
            file_documents_for_claim(con, ent, ctry, period)
        except Exception as e:
            log.warning("vat_refund: file_documents_for_claim failed for %s/%s/%s "
                        "(status change kept): %s", ent, ctry, period, e)
    return True, f"status -> {new}" + (" (invoices locked)" if new in LOCKING and cur not in LOCKING
                                       else " (locks released)" if new == "withdrawn"
                                       else " (invoices stay locked)" if new == "rejected" else "")

# ===================================================================================
# WORKFLOW STATUS CODES  (1A..5)  — a controllable claim lifecycle on top of the engine
# ===================================================================================
# The pre-submission codes (1A/1B/1C/1E) are SYSTEM-CONTROLLED: derived live from the
# adjustable checklist + period end + threshold, never set by hand. The rest are
# advanced manually and map to the coarse engine `status` that drives locks/fees.
STATUS_LABELS = {
    "1A": "Missing documents",
    "1B": "Documents received — period not ended",
    "1C": "Can be submitted",
    "1E": "Ready to submit",
    "2":  "Submitted",
    "2A": "Successfully submitted",
    "2B": "Document request received",
    "3":  "Decision received",
    "3A": "Money received",
    "3B": "Rejection",
    "3D": "Under appeal",
    "3C": "Confiscation by government",
    "4":  "Ready to invoice fee",
    "4A": "Ready to invoice credit",
    "5":  "Closed",
}
AUTO_CODES = ("1A", "1B", "1C", "1E")          # system-derived; never set by a user
MANUAL_CODES = ("2", "2A", "2B", "3", "3A", "3B", "3D", "3C", "4", "4A", "5")
# workflow code -> coarse engine status (drives the lock/fee machinery). 3B/3C/3D map
# to LOCKING states so the invoice locks are KEPT (appeal / invoice the fee); only an
# explicit withdraw releases them.
ENGINE_OF = {
    "2": "submitted", "2A": "submitted", "2B": "submitted", "3D": "submitted",
    "3": "approved", "3B": "approved", "3C": "approved",
    "3A": "paid", "4": "paid", "4A": "paid", "5": "paid",
}

# Human-readable status of an OPEN power-of-attorney document request, used purely to
# annotate (never gate) the PoA checklist item. Keys are customer_master request statuses.
_POA_REQUEST_LABEL = {
    "requested":          "request created",
    "generated":          "draft generated",
    "sent_for_signature": "sent for signature",
    "signed":             "signed, awaiting receipt",
}

def _open_poa_request_note(cm, code, ctry):
    """The status note of the open power-of-attorney document request for (code, country),
    or None if there is none / on any read error. ANNOTATION ONLY — never gates. cm is a
    customers.db connection (a DIFFERENT DB than the claims `con`)."""
    try:
        best = None
        for r in customer_master.list_document_requests(cm, code):
            if (r.get("kind") == "power_of_attorney"
                    and (r.get("refund_country") or None) == (ctry or None)
                    and r.get("status") in _POA_REQUEST_LABEL):
                # Prefer the most-advanced (latest id) open request if several exist.
                if best is None or r["id"] > best["id"]:
                    best = r
        return _POA_REQUEST_LABEL[best["status"]] if best else None
    except Exception:
        applog.get("vat_refund").exception("open PoA request lookup failed for %s/%s",
                                           code, ctry)
        return None

def submission_checklist(con, ent, ctry, period, cache=None):
    """SYSTEM-CONTROLLED checklist for one claim stream: the adjustable customer/country
    requirements (customer_master.checklist_rules) PLUS the claim-level data checks.
    Returns [(label, ok)] — the user cannot tick these; the system verifies each."""
    cache = cache if cache is not None else {}
    cm = cache.get("_cmcon")
    if cm is None:
        cm = cache["_cmcon"] = customer_master.connect()
    codes = cache.setdefault("_codeof", {})
    if ent not in codes:
        codes[ent] = customer_master._code_of(cm, ent)
    code = codes[ent]
    items = []
    if code is not None:
        # If a power-of-attorney document REQUEST is open for this (entity, country),
        # enrich the PoA checklist item's LABEL with the request's status so the user can
        # see "sent for signature" etc. without leaving the claim. This is LABEL TEXT ONLY:
        # the boolean `ok` is untouched — the gate truth stays exactly evaluate_checklist /
        # country_active. A failure to read the request degrades to the plain label.
        poa_note = _open_poa_request_note(cm, code, ctry)
        for k, label, _scope, ok in customer_master.evaluate_checklist(cm, code, ctry):
            if poa_note and k == "power_of_attorney" and _scope == "country" and not ok:
                label = f"{label} — {poa_note}"
            items.append((label, ok))
    invs = stream_invoices(con, ent, ctry, period, cache)
    waived = list_waivers(con, ent, ctry, period)
    scon = cache.get("_scon")
    if scon is None:
        scon = cache["_scon"] = supplier_master.connect()
    # Receipt control names WHICH suppliers are blocking. A supplier is WAIVABLE-MISSING
    # only when its ref is SYNTHETIC *and* it has NO registered invoice for this country
    # (R5 case (a): the invoice isn't coming, so its ref is the INPUT stub). It drops out
    # of the missing list once an admin waives it. A case-(b) UNMATCHED supplier (HAS
    # invoices but no note match) is NOT waivable and falls into `other_bad` below so it
    # stays blocked here — never silenced.
    nri = cache.setdefault("_waivable", {})
    def waivable(s, r):
        if (s, r) not in nri:
            nri[(s, r)] = _waivable_missing(r, s, ctry, scon)
        return nri[(s, r)]
    missing_sup = sorted({s for s, r in invs if waivable(s, r) and s not in waived})
    items.append(("Receipt control: required invoices received"
                  + (" — missing: " + ", ".join(missing_sup) if missing_sup else ""),
                  len(invs) > 0 and not missing_sup))
    # Non-waivable synthetic refs: ALL: aggregates, vat-id INPUT placeholders AND
    # case-(b) UNMATCHED where the supplier HAS invoices (a note-matching fix, not a
    # missing invoice). Every synthetic ref is covered by exactly one of the two items:
    # a WAIVABLE (s, r) (no registered invoice) is named in `missing_sup` until waived;
    # a NON-waivable synthetic blocks here unconditionally. Nothing slips through: a
    # synthetic ref is in `missing_sup` (waivable, unwaived), excluded entirely
    # (waivable, waived), or in `other_bad` (not waivable).
    other_bad = sorted({r for s, r in invs if _synthetic(r) and not waivable(s, r)})
    items.append(("All invoice refs resolved (no INPUT/aggregate placeholders)",
                  len(invs) > 0 and not other_bad))
    docidx = cache.get("_docidx")
    if docidx is None:
        docidx = cache["_docidx"] = docs_index(con)
    # A WAIVED, genuinely-uninvoiced supplier is excluded from the claim (set_status drops
    # its synthetic ref), so it is NOT required to have a document — mirror that exclusion
    # here, else the doc gate would re-block a claim the receipt-control item just cleared.
    nodoc = [(s, r) for s, r in invs
             if (ent, s, r) not in docidx and not (s in waived and waivable(s, r))]
    items.append(("All invoice documents attached", len(invs) > 0 and not nodoc))
    items.append(("Claim period ended", period_ended(period)))
    return items

def derive_stage(con, ent, ctry, period, verdict=None, cache=None):
    """The SYSTEM-derived pre-submission stage. Returns (code, checklist):
      1A missing items · 1B checklist done but period open · 1C can submit (a caveat,
      e.g. below threshold) · 1E ready (all clear). Period-end is a hard gate."""
    items = submission_checklist(con, ent, ctry, period, cache)
    non_period = [(l, ok) for (l, ok) in items if l != "Claim period ended"]
    if not all(ok for _, ok in non_period):
        return "1A", items
    if not period_ended(period):
        return "1B", items
    caveat = bool(verdict and not str(verdict).startswith("READY"))
    return ("1C" if caveat else "1E"), items

def current_code(con, ent, ctry, period, verdict=None, cache=None):
    """The claim's effective workflow code: the stored manual code once one exists,
    otherwise the live system-derived pre-submission stage."""
    r = con.execute("""SELECT status, status_code FROM vat_applications
                       WHERE entity=? AND refund_country=? AND ref_period=?""",
                    (ent, ctry, period)).fetchone()
    if r and r["status_code"]:
        return r["status_code"]
    if r and r["status"] in ("submitted", "approved", "paid"):   # legacy rows, no code
        return {"submitted": "2", "approved": "3", "paid": "3A"}[r["status"]]
    code, _ = derive_stage(con, ent, ctry, period, verdict, cache)
    return code

def _stream_vat(con, ent, ctry, period):
    """The claim's applicable VAT as (vat_eur, vat_local, currency), Decimal-exact.

    Prefers the FROZEN figures on the submitted application row; falls back to the
    period aggregate the verdict path (claim_matrix) uses — SUM of the period's
    transactions for this (entity, country) — so the threshold gate and the displayed
    verdict agree. Returns Decimals (full precision); never raises (a missing/empty
    stream reads as 0). `con` is the claims connection.
    """
    row = con.execute("""SELECT vat_eur, vat_local, currency, status FROM vat_applications
                         WHERE entity=? AND refund_country=? AND ref_period=?""",
                      (ent, ctry, period)).fetchone()
    if row and (row["status"] or "draft") in LOCKING and row["vat_eur"] is not None:
        # frozen at submission over exactly the locked claim_set — the canonical base
        return (money.D(row["vat_eur"] or 0), money.D(row["vat_local"] or 0),
                (row["currency"] or "EUR"))
    months = q_months(period)
    acon = analytics_connect()
    try:
        agg = acon.execute(
            """SELECT ROUND(SUM(vat_eur),2) ve, ROUND(SUM(vat_local),2) vl,
                      MAX(currency) ccy
               FROM transactions WHERE entity=? AND country=? AND period IN (%s)"""
            % ",".join("?" * len(months)), [ent, ctry] + months).fetchone()
    finally:
        acon.close()
    ve = money.D(agg["ve"] or 0) if agg else money.D(0)
    vl = money.D(agg["vl"] or 0) if agg else money.D(0)
    ccy = (agg["ccy"] if agg and agg["ccy"] else (row["currency"] if row else None)) or "EUR"
    return ve, vl, ccy

def below_minimum(con, ent, ctry, period):
    """Is this claim below the refund country's statutory minimum (Dir. 2008/9/EC
    Art. 17), enforced in the country's currency? Returns (below: bool, detail: str).

    A `-YEAR` period uses the ANNUAL minimum, else the QUARTERLY (sub-year) minimum.
    The basis (national-currency vat_local vs EUR vat_eur) is chosen by
    vat_config.min_for: countries with a fixed national amount (NATIONAL_MINIMUMS,
    e.g. Sweden/Denmark) compare in local currency; euro countries and Poland fall
    back to the EUR base on vat_eur. The Decimal threshold comparison uses money.q2
    (the EUR/threshold quantizer). Never raises."""
    try:
        is_annual = str(period).endswith("-YEAR")
        ccy, threshold, basis = vat_config.min_for(ctry, is_annual)
        ve, vl, _stream_ccy = _stream_vat(con, ent, ctry, period)
        amount = vl if basis == "local" else ve
        thr = money.q2(money.D(threshold))
        amt = money.q2(amount)
        kind = "annual" if is_annual else "quarterly"
        detail = (f"below the {ccy} {threshold:,.0f} {kind} minimum "
                  f"(this claim: {ccy} {amt:,.2f}) — defer to the annual claim or override")
        return (amt < thr), detail
    except Exception as e:
        # Never let the threshold check break the submit path; log and treat as not-below
        # (the existing checklist/period gates still protect submission).
        log.warning("vat_refund: below_minimum failed for %s/%s/%s: %s", ent, ctry, period, e)
        return False, ""

def filing_deadline(period):
    """Statutory 2008/9/EC filing deadline for a claim period: 30 September of the
    following year. Returns a datetime.date."""
    import datetime
    return datetime.date(int(str(period)[:4]) + 1, 9, 30)

def suggested_next(code, payout_to=None):
    """The recommended next manual step after `code` (None if terminal/no suggestion).
    After the money arrives (3A) the route decides: refund to customer → 4 invoice the
    fee; refund to us → 4A credit the customer the net."""
    return {
        "2": "2A", "2A": None, "2B": None,
        "3": "3A",
        "3A": ("4A" if (payout_to or "customer") == "us" else "4"),
        "3B": "3D", "3D": None,
        "3C": ("4A" if (payout_to or "customer") == "us" else "4"),
        "4": "5", "4A": "5", "5": None,
    }.get(code)

def set_status_code(con, ent, ctry, period, code, note=None, deadline=None,
                    override_threshold=False):
    """Advance a claim along the controllable workflow. Pre-submission codes are
    system-controlled (rejected here). Submitting (2) is HARD-GATED on the system
    checklist + period end + the refund-country minimum (Dir. 2008/9/EC Art. 17,
    enforced in national currency). 3B/3C/3D keep the invoice locks.

    `note` records WHY (rejection reason, what documents were requested, appeal
    grounds); `deadline` (ISO date) records the open action's deadline — the
    document-request response date (2B) or the appeal deadline (3D).

    `override_threshold` (admin-only at the route layer) lets a below-minimum claim be
    submitted anyway; the override is recorded in the claim's status_note for an audit
    trail. It has no effect on any other gate."""
    code = (code or "").strip()
    if code in AUTO_CODES:
        return False, f"'{code} {STATUS_LABELS.get(code,'')}' is system-controlled — it follows the checklist automatically"
    if code not in MANUAL_CODES:
        return False, f"unknown status '{code}'"
    row = con.execute("""SELECT status FROM vat_applications WHERE entity=? AND
                         refund_country=? AND ref_period=?""", (ent, ctry, period)).fetchone()
    eng_now = row["status"] if row else "draft"
    if eng_now not in LOCKING:
        # not yet locked: the only legal first manual step is Submit (2), and only when
        # the SYSTEM says the checklist is complete and the period has ended.
        if code != "2":
            return False, f"can't set '{STATUS_LABELS[code]}' before the claim is submitted"
        stage, items = derive_stage(con, ent, ctry, period)
        if stage == "1A":
            missing = [l for l, ok in items if not ok and l != "Claim period ended"]
            return False, "BLOCKED — checklist incomplete: " + "; ".join(missing)
        if stage == "1B":
            return False, ("BLOCKED — the claim period has not ended yet (ends "
                           f"{period_end_date(period)})")
        # Hard period-end + checklist gates passed: now the refund-country minimum
        # (Art. 17, in national currency). Below the minimum a claim would be rejected
        # and its invoices locked out of the annual mop-up — block it, unless an admin
        # overrides (recorded in the status_note for traceability).
        below, why = below_minimum(con, ent, ctry, period)
        if below and not override_threshold:
            return False, "BLOCKED — " + why
        if below and override_threshold:
            actor = audit._current_actor() or "admin"
            ovr = f"minimum-threshold overridden by {actor}"
            note = (f"{note}; {ovr}" if note else ovr)
        # Record the WAIVER use on submission (mirror the threshold-override note): the
        # claim was filed excluding genuinely-uninvoiced suppliers (R5 case (a)).
        waived = list_waivers(con, ent, ctry, period)
        if waived:
            wnote = "filed excluding waived suppliers: " + ", ".join(sorted(waived))
            note = (f"{note}; {wnote}" if note else wnote)
    engine = ENGINE_OF.get(code)
    if engine:
        ok, msg = set_status(con, ent, ctry, period, engine, gate_activation=False)
        if not ok:
            return ok, msg
    con.execute("""UPDATE vat_applications SET status_code=?, updated=CURRENT_TIMESTAMP
                   WHERE entity=? AND refund_country=? AND ref_period=?""",
                (code, ent, ctry, period))
    # per-status data: decision date on first decision code; the note; the open-action
    # deadline lives only while a 2B/3D is open (cleared when the claim moves on).
    if code in ("3", "3A", "3B", "3C"):
        con.execute("""UPDATE vat_applications SET decision_date=COALESCE(decision_date, CURRENT_DATE)
                       WHERE entity=? AND refund_country=? AND ref_period=?""", (ent, ctry, period))
    if note:
        con.execute("""UPDATE vat_applications SET status_note=? WHERE entity=? AND
                       refund_country=? AND ref_period=?""", (str(note)[:500], ent, ctry, period))
    con.execute("""UPDATE vat_applications SET action_deadline=? WHERE entity=? AND
                   refund_country=? AND ref_period=?""",
                ((deadline or None) if code in ("2B", "3D") else None, ent, ctry, period))
    con.commit()
    return True, f"status → {code} {STATUS_LABELS[code]}" + (f" (note recorded)" if note else "")

def record_payment(con, ent, ctry, period, amount, date=None):
    """Record the ACTUALLY-REFUNDED amount on a claim and move it to 'money received'
    (3A). The service fee is contingency on the PAID amount, not the full claimed VAT:
    we stamp paid_amount/paid_date, then drive the claim to 3A so the EXISTING paid
    recompute (in set_status, the `new == "paid"` branch) recomputes
    fee_eur = compute_fee(paid_amount, frozen fee_pct, frozen fee_min). The frozen
    rate/minimum are NOT re-derived — only the fee BASE changes from claimed→paid.

    ATOMICITY (M5a): the paid_amount stamp and the paid-transition (status='paid',
    paid_date, fee recompute) commit as ONE transaction. The UPDATE below is left
    UNcommitted; set_status reads it on the SAME connection and commits the stamp +
    the recompute together (the deferred-transaction `connect()` is NOT autocommit, so
    the pending UPDATE does not leak out). On ANY failure set_status rolls back and
    discards the pending stamp, so a crash can never leave paid_amount stamped while the
    status/fee lag. The benign residual: a crash between set_status's commit and the
    status_code='3A' display write below leaves a CONSISTENT money state (paid + fee
    recomputed) with only the 3A display code lagging — recoverable by re-running.
    Returns (ok, message)."""
    try:
        amt = money.f2(amount)
    except (TypeError, ValueError):
        return False, "invalid amount"
    if amt < 0:
        return False, "amount must be >= 0"
    row = con.execute("""SELECT status FROM vat_applications WHERE entity=? AND
                         refund_country=? AND ref_period=?""", (ent, ctry, period)).fetchone()
    if not row:
        return False, "no such claim"
    if (row["status"] or "draft") not in LOCKING:
        return False, "claim must be submitted before a payment can be recorded"
    # Stamp the refunded amount FIRST so the paid-transition recompute reads it (the
    # recompute falls back to the full claimed vat_eur when paid_amount is null).
    con.execute("""UPDATE vat_applications SET paid_amount=?, updated=CURRENT_TIMESTAMP
                   WHERE entity=? AND refund_country=? AND ref_period=?""",
                (amt, ent, ctry, period))
    # NO intermediate commit: the stamp stays pending so it commits ATOMICALLY with the
    # paid-recompute inside set_status (its con.commit()), and set_status's rollback
    # discards this UPDATE on any failure (M5a).
    # Drive to 3A 'Money received' — ENGINE_OF['3A']='paid', so set_status fires the
    # canonical paid recompute on the frozen rate. REUSE that path; never re-derive
    # the fee formula here.
    ok, msg = set_status_code(con, ent, ctry, period, "3A")
    if not ok:
        # Belt-and-suspenders: discard any still-pending paid_amount stamp if an early
        # return from set_status_code did not itself roll back (for code '3A' on a locked
        # claim these guards don't trigger, but never leave a half-write).
        con.rollback()
        return ok, msg
    # Stamp the explicit refund date AFTER the transition (set_status stamps paid_date
    # = CURRENT_DATE; an explicitly supplied date overrides it).
    if date:
        con.execute("""UPDATE vat_applications SET paid_date=? WHERE entity=? AND
                       refund_country=? AND ref_period=?""", (date, ent, ctry, period))
        con.commit()
    r = con.execute("""SELECT fee_eur FROM vat_applications WHERE entity=? AND
                       refund_country=? AND ref_period=?""", (ent, ctry, period)).fetchone()
    fee = money.f2(r["fee_eur"]) if r and r["fee_eur"] is not None else 0.0
    return True, f"payment €{amt:,.2f} recorded — fee €{fee:,.2f}"

def withdraw_claim(con, ent, ctry, period):
    """Admin escape hatch: cancel a claim and RELEASE its invoice locks (the only path
    that frees invoices — rejection/confiscation/appeal keep them)."""
    ok, msg = set_status(con, ent, ctry, period, "withdrawn", gate_activation=False)
    if ok:
        con.execute("""UPDATE vat_applications SET status_code=NULL, updated=CURRENT_TIMESTAMP
                       WHERE entity=? AND refund_country=? AND ref_period=?""", (ent, ctry, period))
        con.commit()
    return ok, msg

def settlement(payout_to, refund_eur, fee_eur):
    """How the fee is settled. payout_to='customer' -> we invoice the fee (receivable);
    payout_to='us' -> we deduct the fee and remit the net to the customer."""
    refund = float(refund_eur or 0); fee = money.f2(fee_eur or 0)
    if payout_to == "us":
        return {"route": "us", "refund": money.f2(refund), "fee": fee,
                "net_to_customer": money.f2(refund - fee), "fee_receivable": 0.0}
    return {"route": "customer", "refund": money.f2(refund), "fee": fee,
            "net_to_customer": 0.0, "fee_receivable": fee}

def issue_fee_invoice(con, ent, ctry, period):
    """Assign a fee-invoice number/date to a paid claim. Returns (ok, number_or_msg)."""
    r = con.execute("""SELECT fee_billed_date, fee_invoice_no FROM vat_applications
                       WHERE entity=? AND refund_country=? AND ref_period=?""",
                    (ent, ctry, period)).fetchone()
    if not r:
        return (False, "no such claim")
    if not r["fee_billed_date"]:
        return (False, "fee not charged yet — the refund must be paid first")
    if r["fee_invoice_no"]:
        return (True, r["fee_invoice_no"])               # already issued (idempotent)
    yr = str(period).split("-")[0]
    n = con.execute("SELECT COUNT(*) FROM vat_applications WHERE fee_invoice_no IS NOT NULL"
                    ).fetchone()[0] + 1
    inv_no = f"F{yr}-{n:04d}"
    con.execute("""UPDATE vat_applications SET fee_invoice_no=?, fee_invoice_date=CURRENT_DATE
                   WHERE entity=? AND refund_country=? AND ref_period=?""",
                (inv_no, ent, ctry, period))
    con.commit()
    return (True, inv_no)

def submission_readiness(con, ent, ctry, period, cache=None):
    """Read-only check of whether a claim CAN be submitted. Returns (ready, [issues]).
    Mirrors the blocking conditions in set_status without writing anything. `cache`
    (optional) is shared across many streams to avoid per-stream connections and
    per-invoice queries (one docs/locks index + memoised activation lookups)."""
    cache = cache if cache is not None else {}
    issues = []
    ia = cache.setdefault("_isactive", {})
    if ent not in ia:
        ia[ent] = customer_master.is_active(ent)
    if ia[ent] is False:
        issues.append("customer not activated")
    ca = cache.setdefault("_ctryactive", {})
    if (ent, ctry) not in ca:
        ca[(ent, ctry)] = customer_master.country_active(ent, ctry)
    if ca[(ent, ctry)] is False:
        issues.append(f"refund country '{ctry}' not activated")
    invs = stream_invoices(con, ent, ctry, period, cache)
    # Receipt-control WAIVERS: a waived, genuinely-uninvoiced supplier is excluded from
    # the claim by set_status; mirror that here so the readiness verdict the UI shows
    # agrees with what set_status will actually do (else a waived stream reads "blocked"
    # while submission would succeed). Only drops a (waivable AND waived) synthetic ref.
    waived = list_waivers(con, ent, ctry, period)
    if waived:
        scon = cache.get("_scon")
        if scon is None:
            scon = cache["_scon"] = supplier_master.connect()
        invs = [(s, r) for s, r in invs
                if not (s in waived and _waivable_missing(r, s, ctry, scon))]
    bad = [r for s, r in invs if _synthetic(r)]
    if bad:
        issues.append(f"{len(bad)} unresolved invoice ref(s)")
    docidx = cache.get("_docidx")
    if docidx is None:
        docidx = cache["_docidx"] = docs_index(con)
    nodoc = [(s, r) for s, r in invs if (ent, s, r) not in docidx]
    if nodoc:
        issues.append(f"{len(nodoc)} invoice(s) missing documents")
    locks = cache.get("_locks")
    if locks is None:
        locks = cache["_locks"] = {
            (r["entity"], r["refund_country"], r["supplier"], r["invoice_ref"]): r["ref_period"]
            for r in con.execute("""SELECT entity, refund_country, supplier, invoice_ref, ref_period
                                    FROM vat_claimed_invoices""")}
    conflicts = [(s, r) for s, r in invs if locks.get((ent, ctry, s, r), period) != period]
    if conflicts:
        issues.append(f"{len(conflicts)} invoice(s) locked by another claim")
    return (len(issues) == 0, issues)

def claims_overview(year):
    """For the VAT-refund 'can we submit?' report: every claimable quarter split
    into TO-SUBMIT (with a readiness verdict) and OPEN (submitted/approved, aging)."""
    import datetime
    con = connect()
    matrix = claim_matrix(con, year, with_portal=False)   # /, /readiness never read m["home"]
    apps = {(r["entity"], r["refund_country"], r["ref_period"]): r
            for r in con.execute("""SELECT entity, refund_country, ref_period, status,
                                    submitted_date, status_code, action_deadline, status_note
                                    FROM vat_applications""")}
    today = datetime.date.today()
    to_submit, open_claims = [], []
    cache = {}   # shared across all streams: one supplier + one analytics connection,
                 # one docs/locks index, memoised activation lookups
    for m in matrix:
        if m["period"].endswith("YEAR") or (m["vat_eur"] or 0) <= 0:
            continue
        key = (m["entity"], m["country"], m["period"])
        a = apps.get(key)
        status = a["status"] if a else "draft"
        if status in ("submitted", "approved"):
            age = ""
            sd = a["submitted_date"] if a else None
            if sd:
                try: age = (today - datetime.date.fromisoformat(sd)).days
                except ValueError as e:
                    log.debug("vat_refund: unparseable submitted_date %r: %s", sd, e)
            code = (a["status_code"] if a else None) or {"submitted": "2", "approved": "3"}[status]
            open_claims.append(dict(entity=m["entity"], country=m["country"], period=m["period"],
                                    vat_eur=m["vat_eur"], status=status, submitted=sd, age_days=age,
                                    code=code, code_label=STATUS_LABELS.get(code, code),
                                    action_deadline=a["action_deadline"] if a else None,
                                    note=a["status_note"] if a else None))
        elif status not in ("paid", "rejected", "withdrawn"):
            ready, issues = submission_readiness(con, m["entity"], m["country"], m["period"], cache)
            if not m["verdict"].startswith("READY"):
                issues = issues + [m["verdict"].split(" (")[0].lower()]
            stage, items = derive_stage(con, m["entity"], m["country"], m["period"],
                                        m["verdict"], cache)
            # surface the failed CUSTOMER-checklist rules (the invoice-level ones are
            # already covered by submission_readiness) and the period-end gate
            overlap = {"All invoice refs resolved (no INPUT/aggregate placeholders)",
                       "All invoice documents attached", "Claim period ended"}
            issues = issues + [l for l, ok in items if not ok and l not in overlap
                               and not l.startswith("Receipt control:")]
            if stage == "1B":
                issues = issues + [f"period not ended (ends {period_end_date(m['period'])})"]
            ready = (len(issues) == 0)
            fdl = filing_deadline(m["period"])
            to_submit.append(dict(entity=m["entity"], country=m["country"], period=m["period"],
                                  vat_eur=m["vat_eur"], verdict=m["verdict"],
                                  ready=ready, issues=issues, missing=m["missing"],
                                  code=stage, code_label=STATUS_LABELS.get(stage, stage),
                                  deadline=fdl.isoformat(),
                                  deadline_days=(fdl - today).days))
    for k in ("_scon", "_acon", "_cmcon"):           # close the shared connections opened lazily
        if cache.get(k) is not None:
            try: cache[k].close()
            except Exception as e:
                log.debug("vat_refund: closing shared %s failed: %s", k, e)
    con.close()
    return {"to_submit": to_submit, "open": open_claims}

def claim_matrix(con, year, with_portal=True):
    """All streams for the year: per (entity, country) give Q1..Q4 + YEAR VAT, currency, status.
    `con` is the claims connection; transactions are read from the analytics DB.

    `home` (the entity's customer-portal URL) is only consumed on /vat and in the
    claim workbook; callers that ignore it (claims_overview → / and /readiness) pass
    with_portal=False to skip the per-entity customer_master.portal() lookups (each is
    a fresh connection + 2 queries). When True the lookup is memoised per entity."""
    acon = analytics_connect()
    rows = acon.execute("""
        SELECT entity, country, currency, period,
               ROUND(SUM(vat_eur),2) ve, ROUND(SUM(vat_local),2) vl,
               COUNT(*) n
        FROM transactions WHERE period LIKE ? GROUP BY entity, country, period""",
        (f"{year}-%",)).fetchall()
    acon.close()
    streams = collections.defaultdict(lambda: {"qs": collections.defaultdict(
                                                   lambda: [money.D(0), money.D(0), 0]),
                                               "ccy": "EUR"})
    loaded_periods = set()
    for r in rows:
        s = streams[(r["entity"], r["country"])]
        q = quarter(r["period"])
        # accumulate VAT exactly as Decimal so the EUR-threshold test below never
        # flips on binary-float noise
        s["qs"][q][0] += money.D(r["ve"]); s["qs"][q][1] += money.D(r["vl"]); s["qs"][q][2] += r["n"]
        s["ccy"] = r["currency"]; loaded_periods.add(r["period"])
    out = []
    portals = {}   # memoise portal(ent) once per entity (used for both quarter & YEAR rows)
    def home_of(ent):
        if not with_portal:
            return None
        if ent not in portals:
            portals[ent] = customer_master.portal(ent)
        return portals[ent]
    # The threshold verdict is computed on the SAME national-currency basis as the
    # submission gate (vat_refund.below_minimum → vat_config.min_for): Sweden/Denmark
    # compare in local currency (vat_local), euro countries and Poland on the EUR base
    # (vat_eur). EUR figures are still shown; only the READY/DEFER/BELOW decision moves.
    for (ent, ctry), s in sorted(streams.items()):
        year_ve = money.q2(sum((v[0] for v in s["qs"].values()), money.D(0)))
        year_vl = money.q2(sum((v[1] for v in s["qs"].values()), money.D(0)))
        q_ccy, q_thr, q_basis = vat_config.min_for(ctry, is_annual=False)
        a_ccy, a_thr, a_basis = vat_config.min_for(ctry, is_annual=True)
        year_amt = year_vl if a_basis == "local" else year_ve
        year_ok = money.q2(year_amt) >= money.q2(money.D(a_thr))
        for q, (ve, vl, n) in sorted(s["qs"].items()):
            ve = money.q2(ve)            # quarterly VAT, exact cents
            vl = money.q2(vl)
            missing = [m for m in q_months(q) if m not in loaded_periods]
            q_amt = vl if q_basis == "local" else ve
            if money.q2(q_amt) >= money.q2(money.D(q_thr)):
                verdict = f"READY (>= {q_ccy} {q_thr:,.0f} quarterly min)"
            elif year_ok:
                verdict = (f"DEFER TO ANNUAL (below {q_ccy} {q_thr:,.0f}, "
                           f"year >= {a_ccy} {a_thr:,.0f})")
            else:
                verdict = "BELOW ANNUAL MIN - accumulate"
            out.append(dict(entity=ent, country=ctry, period=q, vat_eur=money.f2(ve),
                            vat_local=money.f2(vl), currency=s["ccy"], lines=n,
                            verdict=verdict, missing=missing,
                            home=home_of(ent),
                            deadline=DEADLINE_FMT.format(year_plus1=int(year)+1)))
        out.append(dict(entity=ent, country=ctry, period=f"{year}-YEAR",
                        vat_eur=money.f2(year_ve), vat_local=money.f2(year_vl),
                        currency=s["ccy"], lines=sum(v[2] for v in s["qs"].values()),
                        verdict=(f"READY (annual >= {a_ccy} {a_thr:,.0f})" if year_ok
                                 else "BELOW ANNUAL MIN"),
                        missing=[], home=home_of(ent),
                        deadline=DEADLINE_FMT.format(year_plus1=int(year)+1)))
    return out

def invoice_lines(con, ent, ctry, qtr, cache=None):
    """Invoice-level detail for one claim: prefer per-invoice split via the note column,
    fall back to registry invoice(s) carrying the country aggregate.

    `cache` (optional dict) lets a caller rendering many claims share one
    supplier-DB connection and memoize per-(supplier, country) issuer/invoice
    lookups across calls, avoiding an N+1 of connections/queries."""
    own_cache = cache is None
    cache = cache if cache is not None else {}
    scon = cache.get("_scon")
    if scon is None:
        scon = supplier_master.connect(); cache["_scon"] = scon
    acon = cache.get("_acon")             # analytics (transactions) connection
    if acon is None:
        acon = analytics_connect(); cache["_acon"] = acon
    months = q_months(qtr)
    sups = [r[0] for r in acon.execute(
        """SELECT DISTINCT supplier FROM transactions WHERE entity=? AND country=?
           AND period IN (%s)""" % ",".join("?"*len(months)), [ent, ctry]+months)]
    lines = []
    for sup in sups:
        ck = (sup, ctry)
        if ck in cache:
            issuer, vatid, vnote, regs = cache[ck]
        else:
            issuer, vatid, vnote = supplier_master.get_issuer(sup, ctry, con=scon)
            regs = supplier_master.get_invoices(sup, ctry, con=scon)
            cache[ck] = (issuer, vatid, vnote, regs)
        refs = [r[0] for r in regs]
        rows = acon.execute(
            """SELECT note, product_group, ROUND(SUM(net_eur),2) net, ROUND(SUM(vat_eur),2) vat,
                      ROUND(SUM(net_local),2) netl, ROUND(SUM(vat_local),2) vatl, currency
               FROM transactions WHERE entity=? AND country=? AND supplier=?
               AND period IN (%s) GROUP BY note, product_group""" % ",".join("?"*len(months)),
            [ent, ctry, sup]+months).fetchall()
        by_inv = collections.defaultdict(lambda: collections.defaultdict(lambda: [0,0,0,0,""]))
        for r in rows:
            inv = next((ref for ref in refs if r["note"] and r["note"].split("/")[0] in r["note"]
                        and ref.startswith(r["note"].split(" ")[0])), None)
            inv = inv or next((ref for ref in refs if r["note"] and ref.split("/")[0] in r["note"]), None)
            # No note match: resolve to the sole registered invoice if there is
            # exactly one (legitimate); otherwise (zero or several) tag the row
            # UNMATCHED so the gates treat it as a hard block instead of inventing
            # an ALL: aggregate. UNMATCHED carries its VAT into the row.
            inv = inv or (refs[0] if len(refs) == 1 else "UNMATCHED")
            a = by_inv[inv][r["product_group"]]
            a[0] += r["net"]; a[1] += r["vat"]; a[2] += r["netl"]; a[3] += r["vatl"]; a[4] = r["currency"]
        dates = dict(regs)
        for inv, prods in by_inv.items():
            for pg, (net, vat, netl, vatl, ccy) in sorted(prods.items()):
                # Unknown product group -> goods code "10" (Other, recoverable), NOT
                # "9" (luxuries/entertainment, NEVER VAT-recoverable; 2008/9/EC Art. 9,
                # Reg. 79/2012). Defaulting to 9 would silently file an unclassified
                # product under the one non-refundable code.
                code, desc = GOODS_CODE.get(pg, ("10", "Other"))
                lines.append(dict(supplier=sup, issuer=issuer,
                                  vat_id=vatid or "INPUT: " + vnote,
                                  invoice=inv, inv_date=dates.get(inv, ""),
                                  code=code, desc=desc, product=pg, currency=ccy,
                                  net_local=money.f2(netl), vat_local=money.f2(vatl),
                                  net_eur=money.f2(net), vat_eur=money.f2(vat)))
    if own_cache:
        for k in ("_scon", "_acon", "_cmcon"):
            if cache.get(k) is not None:
                cache[k].close()
    return lines

# ---------------------------------------------------------------- Excel pack
def build_workbook(con, year):
    matrix = claim_matrix(con, year)
    wb = Workbook()
    fillH = PatternFill("solid", start_color="5B3A8E")
    fillY = PatternFill("solid", start_color="FFF2CC")
    bw = Font(bold=True, color="FFFFFF", name="Arial", size=9)
    b10 = Font(bold=True, name="Arial", size=10)
    norm = Font(name="Arial", size=9)
    it8 = Font(italic=True, name="Arial", size=8)
    def head(ws, row):
        for c in ws[row]:
            if c.value: c.font = bw; c.fill = fillH; c.alignment = Alignment(horizontal="center", wrap_text=True)

    ws = wb.active; ws.title = "Overview"
    ws["A1"] = f"VAT REFUND APPLICATIONS {year} - EU Directive 2008/9/EC (deadline {int(year)+1}-09-30)"
    ws["A1"].font = Font(bold=True, size=12, name="Arial")
    ws.append([])
    ws.append(["Entity","Refund country","Period","VAT EUR","VAT local","Ccy","Txn lines",
               "Threshold verdict","Months not yet loaded","Home portal","Status"])
    head(ws, 3)
    r = 4
    for m in matrix:
        st = con.execute("""SELECT status FROM vat_applications WHERE entity=? AND
                            refund_country=? AND ref_period=?""",
                         (m["entity"], m["country"], m["period"])).fetchone()
        status = st["status"] if st else "draft"
        # Refresh the all-period recompute into the DRAFT row (or INSERT a new draft),
        # but NEVER overwrite a SUBMITTED/approved/paid stream: its vat_eur/vat_local was
        # FROZEN at submission over EXACTLY the locked claim_set (see set_status freeze,
        # ~:464). claim_matrix sums over ALL period transactions, so a refresh here would
        # clobber the frozen base (and the fee base downstream) whenever some invoices are
        # locked to another claim. Guard the UPDATE half on status; the INSERT half is
        # untouched (a not-yet-existing claim still gets its draft row created).
        con.execute("""INSERT INTO vat_applications (entity, refund_country, ref_period,
                       vat_eur, vat_local, currency, status) VALUES (?,?,?,?,?,?,?)
                       ON CONFLICT(entity, refund_country, ref_period) DO UPDATE SET
                       vat_eur=excluded.vat_eur, vat_local=excluded.vat_local
                       WHERE vat_applications.status NOT IN ('submitted','approved','paid')""",
                    (m["entity"], m["country"], m["period"], m["vat_eur"], m["vat_local"],
                     m["currency"], status))
        ws.append([m["entity"], m["country"], m["period"], m["vat_eur"], m["vat_local"],
                   m["currency"], m["lines"], m["verdict"],
                   ", ".join(m["missing"]) if m["missing"] else "", m["home"], status])
        for c in ws[r]: c.font = norm
        ws.cell(row=r, column=4).number_format = "#,##0.00"
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        if m["missing"]:
            ws.cell(row=r, column=9).fill = fillY
        if (m["entity"], m["country"]) in LOCAL_CCY_INPUT and not m["period"].endswith("YEAR"):
            ws.cell(row=r, column=5).fill = fillY
            ws.cell(row=r, column=5).value = f"INPUT {LOCAL_CCY_INPUT[(m['entity'], m['country'])]}"
        r += 1
    con.commit()
    r += 1
    ws[f"A{r}"] = "Compliance notes:"; ws[f"A{r}"].font = b10
    for i, n in enumerate(COMPLIANCE_NOTES):
        ws[f"B{r+1+i}"] = "- " + n; ws[f"B{r+1+i}"].font = it8
        ws[f"B{r+1+i}"].alignment = Alignment(wrap_text=True)
    for col, w in zip("ABCDEFGHIJK",[26,13,10,11,12,6,9,34,18,30,10]):
        ws.column_dimensions[col].width = w

    # one sheet per quarterly claim stream with VAT > 0
    done = set()
    pack_cache = {}   # share one supplier + analytics connection across all claim packs
    for m in matrix:
        if m["period"].endswith("YEAR") or m["vat_eur"] <= 0: continue
        key = (m["entity"], m["country"], m["period"])
        if key in done: continue
        done.add(key)
        toks = m["entity"].replace("UAB ","").replace("SIA ","").split()
        sheetname = f"{toks[0][:12]}-{m['country'][:8]}-{m['period'][-2:]}"
        ws2 = wb.create_sheet(sheetname[:31])
        ws2["A1"] = f"CLAIM PACK: {m['entity']} -> {m['country']} -> {m['period']}"
        ws2["A1"].font = Font(bold=True, size=11, name="Arial")
        ws2["A2"] = (f"File via {m['home']}; claim currency {m['currency']}; "
                     f"deadline {m['deadline']}; verdict: {m['verdict']}"
                     + ("; QUARTER INCOMPLETE - missing " + ", ".join(m["missing"]) if m["missing"] else ""))
        ws2["A2"].font = it8
        cust = customer_master.get_customer(m["entity"])
        ws2["A3"] = (f"APPLICANT: {cust['company_name']} | Reg. no: {cust['reg_number']} | "
                     f"VAT: {cust['vat_number']} | Legal address: {cust['legal_address']}")
        ws2["A4"] = f"Refund payout account: {cust['payout']}"
        for cell in ("A3","A4"):
            ws2[cell].font = norm
            if "INPUT" in str(ws2[cell].value):
                ws2[cell].fill = fillY
        ws2.append([])
        ws2.append(["Supplier","Issuer (invoice party)","Issuer VAT ID","Invoice ref","Invoice date",
                    "Goods code","Code description","Product","Ccy","Taxable base (local)",
                    "VAT (local)","Net EUR","VAT EUR","Duplicate control","Document(s) attached"])
        head(ws2, 6)
        rr = 7
        lines = invoice_lines(con, m["entity"], m["country"], m["period"], pack_cache)
        synth = [L for L in lines if _synthetic(L["invoice"], L.get("vat_id"))]
        if synth:
            # Refuse to emit data rows for a pack with any line not tied to ONE
            # documented invoice (INPUT / ALL: / UNMATCHED). Render a single bold
            # red BLOCKED banner and a zero TOTAL so the pack can never be filed
            # as-is. RENDER-ONLY: the frozen amount in vat_applications is untouched.
            cell = ws2.cell(row=rr, column=1)
            cell.value = (f"BLOCKED - {len(synth)} line(s) not tied to a documented "
                          "invoice; resolve before filing")
            cell.font = Font(bold=True, color="C00000", name="Arial", size=10)
            rr += 1
            ws2.append(["TOTAL","","","","","","","","", 0, 0, 0, 0])
            for c in ws2[rr]: c.font = b10
            for col in (10,11,12,13): ws2.cell(row=rr, column=col).number_format = "#,##0.00"
            for col, w in zip("ABCDEFGHIJKLMNO",[8,30,22,22,11,7,26,11,5,13,11,11,11,20,40]):
                ws2.column_dimensions[col].width = w
            continue
        for L in lines:
            other = lock_state(con, m["entity"], m["country"], L["supplier"], L["invoice"])
            L["lock"] = ("LOCKED here" if other == m["period"]
                         else f"EXCLUDE - claimed in {other}" if other else "free")
            ws2.append([L["supplier"], L["issuer"], L["vat_id"], L["invoice"], L["inv_date"],
                        L["code"], L["desc"], L["product"], L["currency"],
                        L["net_local"], L["vat_local"], L["net_eur"], L["vat_eur"], L["lock"],
                        "; ".join(f"{d['filename']} [sha {d['sha256'][:8]}]"
                                  for d in docs_for(con, m["entity"], L["supplier"], L["invoice"]))
                        or "MISSING"])
            if not docs_for(con, m["entity"], L["supplier"], L["invoice"]):
                ws2.cell(row=rr, column=15).fill = PatternFill("solid", start_color="F8CBAD")
            if L["lock"].startswith("EXCLUDE"):
                ws2.cell(row=rr, column=14).fill = PatternFill("solid", start_color="F8CBAD")
            for c in ws2[rr]: c.font = norm
            for col in (10,11,12,13): ws2.cell(row=rr, column=col).number_format = "#,##0.00"
            if "INPUT" in str(L["vat_id"]) or "INPUT" in str(L["invoice"]):
                ws2.cell(row=rr, column=3 if "INPUT" in str(L["vat_id"]) else 4).fill = fillY
            rr += 1
        ws2.append(["TOTAL","","","","","","","","",
                    f"=SUM(J7:J{rr-1})", f"=SUM(K7:K{rr-1})", f"=SUM(L7:L{rr-1})", f"=SUM(M7:M{rr-1})"])
        for c in ws2[rr]: c.font = b10
        for col in (10,11,12,13): ws2.cell(row=rr, column=col).number_format = "#,##0.00"
        for col, w in zip("ABCDEFGHIJKLMNO",[8,30,22,22,11,7,26,11,5,13,11,11,11,20,40]):
            ws2.column_dimensions[col].width = w

    path = f"{WORKDIR}/VAT_Refund_Claims_{year}.xlsx"
    for k in ("_scon", "_acon", "_cmcon"):
        if pack_cache.get(k) is not None:
            try: pack_cache[k].close()
            except Exception as e:
                log.debug("vat_refund: closing shared %s failed: %s", k, e)
    wb.save(path)
    return path, matrix

if __name__ == "__main__":
    con = connect()
    if len(sys.argv) >= 5 and sys.argv[1] == "--set-status":
        ent, ctry, per, st = sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5]
        ok, msg = set_status(con, ent, ctry, per, st)
        print(("OK: " if ok else "") + msg)
    else:
        year = sys.argv[1] if len(sys.argv) > 1 else "2026"
        path, matrix = build_workbook(con, year)
        print(f"{'Entity':28}{'Country':10}{'Period':10}{'VAT EUR':>10}  Verdict")
        for m in matrix:
            print(f"{m['entity'][:26]:28}{m['country'][:9]:10}{m['period']:10}{m['vat_eur']:>10,.2f}  {m['verdict']}"
                  + ("  [missing: " + ",".join(m['missing']) + "]" if m['missing'] else ""))
        print("saved:", path)
    con.close()


def recovery_report(year=None):
    """Submitted vs approved vs paid, with aging of unpaid submitted claims."""
    con = connect()
    rows = con.execute("""SELECT entity, refund_country, ref_period, vat_eur, status,
        submitted_date, approved_date, paid_date, paid_amount,
        fee_eur, fee_pct, fee_min, fee_billed_date, payout_to, fee_invoice_no, fee_invoice_date,
        status_code, decision_date, status_note, action_deadline
        FROM vat_applications WHERE status IN ('submitted','approved','paid')
        AND (? IS NULL OR ref_period LIKE ?) ORDER BY submitted_date""",
        (year, f"{year}-%" if year else None)).fetchall()
    import datetime
    today = datetime.date.today()
    out = []
    for r in rows:
        age = ""
        if r["status"] in ("submitted", "approved") and r["submitted_date"]:
            try:
                age = (today - datetime.date.fromisoformat(r["submitted_date"])).days
            except ValueError as e:
                log.debug("vat_refund: unparseable submitted_date %r: %s", r["submitted_date"], e)
        code = r["status_code"] or {"submitted": "2", "approved": "3", "paid": "3A"}[r["status"]]
        out.append(dict(entity=r["entity"], country=r["refund_country"], period=r["ref_period"],
                        vat_eur=r["vat_eur"], status=r["status"], submitted=r["submitted_date"],
                        paid=r["paid_date"], paid_amount=r["paid_amount"], age_days=age,
                        fee_eur=r["fee_eur"], fee_pct=r["fee_pct"], fee_min=r["fee_min"],
                        fee_billed_date=r["fee_billed_date"], payout_to=r["payout_to"],
                        fee_invoice_no=r["fee_invoice_no"], fee_invoice_date=r["fee_invoice_date"],
                        status_code=code, next_code=suggested_next(code, r["payout_to"]),
                        decision_date=r["decision_date"], status_note=r["status_note"],
                        action_deadline=r["action_deadline"]))
    con.close()
    summary = {s: money.fsum(o["vat_eur"] or 0 for o in out if o["status"] == s)
               for s in ("submitted", "approved", "paid")}
    summary["outstanding"] = money.fsum(o["vat_eur"] or 0 for o in out
                                        if o["status"] in ("submitted", "approved"))
    return out, summary


def _aging_band(days):
    """Aging band for an open (unpaid submitted/approved) receivable, by days since
    submission. Returns one of '0-30','30-60','60-90','90+' (a 45-day-old claim is
    '30-60'; the lower bound is inclusive, the upper exclusive)."""
    if not isinstance(days, int):
        return ""
    if days < 30:
        return "0-30"
    if days < 60:
        return "30-60"
    if days < 90:
        return "60-90"
    return "90+"


def _median(values):
    """Median of a list of numbers (None for an empty list). No external deps so the
    module stays import-light; sorts and averages the two middle values for an even
    count."""
    xs = sorted(v for v in values if v is not None)
    if not xs:
        return None
    n = len(xs)
    mid = n // 2
    if n % 2:
        return xs[mid]
    return (xs[mid - 1] + xs[mid]) / 2.0


AGING_BANDS = ("0-30", "30-60", "60-90", "90+")


def receivables_forecast(year=None):
    """VAT-receivable / financing-ready view for the ADMIN VAT surface — an INTERNAL,
    data-only forecast (no lending, no outward send).

    Builds on the same submitted/approved/paid claim base + frozen fee fields as
    recovery_report (so the fee math is never recomputed differently) and adds the
    under-used VAT-lifecycle analytics flagged in docs/DATA_ARCHITECTURE.md:
      * #2 cycle-time + payout forecasting — median submitted->paid days per country
        and overall; aging of open receivables by EUR and count;
      * #9 realization rate — paid_amount / vat_eur per refund country (which
        jurisdictions haircut claims).

    ROUTE-AWARE (mirrors the Recovery page): the per-claim figures come from
    settlement(payout_to, vat_eur, fee_eur) — the SAME helper the recovery page uses —
    so the two economically different cash flows are never conflated:
      * refund_receivable_eur = vat_eur — the refund owed BY THE STATE, route-independent
        (this is what's aged until the state pays the claim);
      * fee_eur                = the frozen service fee (agency receivable);
      * net_to_customer        = settlement()'s net (0 on the DEFAULT 'customer' route
        where the customer collects the full refund and we invoice the fee separately;
        vat − fee on the 'us' deduct route where we remit the net);
      * route                  = payout_to ('customer' default, or 'us').

    EUR figures via money.f2; the realization ratio is a fraction (format as % at
    display). Never invents a figure — fee_eur is the frozen value (falls back to
    compute_fee on the frozen/derived rate for legacy rows predating fee-freezing,
    identical to the recovery page).

    Returns a single dict:
      rows           per-claim receivable rows (see below)
      cycle_time     {"overall": median days|None, "by_country": {ctry: median}}
      aging          {"by_band": {band: {"eur":.., "count":..}}, "total_eur":.., "total_count":..}
      realization    {ctry: {"claimed":.., "paid":.., "rate": fraction|None}, ...} + "overall"
      forecast       open expected-cash view — the refund receivable owed by the state
                     (vat) and the agency fee receivable, kept separate (see below)
    """
    import datetime
    con = connect()
    rows = con.execute("""SELECT entity, refund_country, ref_period, vat_eur, status,
        status_code, submitted_date, approved_date, paid_date, paid_amount,
        fee_eur, fee_pct, fee_min, payout_to
        FROM vat_applications WHERE status IN ('submitted','approved','paid')
        AND (? IS NULL OR ref_period LIKE ?) ORDER BY refund_country, submitted_date""",
        (year, f"{year}-%" if year else None)).fetchall()
    today = datetime.date.today()
    out = []
    # cycle-time + realization accumulators, keyed by refund country
    cycle_by_ctry = collections.defaultdict(list)   # paid claims: submitted->paid days
    cycle_all = []
    realiz = collections.defaultdict(lambda: {"claimed": 0.0, "paid": 0.0})
    aging = {b: {"eur": 0.0, "count": 0} for b in AGING_BANDS}
    for r in rows:
        vat = r["vat_eur"] or 0
        # Frozen fee_eur is canonical; legacy rows predating fee-freezing derive the
        # fee from the customer's rate (same fallback the recovery page uses) so we
        # never silently treat the fee as zero.
        if r["fee_eur"] is None:
            fpct, fmin = customer_master.fee_for(r["entity"], r["refund_country"])
            fee, _b = customer_master.compute_fee(vat, fpct, fmin)
        else:
            fee = r["fee_eur"]
        # Route-aware settlement (same helper as the recovery page). The REFUND
        # RECEIVABLE owed by the state is vat (route-independent — what's aged until the
        # state pays); net_to_customer / fee_receivable depend on payout_to.
        route = r["payout_to"] or "customer"
        st = settlement(route, vat, fee)
        refund_receivable = st["refund"]            # == vat, route-independent
        net_to_customer = st["net_to_customer"]     # 0 on 'customer', vat-fee on 'us'
        # AGING of OPEN receivables (submitted/approved, not yet paid) by days since
        # submission — on the REFUND RECEIVABLE (the cash owed by the state).
        age = ""
        if r["status"] in ("submitted", "approved") and r["submitted_date"]:
            try:
                age = (today - datetime.date.fromisoformat(r["submitted_date"])).days
            except ValueError as e:
                log.debug("vat_refund: unparseable submitted_date %r: %s", r["submitted_date"], e)
        band = _aging_band(age) if r["status"] in ("submitted", "approved") else ""
        if band:
            aging[band]["eur"] = money.f2(aging[band]["eur"] + refund_receivable)
            aging[band]["count"] += 1
        # cycle time + realization on PAID claims only
        if r["status"] == "paid":
            if r["submitted_date"] and r["paid_date"]:
                try:
                    d = (datetime.date.fromisoformat(r["paid_date"])
                         - datetime.date.fromisoformat(r["submitted_date"])).days
                    cycle_by_ctry[r["refund_country"]].append(d)
                    cycle_all.append(d)
                except ValueError as e:
                    log.debug("vat_refund: unparseable cycle date(s) sub=%r paid=%r: %s",
                              r["submitted_date"], r["paid_date"], e)
            if vat:
                realiz_ct = realiz[r["refund_country"]]
                realiz_ct["claimed"] = money.f2(realiz_ct["claimed"] + vat)
                realiz_ct["paid"] = money.f2(realiz_ct["paid"] + (r["paid_amount"] or 0))
        code = r["status_code"] or {"submitted": "2", "approved": "3", "paid": "3A"}[r["status"]]
        out.append(dict(entity=r["entity"], country=r["refund_country"], period=r["ref_period"],
                        status=r["status"], status_code=code,
                        status_label=STATUS_LABELS.get(code, code),
                        vat_eur=vat, fee_eur=fee, route=route,
                        refund_receivable_eur=refund_receivable,
                        net_to_customer_eur=net_to_customer,
                        paid_amount=r["paid_amount"],
                        submitted=r["submitted_date"], approved=r["approved_date"],
                        paid=r["paid_date"], age_days=age, aging_band=band))
    con.close()
    # cycle-time medians
    cycle_time = {"overall": _median(cycle_all),
                  "by_country": {c: _median(v) for c, v in sorted(cycle_by_ctry.items())}}
    # realization rate per country (paid/claimed) + overall
    realization = {}
    tot_claimed = tot_paid = 0.0
    for c, v in sorted(realiz.items()):
        rate = (v["paid"] / v["claimed"]) if v["claimed"] else None
        realization[c] = {"claimed": money.f2(v["claimed"]), "paid": money.f2(v["paid"]),
                          "rate": rate}
        tot_claimed = money.f2(tot_claimed + v["claimed"])
        tot_paid = money.f2(tot_paid + v["paid"])
    realization["overall"] = {"claimed": tot_claimed, "paid": tot_paid,
                              "rate": (tot_paid / tot_claimed) if tot_claimed else None}
    # CASH FORECAST (route-aware, two SEPARATE flows — never summed across routes):
    #   * REFUND RECEIVABLE = the vat of open (unpaid submitted/approved) claims — the
    #     cash owed BY THE STATE. Route-independent: on BOTH routes the state pays this
    #     amount; on 'customer' it goes to the customer, on 'us' it comes to us. This is
    #     what's aged and realization-weighted (each open claim scaled by its refund
    #     country's historical paid/claimed rate; no history -> 1.0, no haircut yet).
    #   * AGENCY FEE RECEIVABLE = the frozen fee of open claims — the agency's own
    #     receivable, invoiced separately on the default 'customer' route and deducted on
    #     the 'us' route. Kept apart so the two economically different cash flows are not
    #     conflated.
    open_rows = [o for o in out if o["status"] in ("submitted", "approved")]
    open_refund = money.fsum(o["refund_receivable_eur"] for o in open_rows)
    open_fee = money.fsum(o["fee_eur"] for o in open_rows)
    weighted = 0.0
    for o in open_rows:
        cr = realization.get(o["country"], {})
        w = cr["rate"] if cr.get("rate") is not None else 1.0
        weighted = money.f2(weighted + o["refund_receivable_eur"] * w)
    aging_total_eur = money.fsum(aging[b]["eur"] for b in AGING_BANDS)
    aging_total_count = sum(aging[b]["count"] for b in AGING_BANDS)
    forecast = {"open_count": len(open_rows),
                "open_refund_receivable_eur": open_refund,
                "open_fee_receivable_eur": open_fee,
                "open_weighted_refund_eur": weighted,
                "aging_by_band": {b: aging[b]["eur"] for b in AGING_BANDS}}
    return {"rows": out, "cycle_time": cycle_time,
            "aging": {"by_band": aging, "total_eur": aging_total_eur,
                      "total_count": aging_total_count},
            "realization": realization, "forecast": forecast}
