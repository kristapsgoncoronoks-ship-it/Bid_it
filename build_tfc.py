import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

WORKDIR = os.path.dirname(os.path.abspath(__file__))

# (card, date, time, receipt, plate, product, location, qty, loc_price, disc, net_price, amount_net)
# disc=None & net_price=None where invoice shows "-" (no discount; amount = qty x loc_price)
T = [
("108019237","27-05-26","11:42",110795,"NCK 218","Diesel","Romac Houdeng",456.00,1.5944,None,None,727.05),
("108019237","27-05-26","11:50",110817,"NCK 218","Diesel","Romac Houdeng",51.00,1.5944,None,None,81.32),
("108019237","27-05-26","11:54",110829,"NCK 218","Diesel","Romac Houdeng",24.00,1.5944,None,None,38.27),
("108019237","29-05-26","15:10",117360,"NCK 218","Diesel","Tfc Dcb Energy Hub Meer",271.00,1.6843,0.2050,1.4793,400.90),
("108019237","29-05-26","15:19",117410,"NCK 218","Diesel","Tfc Dcb Energy Hub Meer",26.00,1.6843,0.2050,1.4793,38.47),
("108019237","29-05-26","15:21",117425,"NCK 218","Diesel","Tfc Dcb Energy Hub Meer",137.01,1.6843,0.2050,1.4793,202.68),
("108019237","27-05-26","11:47",110809,"NCK 218","AdBlue","Romac Houdeng",30.00,0.6750,None,None,20.25),
("108019237","29-05-26","15:15",117389,"NCK 218","AdBlue","Tfc Dcb Energy Hub Meer",48.11,0.6000,0.0,0.6000,28.87),
("145993502","18-05-26","13:38",84860,"MSZ552","Diesel","Tfc Dcb Energy Hub Meer",445.00,1.7488,0.2050,1.5438,687.00),
("145993502","18-05-26","13:51",84902,"MSZ552","Diesel","Tfc Dcb Energy Hub Meer",55.00,1.7488,0.2050,1.5438,84.91),
("145993502","18-05-26","13:46",84891,"MSZ552","AdBlue","Tfc Dcb Energy Hub Meer",44.00,0.6000,0.0,0.6000,26.40),
("168404872","21-05-26","13:25",95615,"AII 954","Diesel","Tfc Dcb Energy Hub Meer 2",495.00,1.8116,0.2050,1.6066,795.27),
("168404872","21-05-26","13:32",95632,"AII 954","Diesel","Tfc Dcb Energy Hub Meer 2",260.02,1.8116,0.2050,1.6066,417.75),
("168404872","21-05-26","13:39",95650,"AII 954","Diesel","Tfc Dcb Energy Hub Meer 2",129.00,1.8116,0.2050,1.6066,207.26),
("168404872","21-05-26","13:42",95664,"AII 954","AdBlue","Tfc Dcb Energy Hub Meer 2",81.02,0.6000,0.0,0.6000,48.62),
("280623763","18-05-26","08:25",84064,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer",602.43,1.7488,0.2050,1.5438,930.04),
("280623763","18-05-26","08:41",84092,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer",303.85,1.7488,0.2050,1.5438,469.09),
("280623763","18-05-26","08:45",84104,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer",16.77,1.7488,0.2050,1.5438,25.89),
("280623763","22-05-26","14:19",98837,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer 2",301.16,1.7694,0.2050,1.5644,471.14),
("280623763","22-05-26","14:24",98855,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer 2",157.64,1.7694,0.2050,1.5644,246.62),
("280623763","22-05-26","14:29",98870,"NCK 230","Diesel","Tfc Dcb Energy Hub Meer 2",16.23,1.7694,0.2050,1.5644,25.40),
("280623763","18-05-26","08:33",84077,"NCK 230","AdBlue","Tfc Dcb Energy Hub Meer",79.77,0.6000,0.0,0.6000,47.87),
("280623763","22-05-26","14:32",98877,"NCK 230","AdBlue","Tfc Dcb Energy Hub Meer 2",30.39,0.6000,0.0,0.6000,18.24),
("332206574","29-05-26","14:00",251637,"MTZ159","Diesel","Dcb Energy Hub Maasmechelen",620.35,1.6843,0.1900,1.4943,926.99),
("332206574","29-05-26","14:28",251689,"MTZ159","Diesel","Dcb Energy Hub Maasmechelen",154.53,1.6843,0.1900,1.4943,230.92),
("332206574","29-05-26","14:12",251662,"MTZ159","AdBlue","Dcb Energy Hub Maasmechelen",58.93,0.6750,None,None,39.78),
("368425169","23-05-26","03:55",100414,"MSZ543","Diesel","Tfc Dcb Energy Hub Meer 2",555.00,1.7694,0.2050,1.5644,868.25),
("368425169","23-05-26","04:03",100424,"MSZ543","Diesel","Tfc Dcb Energy Hub Meer 2",300.00,1.7694,0.2050,1.5644,469.32),
("368425169","23-05-26","04:13",100438,"MSZ543","Diesel","Tfc Dcb Energy Hub Meer 2",160.00,1.7694,0.2050,1.5644,250.31),
("368425169","23-05-26","04:08",100431,"MSZ543","AdBlue","Tfc Dcb Energy Hub Meer 2",57.00,0.6000,0.0,0.6000,34.20),
("390924019","26-05-26","12:50",107916,"NNJ 874","Diesel","Tfc Dcb Energy Hub Meer 2",580.00,1.7694,0.2050,1.5644,907.36),
("390924019","26-05-26","12:58",107938,"NNJ 874","Diesel","Tfc Dcb Energy Hub Meer 2",310.00,1.7694,0.2050,1.5644,484.97),
("390924019","26-05-26","13:03",107952,"NNJ 874","Diesel","Tfc Dcb Energy Hub Meer 2",120.00,1.7694,0.2050,1.5644,187.73),
("390924019","26-05-26","13:07",107966,"NNJ 874","AdBlue","Tfc Dcb Energy Hub Meer 2",56.71,0.6000,0.0,0.6000,34.03),
("420982795","22-05-26","12:39",98533,"MND 992","Diesel","Tfc Dcb Energy Hub Meer",120.00,1.7694,0.2050,1.5644,187.73),
("420982795","22-05-26","12:45",98555,"MND 992","Diesel","Tfc Dcb Energy Hub Meer",400.00,1.7694,0.2050,1.5644,625.76),
("420982795","28-05-26","01:03",248135,"MND 992","Diesel","Romac Houdeng",205.10,1.5944,None,None,327.02),
("420982795","28-05-26","01:08",248142,"MND 992","Diesel","Romac Houdeng",340.00,1.5944,None,None,542.10),
("420982795","28-05-26","01:14",248149,"MND 992","Diesel","Romac Houdeng",550.00,1.5944,None,None,876.92),
("420982795","28-05-26","01:21",248157,"MND 992","AdBlue","Romac Houdeng",78.00,0.6750,None,None,52.65),
("431120840","22-05-26","13:14",98635,"MFF 264","Diesel","Tfc Dcb Energy Hub Meer",354.49,1.7694,0.2050,1.5644,554.57),
("431120840","22-05-26","13:20",98648,"MFF 264","Diesel","Tfc Dcb Energy Hub Meer",384.26,1.7694,0.2050,1.5644,601.14),
("431120840","22-05-26","13:28",98663,"MFF 264","AdBlue","Tfc Dcb Energy Hub Meer",26.75,0.6000,0.0,0.6000,16.05),
("492309405","28-05-26","15:46",114524,"NCK 206","Diesel","Tfc Dcb Energy Hub Meer 2",235.06,1.7694,0.2050,1.5644,367.73),
("492309405","28-05-26","15:51",114532,"NCK 206","Diesel","Tfc Dcb Energy Hub Meer 2",258.85,1.7694,0.2050,1.5644,404.95),
("492309405","28-05-26","15:55",114543,"NCK 206","Diesel","Tfc Dcb Energy Hub Meer 2",433.73,1.7694,0.2050,1.5644,678.53),
("492309405","28-05-26","16:02",114563,"NCK 206","AdBlue","Tfc Dcb Energy Hub Meer 2",48.40,0.6000,0.0,0.6000,29.04),
("496838029","20-05-26","15:30",92552,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",467.67,1.8116,0.2050,1.6066,751.36),
("496838029","20-05-26","15:42",92578,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",250.40,1.8116,0.2050,1.6066,402.30),
("496838029","20-05-26","15:46",92584,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",99.57,1.8116,0.2050,1.6066,159.97),
("496838029","30-05-26","09:35",119873,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",563.10,1.6843,0.2050,1.4793,833.00),
("496838029","30-05-26","09:47",119911,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",288.72,1.6843,0.2050,1.4793,427.11),
("496838029","30-05-26","09:52",119930,"MTG172","Diesel","Tfc Dcb Energy Hub Meer 2",188.34,1.6843,0.2050,1.4793,278.62),
("496838029","20-05-26","15:37",92561,"MTG172","AdBlue","Tfc Dcb Energy Hub Meer 2",49.72,0.6000,0.0,0.6000,29.84),
("496838029","30-05-26","09:42",119895,"MTG172","AdBlue","Tfc Dcb Energy Hub Meer 2",55.80,0.6000,0.0,0.6000,33.48),
("575148137","26-05-26","12:28",107866,"MTG196","Diesel","Tfc Dcb Energy Hub Meer 2",360.00,1.7694,0.2050,1.5644,563.19),
("575148137","26-05-26","12:35",107881,"MTG196","Diesel","Tfc Dcb Energy Hub Meer 2",185.00,1.7694,0.2050,1.5644,289.42),
("575148137","26-05-26","12:40",107893,"MTG196","Diesel","Tfc Dcb Energy Hub Meer 2",51.00,1.7694,0.2050,1.5644,79.79),
("575148137","26-05-26","12:45",107906,"MTG196","AdBlue","Tfc Dcb Energy Hub Meer 2",33.89,0.6000,0.0,0.6000,20.34),
("647718447","23-05-26","15:35",101833,"NRH 225","Diesel","Tfc Dcb Energy Hub Meer 2",34.65,1.7694,0.2050,1.5644,54.21),
("649128539","30-05-26","06:01",253016,"MSZ557","Diesel","De Warande Oostkamp",205.01,1.5093,None,None,309.43),
("649128539","30-05-26","06:09",253025,"MSZ557","Diesel","De Warande Oostkamp",430.00,1.5093,None,None,649.00),
("649128539","30-05-26","06:16",253033,"MSZ557","AdBlue","De Warande Oostkamp",44.00,0.6750,None,None,29.70),
("707758061","30-05-26","07:20",119469,"MOS360","Diesel","Tfc Dcb Energy Hub Meer 2",590.00,1.6843,0.2050,1.4793,872.79),
("707758061","30-05-26","07:29",119491,"MOS360","Diesel","Tfc Dcb Energy Hub Meer 2",315.01,1.6843,0.2050,1.4793,466.00),
("707758061","30-05-26","07:37",119511,"MOS360","Diesel","Tfc Dcb Energy Hub Meer 2",118.01,1.6843,0.2050,1.4793,174.58),
("707758061","30-05-26","07:41",119528,"MOS360","AdBlue","Tfc Dcb Energy Hub Meer 2",78.08,0.6000,0.0,0.6000,46.85),
("771666683","29-05-26","16:50",117693,"AIE 819","Diesel","Tfc Dcb Energy Hub Meer",528.26,1.6843,0.2050,1.4793,781.46),
("771666683","29-05-26","17:04",117730,"AIE 819","Diesel","Tfc Dcb Energy Hub Meer",268.39,1.6843,0.2050,1.4793,397.03),
("771666683","29-05-26","17:09",117745,"AIE 819","Diesel","Tfc Dcb Energy Hub Meer",52.45,1.6843,0.2050,1.4793,77.59),
("771666683","29-05-26","17:12",117753,"AIE 819","Diesel","Tfc Dcb Energy Hub Meer",3.88,1.6843,0.2050,1.4793,5.74),
("771666683","29-05-26","16:59",117715,"AIE 819","AdBlue","Tfc Dcb Energy Hub Meer",76.21,0.6000,0.0,0.6000,45.73),
("775094081","18-05-26","10:15",227217,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",625.67,1.7488,0.1900,1.5588,975.30),
("775094081","18-05-26","10:29",227241,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",323.08,1.7488,0.1900,1.5588,503.62),
("775094081","18-05-26","10:35",227252,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",201.00,1.7488,0.1900,1.5588,313.32),
("775094081","27-05-26","10:25",246639,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",403.72,1.7694,0.1900,1.5794,637.64),
("775094081","27-05-26","10:32",246656,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",217.40,1.7694,0.1900,1.5794,343.37),
("775094081","27-05-26","10:36",246671,"MSB265","Diesel","Dcb Energy Hub Maasmechelen",119.23,1.7694,0.1900,1.5794,188.32),
("775094081","18-05-26","10:24",227231,"MSB265","AdBlue","Dcb Energy Hub Maasmechelen",45.03,0.6750,None,None,30.40),
("775094081","27-05-26","10:40",246679,"MSB265","AdBlue","Dcb Energy Hub Maasmechelen",22.45,0.6750,None,None,15.16),
("823016623","18-05-26","14:57",85096,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer 2",675.00,1.7488,0.2050,1.5438,1042.07),
("823016623","18-05-26","15:07",85127,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer 2",293.00,1.7488,0.2050,1.5438,452.34),
("823016623","18-05-26","15:12",85146,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer 2",85.00,1.7488,0.2050,1.5438,131.23),
("823016623","27-05-26","15:28",111452,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer",661.00,1.7694,0.2050,1.5644,1034.07),
("823016623","27-05-26","15:36",111473,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer",290.00,1.7694,0.2050,1.5644,453.68),
("823016623","27-05-26","15:42",111486,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer",151.01,1.7694,0.2050,1.5644,236.25),
("823016623","29-05-26","15:17",117399,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer",100.00,1.6843,0.2050,1.4793,147.93),
("823016623","29-05-26","15:22",117427,"LRN 472","Diesel","Tfc Dcb Energy Hub Meer",52.00,1.6843,0.2050,1.4793,76.93),
("823016623","18-05-26","14:50",85075,"LRN 472","AdBlue","Tfc Dcb Energy Hub Meer 2",49.00,0.6000,0.0,0.6000,29.40),
("823016623","27-05-26","15:23",111434,"LRN 472","AdBlue","Tfc Dcb Energy Hub Meer",48.86,0.6000,0.0,0.6000,29.32),
("823016623","29-05-26","15:14",117373,"LRN 472","AdBlue","Tfc Dcb Energy Hub Meer",8.00,0.6000,0.0,0.6000,4.80),
("856166744","26-05-26","13:31",108012,"MFF 265","Diesel","Romac Houdeng",250.00,1.5944,None,None,398.60),
("856166744","26-05-26","13:38",108025,"MFF 265","Diesel","Romac Houdeng",370.00,1.5944,None,None,589.93),
("856166744","26-05-26","13:48",108054,"MFF 265","Diesel","Romac Houdeng",45.81,1.5944,None,None,73.04),
("856166744","26-05-26","13:44",108042,"MFF 265","AdBlue","Romac Houdeng",19.92,0.6750,None,None,13.45),
("889167685","16-05-26","07:11",79582,"MDU 054","Diesel","Tfc Dcb Energy Hub Meer 2",400.00,1.7488,0.2050,1.5438,617.52),
("889167685","26-05-26","15:29",108328,"MDU 054","Diesel","Romac Houdeng",525.00,1.5944,None,None,837.06),
("889167685","26-05-26","15:57",108402,"MDU 054","Diesel","Romac Houdeng",179.00,1.5944,None,None,285.40),
("898603163","20-05-26","19:49",93346,"NCK 186","Diesel","Texaco Kruishoutem",632.02,1.6366,None,None,1034.37),
("898603163","20-05-26","20:05",93402,"NCK 186","Diesel","Texaco Kruishoutem",108.00,1.6366,None,None,176.76),
("898603163","20-05-26","20:01",93388,"NCK 186","AdBlue","Texaco Kruishoutem",31.49,0.6750,None,None,21.26),
("924148677","20-05-26","18:12",93019,"LPV 991","Diesel","Tfc Dcb Energy Hub Meer 2",400.00,1.8116,0.2050,1.6066,642.64),
("924148677","28-05-26","06:12",113223,"LPV 991","Diesel","Tfc Dcb Energy Hub Meer",720.02,1.7694,0.2050,1.5644,1126.40),
("924148677","28-05-26","06:21",113251,"LPV 991","Diesel","Tfc Dcb Energy Hub Meer",150.00,1.7694,0.2050,1.5644,234.66),
("924148677","28-05-26","06:25",113263,"LPV 991","Diesel","Tfc Dcb Energy Hub Meer",295.96,1.7694,0.2050,1.5644,463.00),
("936582681","23-05-26","07:54",100933,"LDI 991","Diesel","Tfc Dcb Energy Hub Meer 2",550.66,1.7694,0.2050,1.5644,861.46),
("936582681","23-05-26","08:01",100956,"LDI 991","Diesel","Tfc Dcb Energy Hub Meer 2",300.06,1.7694,0.2050,1.5644,469.42),
("936582681","23-05-26","08:07",100971,"LDI 991","AdBlue","Tfc Dcb Energy Hub Meer 2",49.40,0.6000,0.0,0.6000,29.64),
("952457222","28-05-26","13:15",249366,"NCS 504","Diesel","Dcb Energy Hub Maasmechelen",543.00,1.7694,0.1900,1.5794,857.62),
("952457222","28-05-26","13:26",249389,"NCS 504","Diesel","Dcb Energy Hub Maasmechelen",129.64,1.7694,0.1900,1.5794,204.76),
("952457222","28-05-26","13:31",249399,"NCS 504","AdBlue","Dcb Energy Hub Maasmechelen",45.28,0.6750,None,None,30.57),
("959770118","29-05-26","15:11",117362,"MFI170","Diesel","Tfc Dcb Energy Hub Meer",567.19,1.6843,0.2050,1.4793,839.05),
("959770118","29-05-26","15:20",117417,"MFI170","Diesel","Tfc Dcb Energy Hub Meer",289.94,1.6843,0.2050,1.4793,428.91),
("959770118","29-05-26","15:26",117450,"MFI170","Diesel","Tfc Dcb Energy Hub Meer",152.01,1.6843,0.2050,1.4793,224.87),
("959770118","29-05-26","15:31",117468,"MFI170","AdBlue","Tfc Dcb Energy Hub Meer",44.89,0.6000,0.0,0.6000,26.94),
("991011529","23-05-26","23:10",102383,"MSZ547","Diesel","Tfc Dcb Energy Hub Meer 2",215.65,1.7694,0.2050,1.5644,337.37),
("991011529","23-05-26","23:17",102391,"MSZ547","Diesel","Tfc Dcb Energy Hub Meer 2",36.01,1.7694,0.2050,1.5644,56.34),
("991011529","23-05-26","23:19",102397,"MSZ547","Diesel","Tfc Dcb Energy Hub Meer 2",600.13,1.7694,0.2050,1.5644,938.85),
("991011529","23-05-26","23:33",102413,"MSZ547","Diesel","Tfc Dcb Energy Hub Meer 2",268.91,1.7694,0.2050,1.5644,420.69),
("991011529","23-05-26","23:25",102405,"MSZ547","AdBlue","Tfc Dcb Energy Hub Meer 2",88.20,0.6000,0.0,0.6000,52.92),
]
card_totals = {"108019237":1537.81,"145993502":798.31,"168404872":1468.90,"280623763":2234.29,
"332206574":1197.69,"368425169":1622.08,"390924019":1614.09,"420982795":2612.18,
"431120840":1171.76,"492309405":1480.25,"496838029":2915.68,"575148137":952.74,
"647718447":54.21,"649128539":988.13,"707758061":1560.22,"771666683":1307.55,
"775094081":3007.13,"823016623":3638.02,"856166744":1075.02,"889167685":1739.98,
"898603163":1232.39,"924148677":2466.70,"936582681":1360.52,"952457222":1092.95,
"959770118":1519.77,"991011529":1806.17}

wb = Workbook()
ws = wb.active
ws.title = "Transactions"
hdr = ["Card","Date","Time","Receipt","Plate","Product","Location","Volume L",
       "Location price EUR/L","Discount EUR/L","Net price EUR/L (doc)","Net amount EUR (doc)",
       "Eff. price used","Check1: price math","Calc amount","Diff","Check2: amount"]
ws.append(hdr)
fill = PatternFill("solid", start_color="C2410C")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    c.fill = fill; c.alignment = Alignment(horizontal="center", wrap_text=True)
norm = Font(name="Arial", size=10)
r = 2
for card, d, t, rec, plate, prod, loc, vol, lp, disc, np, amt in T:
    ws.append([card, d, t, rec, plate, prod, loc, vol, lp, disc, np, amt,
        f"=IF(K{r}<>\"\",K{r},I{r})",
        (f'=IF(K{r}="","n/a (no discount)",IF(ABS(I{r}-J{r}-K{r})<=0.0001,"OK","CHECK"))'),
        f"=ROUND(M{r}*H{r},2)", f"=O{r}-L{r}", f'=IF(ABS(P{r})<=0.011,"OK","CHECK")'])
    if disc is None:
        ws[f"J{r}"] = ""; ws[f"K{r}"] = ""
    r += 1
last = r - 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = norm
for col, f in (("H","#,##0.00"),("I","0.0000"),("J","0.0000"),("K","0.0000"),
               ("L","#,##0.00"),("M","0.0000"),("O","#,##0.00"),("P","0.000")):
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=ord(col)-64, max_col=ord(col)-64):
        row[0].number_format = f
for col, w in zip("ABCDEFGHIJKLMNOPQ",[12,10,7,9,9,8,26,10,11,10,11,12,10,13,11,8,9]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:Q{last}"

# Card totals check
ws2 = wb.create_sheet("Card totals check")
ws2.append(["Card","Stated total EUR (net)","Calc sum of lines","Diff","Check"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10); c.fill = fill
r2 = 2
for card, tot in card_totals.items():
    ws2.append([card, tot,
        f'=ROUND(SUMIF(Transactions!$A$2:$A${last},A{r2}&"",Transactions!$L$2:$L${last}),2)',
        f"=C{r2}-B{r2}", f'=IF(ABS(D{r2})<=0.01,"OK","CHECK")'])
    r2 += 1
ws2.append(["TOTAL", f"=SUM(B2:B{r2-1})", f"=SUM(C2:C{r2-1})", f"=C{r2}-B{r2}",
            f'=IF(ABS(D{r2})<=0.01,"OK","CHECK")'])
last2 = r2
for row in ws2.iter_rows(min_row=2, max_row=last2):
    for c in row: c.font = norm
    for c in row[1:4]: c.number_format = "#,##0.00"
for c in ws2[last2]: c.font = Font(bold=True, name="Arial", size=10)
for col, w in zip("ABCDE",[13,19,17,10,8]):
    ws2.column_dimensions[col].width = w

# Invoice totals + price analysis
ws3 = wb.create_sheet("Invoice totals check")
ws3.append(["Check item","Invoice states","Calculated","Diff","Check"])
for c in ws3[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10); c.fill = fill
items = [
 ("Total excl. VAT EUR", 42454.54, f"=ROUND(SUM(Transactions!$L$2:$L${last}),2)", 0.01),
 ("VAT 21% EUR", 8915.45, "=ROUND(C2*0.21,2)", 0.02),
 ("Invoice total EUR", 51369.99, "=C2+C3", 0.03),
 ("Diesel litres", None, f'=ROUND(SUMIF(Transactions!$F$2:$F${last},"Diesel",Transactions!$H$2:$H${last}),2)', None),
 ("Diesel net EUR", None, f'=ROUND(SUMIF(Transactions!$F$2:$F${last},"Diesel",Transactions!$L$2:$L${last}),2)', None),
 ("AdBlue litres", None, f'=ROUND(SUMIF(Transactions!$F$2:$F${last},"AdBlue",Transactions!$H$2:$H${last}),2)', None),
 ("AdBlue net EUR", None, f'=ROUND(SUMIF(Transactions!$F$2:$F${last},"AdBlue",Transactions!$L$2:$L${last}),2)', None),
]
r3 = 2
for name, doc, calc, tol in items:
    if doc is not None:
        ws3.append([name, doc, calc, f"=C{r3}-B{r3}", f'=IF(ABS(D{r3})<={tol},"OK","CHECK")'])
    else:
        ws3.append([name, "(not stated)", calc, "", ""])
    r3 += 1
for row in ws3.iter_rows(min_row=2, max_row=r3-1):
    for c in row: c.font = norm
    for c in row[1:4]:
        if not isinstance(c.value, str) or c.value.startswith("="): c.number_format = "#,##0.00"
for col, w in zip("ABCDE",[24,15,15,10,8]):
    ws3.column_dimensions[col].width = w

ws4 = wb.create_sheet("Station price analysis")
ws4.append(["Station","Diesel litres","Diesel net EUR","Avg eff. net EUR/L","Discount shown?"])
for c in ws4[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    c.fill = fill; c.alignment = Alignment(wrap_text=True, horizontal="center")
stations = [("Tfc Dcb Energy Hub Meer","yes (-0.2050)"),("Tfc Dcb Energy Hub Meer 2","yes (-0.2050)"),
            ("Dcb Energy Hub Maasmechelen","yes (-0.1900)"),("Romac Houdeng","no"),
            ("De Warande Oostkamp","no"),("Texaco Kruishoutem","no")]
r4 = 2
for stn, disc in stations:
    ws4.append([stn,
      f'=ROUND(SUMPRODUCT((Transactions!$G$2:$G${last}=A{r4})*(Transactions!$F$2:$F${last}="Diesel")*Transactions!$H$2:$H${last}),2)',
      f'=ROUND(SUMPRODUCT((Transactions!$G$2:$G${last}=A{r4})*(Transactions!$F$2:$F${last}="Diesel")*Transactions!$L$2:$L${last}),2)',
      f"=C{r4}/B{r4}", disc])
    r4 += 1
ws4.append(["TOTAL", f"=SUM(B2:B{r4-1})", f"=SUM(C2:C{r4-1})", f"=C{r4}/B{r4}", ""])
for row in ws4.iter_rows(min_row=2, max_row=r4):
    for c in row: c.font = norm
    row[1].number_format = "#,##0.00"; row[2].number_format = "#,##0.00"; row[3].number_format = "0.0000"
for c in ws4[r4]: c.font = Font(bold=True, name="Arial", size=10)
for col, w in zip("ABCDE",[28,13,14,15,15]):
    ws4.column_dimensions[col].width = w

wb.save(os.path.join(WORKDIR, "TFC_26056012270_transactions.xlsx"))
print("lines:", last-1, "cards:", len(card_totals))
