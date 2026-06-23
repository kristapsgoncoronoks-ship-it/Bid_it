"""Supplier overcharge CLAIM-BACK (overcharge.py): the lifecycle store + evidence packet over
the contract-audit overcharges. detected → packaged → claimed → recovered."""
import io

import pytest

import overcharge as OC


@pytest.fixture(autouse=True)
def _tmp_db(tmp_path, monkeypatch):
    monkeypatch.setattr(OC, "DB", str(tmp_path / "overcharge.db"))
    monkeypatch.setattr(OC, "_READY", set())


_FLAGS = [
    {"supplier": "DKV", "country": "Belgium", "station": "BE12", "period": "2026-05",
     "product": "Diesel", "litres": 1000.0, "issue": "short discount",
     "expected": 0.05, "actual": 0.02, "recover_eur": 30.0, "note": "contract -5c/L"},
    {"supplier": "DKV", "country": "France", "station": "FR9", "period": "2026-05",
     "product": "Diesel", "litres": 500.0, "issue": "over ceiling",
     "expected": 1.50, "actual": 1.60, "recover_eur": 50.0, "note": ""},
    {"supplier": "SHELL", "country": "Spain", "station": "ES1", "period": "2026-05",
     "product": "Diesel", "litres": 200.0, "issue": "short discount",
     "expected": 0.04, "actual": 0.01, "recover_eur": 6.0, "note": ""},
]


def _stub_audit(monkeypatch, flags=_FLAGS):
    import contract_audit
    monkeypatch.setattr(contract_audit, "audit",
                        lambda period=None, **k: (flags, {"total_recover": sum(f["recover_eur"] for f in flags),
                                                          "flags": len(flags), "by_supplier": {}, "rules": 1}))


def test_lifecycle_and_recovered_total():
    OC.ensure_claim("DKV", "2026-05", 80.0, actor="t")
    ok, _ = OC.set_status("DKV", "2026-05", "packaged", actor="t")
    assert ok
    OC.set_status("DKV", "2026-05", "claimed", actor="t")
    ok, err = OC.record_recovery("DKV", "2026-05", 64.0, actor="t")
    assert ok, err
    assert OC.recovered_total() == 64.0
    # a bad status / amount is rejected
    assert OC.set_status("DKV", "2026-05", "nonsense", actor="t")[0] is False
    assert OC.record_recovery("DKV", "2026-05", "abc", actor="t")[0] is False


def test_overview_aggregates_per_supplier_and_joins_status(monkeypatch):
    _stub_audit(monkeypatch)
    OC.set_status("DKV", "2026-05", "claimed", actor="t", detected_eur=80.0)
    ov = {(r["supplier"], r["period"]): r for r in OC.overview()}
    dkv = ov[("DKV", "2026-05")]
    assert dkv["detected_eur"] == 80.0 and dkv["flags"] == 2   # two DKV lines aggregated
    assert dkv["status"] == "claimed"
    assert ov[("SHELL", "2026-05")]["status"] == "detected"    # untouched -> detected


def test_recovered_total_only_counts_recovered(monkeypatch):
    _stub_audit(monkeypatch)
    OC.set_status("DKV", "2026-05", "claimed", actor="t", detected_eur=80.0)
    assert OC.recovered_total() == 0.0
    OC.record_recovery("SHELL", "2026-05", 6.0, actor="t", detected_eur=6.0)
    assert OC.recovered_total() == 6.0


def test_evidence_workbook_is_xlsx_with_lines(monkeypatch):
    _stub_audit(monkeypatch)
    data = OC.evidence_workbook("DKV", "2026-05")
    assert data and data[:2] == b"PK"                          # xlsx (zip) magic
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data)).active
    txt = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "DKV" in txt and "Total recoverable EUR" in txt and "short discount" in txt
    # a supplier with no lines packages nothing
    assert OC.evidence_workbook("NOBODY", "2026-05") is None


def test_overcharges_page_renders(client):
    r = client.get("/overcharges")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "Supplier overcharge claim-back" in body and "open claim-backs" in body
