"""Company expense / cost-allocation report (the Phase-2 "easy expense reports").

Covers:
- queries.q_expense: a hand-verifiable per-vehicle rollup (net/vat/gross add up, gross
  == net+vat, the effective NET €/L); the entity filter restricting all three rollups;
  the by_entity totals reconciling with the totals dict; and the never-raise contract
  (a bogus period / a connection with no table returns the shaped-empty dict).
- reports.expense_report_workbook: a valid styled workbook for the seeded period, with
  and without an entity filter, carrying the NET-EUR basis caveat on the title rows.
- /expenses + /export/expenses web smoke: the page renders for an authenticated user
  and the Excel downloads.

BASIS: NET EUR, final (rebates applied); VAT shown separately; gross = net + VAT.
Pure READ-ONLY over the engine-owned transactions.
"""
import sqlite3

import money
import queries
import reports


def _seed(rows):
    """rows: (period, entity, country, vehicle, product_group, qty, net_eur, vat_eur,
    net_eur_eff)."""
    con = sqlite3.connect(":memory:"); con.row_factory = sqlite3.Row
    con.execute("""CREATE TABLE transactions (period TEXT, entity TEXT, country TEXT,
                   vehicle TEXT, product_group TEXT, qty REAL, net_eur REAL,
                   vat_eur REAL, net_eur_eff REAL)""")
    con.executemany("INSERT INTO transactions VALUES (?,?,?,?,?,?,?,?,?)", rows)
    con.commit()
    return con


def test_expense_hand_verifiable_vehicle():
    # One entity, one vehicle, two fuellings: 100L net 130 vat 27.3 eff 128;
    # 50L net 70 vat 14.7 eff 69. Totals: net 200, vat 42, gross 242, litres 150.
    # NET €/L = (128+69)/150 = 197/150 = 1.31333...
    con = _seed([
        ("2026-05", "ENT", "Latvia", "V1", "Diesel", 100.0, 130.0, 27.3, 128.0),
        ("2026-05", "ENT", "Estonia", "V1", "Diesel", 50.0, 70.0, 14.7, 69.0),
    ])
    d = queries.q_expense(con, "2026-05")
    con.close()
    assert len(d["by_vehicle"]) == 1
    v = d["by_vehicle"][0]
    assert v["entity"] == "ENT" and v["vehicle"] == "V1"
    assert v["fuellings"] == 2
    assert v["litres"] == 150.0
    assert v["net_eur"] == 200.0
    assert v["vat_eur"] == 42.0
    assert v["gross_eur"] == 242.0
    # gross == net + vat exactly
    assert v["gross_eur"] == v["net_eur"] + v["vat_eur"]
    # two distinct countries on the vehicle
    assert v["n_countries"] == 2
    # effective NET €/L = SUM(net_eur_eff)/SUM(qty)
    assert round(v["net_eur_l"], 4) == round(197.0 / 150.0, 4)


def test_expense_gross_is_net_plus_vat_everywhere():
    con = _seed([
        ("2026-05", "A", "Latvia", "V1", "Diesel", 100.0, 130.0, 27.3, 128.0),
        ("2026-05", "B", "Estonia", "V2", "AdBlue", 20.0, 18.0, 3.78, 17.5),
    ])
    d = queries.q_expense(con, "2026-05")
    con.close()
    # gross = net + VAT, quantized HALF_UP at the boundary (money.f2) like the code.
    for r in d["by_entity"] + d["by_vehicle"] + d["by_product"]:
        assert r["gross_eur"] == money.f2(r["net_eur"] + r["vat_eur"]), r
    t = d["totals"]
    assert t["gross_eur"] == money.f2(t["net_eur"] + t["vat_eur"])


def test_expense_entity_filter_restricts_all_rollups():
    con = _seed([
        ("2026-05", "A", "Latvia", "V1", "Diesel", 100.0, 130.0, 27.3, 128.0),
        ("2026-05", "A", "Latvia", "V2", "AdBlue", 20.0, 18.0, 3.78, 17.5),
        ("2026-05", "B", "Estonia", "V9", "Diesel", 200.0, 260.0, 54.6, 256.0),
    ])
    d = queries.q_expense(con, "2026-05", entity="A")
    con.close()
    # only entity A appears, in every rollup
    assert [e["entity"] for e in d["by_entity"]] == ["A"]
    assert {v["entity"] for v in d["by_vehicle"]} == {"A"}
    assert {v["vehicle"] for v in d["by_vehicle"]} == {"V1", "V2"}
    # B's diesel must not leak into the product split or totals
    assert d["totals"]["net_eur"] == 148.0     # 130 + 18
    assert d["totals"]["litres"] == 120.0
    assert d["totals"]["fuellings"] == 2


def test_expense_by_entity_reconciles_with_totals():
    con = _seed([
        ("2026-05", "A", "Latvia", "V1", "Diesel", 100.0, 130.0, 27.3, 128.0),
        ("2026-05", "A", "Latvia", "V2", "AdBlue", 20.0, 18.0, 3.78, 17.5),
        ("2026-05", "B", "Estonia", "V9", "Diesel", 200.0, 260.55, 54.6, 256.0),
    ])
    d = queries.q_expense(con, "2026-05")
    con.close()
    t = d["totals"]
    # sum of per-entity figures reconciles with the period totals (to the cent)
    assert round(sum(e["net_eur"] for e in d["by_entity"]), 2) == round(t["net_eur"], 2)
    assert round(sum(e["vat_eur"] for e in d["by_entity"]), 2) == round(t["vat_eur"], 2)
    assert round(sum(e["gross_eur"] for e in d["by_entity"]), 2) == round(t["gross_eur"], 2)
    assert sum(e["fuellings"] for e in d["by_entity"]) == t["fuellings"]
    assert round(sum(e["litres"] for e in d["by_entity"]), 4) == round(t["litres"], 4)


def test_expense_by_entity_sorted_net_desc():
    con = _seed([
        ("2026-05", "SMALL", "Latvia", "V1", "Diesel", 10.0, 13.0, 2.7, 12.8),
        ("2026-05", "BIG", "Estonia", "V9", "Diesel", 200.0, 260.0, 54.6, 256.0),
    ])
    d = queries.q_expense(con, "2026-05")
    con.close()
    nets = [e["net_eur"] for e in d["by_entity"]]
    assert nets == sorted(nets, reverse=True)
    assert d["by_entity"][0]["entity"] == "BIG"


def test_expense_zero_litres_no_divide_error():
    # A vehicle with zero litres (e.g. a toll-only line) must not blow up net_eur_l.
    con = _seed([
        ("2026-05", "A", "Latvia", "V1", "Toll", 0.0, 12.0, 2.52, 12.0),
    ])
    d = queries.q_expense(con, "2026-05")
    con.close()
    assert d["by_vehicle"][0]["net_eur_l"] == 0.0


def test_expense_never_raises_on_bad_period():
    con = _seed([
        ("2026-05", "A", "Latvia", "V1", "Diesel", 100.0, 130.0, 27.3, 128.0),
    ])
    d = queries.q_expense(con, "1999-01")    # no such period
    con.close()
    assert d == {"by_entity": [], "by_vehicle": [], "by_product": [],
                 "totals": {"net_eur": 0.0, "vat_eur": 0.0, "gross_eur": 0.0,
                            "litres": 0.0, "fuellings": 0}}


def test_expense_never_raises_on_missing_table():
    con = sqlite3.connect(":memory:")        # no transactions table
    d = queries.q_expense(con, "2026-05")
    con.close()
    assert d["by_entity"] == [] and d["totals"]["fuellings"] == 0


def test_workbook_builds(tmp_path):
    from openpyxl import load_workbook
    path = reports.expense_report_workbook(
        "2026-05", path=str(tmp_path / "exp.xlsx"))
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames and "Vehicles" in wb.sheetnames
    # NET-EUR basis caveat printed on the title rows.
    cap = str(wb["Summary"]["A2"].value)
    assert "NET EUR" in cap
    assert "gross = net + VAT" in cap
    assert "gross = net + VAT" in str(wb["Vehicles"]["A2"].value)


def test_workbook_entity_filter(tmp_path):
    from openpyxl import load_workbook
    path = reports.expense_report_workbook(
        "2026-05", entity="Jupiter Plus AS", path=str(tmp_path / "exp_ent.xlsx"))
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "Entity: Jupiter Plus AS" in str(wb["Summary"]["A2"].value)


def test_expenses_page_renders(client):
    r = client.get("/expenses?period=2026-05")
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert "Expense by entity" in html
    assert "Expense by vehicle" in html
    # NET-EUR basis stated on the page.
    assert "gross = net + VAT" in html
    assert "Download expense report (Excel)" in html


def test_export_expenses_route(client):
    r = client.get("/export/expenses?period=2026-05")
    assert r.status_code == 200
    assert r.get_data()[:2] == b"PK"     # valid xlsx (zip)
