"""Overpay -> supplier price-review packet (the analytics->act-on-it gap).

Covers:
- queries.q_savings_lines: the hand-verifiable per-day overpay line, the <2-suppliers
  edge, the optional supplier filter, and the KEY consistency check — the per-supplier
  sum of q_savings_lines reconciles with q_savings' by_supplier aggregate.
- reports.overpay_review_workbook: a valid styled workbook for a seeded period, and for
  a single-supplier filter.
- /savings web smoke: the per-supplier review table renders (escaped) and the
  price-review packet downloads.

This is a price-COMPETITIVENESS / negotiation review (supplier X charged more than the
cheapest same-day, same-country rival) — NOT a contractual claim. NET EUR/L basis.
"""
import sqlite3

import queries
import reports


def _seed(rows):
    con = sqlite3.connect(":memory:"); con.row_factory = sqlite3.Row
    con.execute("""CREATE TABLE transactions (period TEXT, date TEXT, country TEXT,
                   supplier TEXT, product_group TEXT, qty REAL, net_eur_eff REAL)""")
    con.executemany("INSERT INTO transactions VALUES (?,?,?,?,?,?,?)", rows)
    con.commit()
    return con


def test_lines_hand_verifiable():
    # A 100L@1.40, B 100L@1.50 same day+country -> ONE line: B overpays 10.00 vs A,
    # delta 0.10, cheapest A. The cheaper supplier yields NO line.
    con = _seed([
        ("2026-05", "2026-05-10", "Latvia", "A", "Diesel", 100.0, 140.0),
        ("2026-05", "2026-05-10", "Latvia", "B", "Diesel", 100.0, 150.0),
    ])
    lines = queries.q_savings_lines(con, "2026-05")
    con.close()
    assert len(lines) == 1
    ln = lines[0]
    assert ln["supplier"] == "B"
    assert ln["date"] == "2026-05-10" and ln["country"] == "Latvia"
    assert ln["litres"] == 100.0
    assert round(ln["eur_l"], 4) == 1.50
    assert round(ln["cheapest_eur_l"], 4) == 1.40
    assert ln["cheapest_supplier"] == "A"
    assert round(ln["delta_eur_l"], 4) == 0.10
    assert ln["overpay_eur"] == 10.00


def test_lines_single_supplier_no_line():
    # Only one supplier in a day+country -> no rival -> no line.
    con = _seed([
        ("2026-05", "2026-05-10", "Latvia", "A", "Diesel", 100.0, 140.0),
    ])
    lines = queries.q_savings_lines(con, "2026-05")
    con.close()
    assert lines == []


def test_lines_supplier_filter():
    con = _seed([
        ("2026-05", "2026-05-10", "Latvia", "CHEAP", "Diesel", 100.0, 130.0),
        ("2026-05", "2026-05-10", "Latvia", "MID", "Diesel", 100.0, 140.0),
        ("2026-05", "2026-05-10", "Latvia", "DEAR", "Diesel", 100.0, 150.0),
    ])
    all_lines = queries.q_savings_lines(con, "2026-05")
    dear = queries.q_savings_lines(con, "2026-05", supplier="DEAR")
    cheap = queries.q_savings_lines(con, "2026-05", supplier="CHEAP")
    con.close()
    # CHEAP never overpays; MID and DEAR each produce one line.
    assert {ln["supplier"] for ln in all_lines} == {"MID", "DEAR"}
    assert len(dear) == 1 and dear[0]["supplier"] == "DEAR"
    assert dear[0]["overpay_eur"] == 20.00     # 100*(1.50-1.30)
    assert cheap == []                          # cheapest -> filtered to nothing


def test_lines_sorted_overpay_desc():
    con = _seed([
        ("2026-05", "2026-05-10", "Latvia", "CHEAP", "Diesel", 100.0, 130.0),
        ("2026-05", "2026-05-10", "Latvia", "MID", "Diesel", 100.0, 140.0),
        ("2026-05", "2026-05-10", "Latvia", "DEAR", "Diesel", 100.0, 150.0),
    ])
    lines = queries.q_savings_lines(con, "2026-05")
    con.close()
    overs = [ln["overpay_eur"] for ln in lines]
    assert overs == sorted(overs, reverse=True)


def test_lines_sum_matches_aggregate():
    # KEY correctness: the per-supplier sum of q_savings_lines reconciles with
    # q_savings' by_supplier (the detail must not fork the math from the aggregate).
    con = _seed([
        ("2026-05", "2026-05-10", "Latvia", "A", "Diesel", 100.0, 140.0),
        ("2026-05", "2026-05-10", "Latvia", "B", "Diesel", 100.0, 150.0),
        ("2026-05", "2026-05-11", "Estonia", "A", "Diesel", 50.0, 60.0),
        ("2026-05", "2026-05-11", "Estonia", "C", "Diesel", 50.0, 75.0),
        ("2026-05", "2026-05-12", "Latvia", "B", "Diesel", 200.0, 300.0),
        ("2026-05", "2026-05-12", "Latvia", "C", "Diesel", 200.0, 280.0),
    ])
    agg = queries.q_savings(con, "2026-05")
    lines = queries.q_savings_lines(con, "2026-05")
    con.close()
    by_sup = {}
    for ln in lines:
        by_sup[ln["supplier"]] = by_sup.get(ln["supplier"], 0.0) + ln["overpay_eur"]
    agg_by_sup = dict(agg["by_supplier"])
    assert set(by_sup) == set(agg_by_sup)
    for sup, v in agg_by_sup.items():
        # aggregate keeps full precision; detail quantizes per-line. Reconcile to the cent.
        assert abs(round(by_sup[sup], 2) - round(v, 2)) < 0.01, (sup, by_sup[sup], v)


def test_lines_never_raises_on_bad_con():
    # q_savings_lines is wrapped to never raise (returns []).
    con = sqlite3.connect(":memory:")   # no transactions table
    assert queries.q_savings_lines(con, "2026-05") == []
    con.close()


def test_workbook_builds(tmp_path):
    from openpyxl import load_workbook
    path = reports.overpay_review_workbook(
        "2026-05", path=str(tmp_path / "ov.xlsx"))
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames and "Detail" in wb.sheetnames
    # caveat + basis stated on the title rows.
    assert "competitiveness" in str(wb["Summary"]["A2"].value).lower()
    assert "not a contractual claim" in str(wb["Detail"]["A2"].value).lower()
    assert "NET EUR/L" in str(wb["Summary"]["A2"].value)


def test_workbook_single_supplier(tmp_path):
    from openpyxl import load_workbook
    path = reports.overpay_review_workbook(
        "2026-05", supplier="DKV", path=str(tmp_path / "ov_dkv.xlsx"))
    wb = load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "Supplier: DKV" in str(wb["Summary"]["A2"].value)


def test_savings_page_renders_review(client):
    r = client.get("/savings?period=2026-05")
    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert "price-review" in html.lower()
    # caveat present on the page
    assert "not a contractual claim" in html.lower()
    assert "Download price-review packet" in html


def test_export_overpay_route(client):
    r = client.get("/export/overpay?period=2026-05")
    assert r.status_code == 200
    assert r.get_data()[:2] == b"PK"     # valid xlsx (zip)
