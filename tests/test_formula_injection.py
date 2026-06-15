"""CSV / spreadsheet formula injection (CWE-1236) neutralization in the exports.

Free-text fields that can originate from INGESTED SUPPLIER DATA (supplier name, station,
vehicle, note, entity, customer) are written into CSV cells and openpyxl workbook cells.
A value whose first char is `= + - @`, TAB, CR or LF is interpreted as a FORMULA when the
client opens the file in Excel / Google Sheets -> arbitrary calc / data exfiltration on the
client's machine. `reports._formula_safe` prefixes a single apostrophe (OWASP mitigation)
so the spreadsheet treats it as text; numbers/dates/code-authored `=SUM(...)` formulas are
left untouched.

Covers:
- _formula_safe: dangerous leading chars get prefixed; normal strings + non-str pass through.
- Accounting CSV: a malicious supplier/note round-trips as a neutralized `'=...` text cell
  while numeric cells + the reconciliation are intact.
- expense_report_workbook: a malicious entity is written as a TEXT cell (data_type != 'f'),
  while the code-authored `=SUM(...)` TOTAL cells remain real formulas (data_type == 'f').
"""
import csv
import io
import datetime
import decimal

from openpyxl import load_workbook

import queries
import reports


# ----------------------------------------------------------------- _formula_safe

def test_formula_safe_prefixes_dangerous_leads():
    assert reports._formula_safe("=cmd") == "'=cmd"
    assert reports._formula_safe("+1") == "'+1"
    assert reports._formula_safe("-1") == "'-1"
    assert reports._formula_safe("@x") == "'@x"
    assert reports._formula_safe("\tx") == "'\tx"
    assert reports._formula_safe("\rx") == "'\rx"
    assert reports._formula_safe("\nx") == "'\nx"


def test_formula_safe_leaves_normal_strings():
    assert reports._formula_safe("Q8 Belgium") == "Q8 Belgium"
    assert reports._formula_safe("") == ""
    assert reports._formula_safe("Diesel") == "Diesel"
    # a dangerous char that is NOT the first char is left alone
    assert reports._formula_safe("a=b") == "a=b"


def test_formula_safe_passes_non_str_untouched():
    assert reports._formula_safe(5) == 5
    assert reports._formula_safe(3.14) == 3.14
    assert reports._formula_safe(decimal.Decimal("2.50")) == decimal.Decimal("2.50")
    assert reports._formula_safe(None) is None
    d = datetime.date(2026, 5, 1)
    assert reports._formula_safe(d) is d


# ----------------------------------------------------------- accounting CSV

def test_csv_neutralizes_malicious_freetext(monkeypatch):
    """A supplier/note beginning with `=` is written as a `'=...` TEXT cell; the numeric
    cells are unchanged and the net/vat/gross still reconcile per row."""
    malicious = {
        "date": "2026-05-01", "period": "2026-05", "entity": "A",
        "supplier": "=cmd|'/c calc'!A1", "country": "Latvia", "vehicle": "V1",
        "station": "ST", "product": "Diesel", "product_group": "Diesel", "qty": 100.0,
        "currency": "EUR", "net_local": 130.0, "vat_local": 27.3, "gross_local": 157.3,
        "net_eur": 130.0, "vat_eur": 27.3, "gross_eur": 157.3, "vat_rate_pct": 21.0,
        "note": '=HYPERLINK("http://evil","click")',
    }
    monkeypatch.setattr(queries, "q_ledger", lambda con, period, entity=None: [malicious])

    name, data = reports.accounting_ledger_csv("2026-05")
    parsed = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    header, row = parsed[0], parsed[1]
    cols = {h: i for i, h in enumerate(header)}

    # the dangerous free-text cells are neutralized (leading apostrophe), no live formula
    assert row[cols["Supplier"]] == "'=cmd|'/c calc'!A1"
    assert row[cols["Note"]].startswith("'=HYPERLINK")

    # numeric cells are intact and still reconcile
    net = float(row[cols["Net EUR"]]); vat = float(row[cols["VAT EUR"]])
    gross = float(row[cols["Gross EUR"]])
    assert net == 130.0 and vat == 27.3
    assert round(gross, 2) == round(net + vat, 2)
    # a non-dangerous free-text cell is left exactly as-is
    assert row[cols["Country"]] == "Latvia"


# ------------------------------------------------------------- xlsx workbook

def test_expense_workbook_neutralizes_freetext_keeps_sum_formulas(monkeypatch, tmp_path):
    """A malicious entity name is written as a TEXT cell (not a live formula) while the
    code-authored `=SUM(...)` TOTAL cells stay real formulas."""
    data = {
        "by_entity": [{
            "entity": "=cmd|'/c calc'!A1", "fuellings": 3, "n_vehicles": 1,
            "litres": 100.0, "net_eur": 130.0, "vat_eur": 27.3, "gross_eur": 157.3,
        }],
        "by_product": [{
            "product_group": "Diesel", "litres": 100.0, "net_eur": 130.0,
            "vat_eur": 27.3, "gross_eur": 157.3,
        }],
        "by_vehicle": [{
            "entity": "=cmd|'/c calc'!A1", "vehicle": "+44 plate", "fuellings": 3,
            "litres": 100.0, "net_eur": 130.0, "vat_eur": 27.3, "gross_eur": 157.3,
            "net_eur_l": 1.3, "n_countries": 1,
        }],
        "totals": {"net_eur": 130.0, "vat_eur": 27.3, "gross_eur": 157.3,
                   "litres": 100.0, "fuellings": 3},
    }
    monkeypatch.setattr(queries, "q_expense", lambda con, period, entity=None: data)

    path = reports.expense_report_workbook("2026-05", path=str(tmp_path / "m.xlsx"))
    wb = load_workbook(path)

    ws = wb["Summary"]
    # the malicious entity cell (col A, first data row at r0+2 == row 6) is neutralized TEXT
    ent_cell = ws.cell(6, 1)
    assert ent_cell.value == "'=cmd|'/c calc'!A1"
    assert ent_cell.data_type != "f"        # NOT a live formula

    # at least one code-authored =SUM(...) TOTAL cell survives as a real formula
    sum_cells = [c for col in ws.iter_cols() for c in col
                 if isinstance(c.value, str) and c.value.startswith("=SUM(")]
    assert sum_cells, "expected code-authored =SUM(...) total cells"
    for c in sum_cells:
        assert c.data_type == "f"

    # the Vehicles sheet: a `+`-leading vehicle plate is neutralized too
    wv = wb["Vehicles"]
    veh_cell = wv.cell(5, 2)
    assert veh_cell.value == "'+44 plate"
    assert veh_cell.data_type != "f"
