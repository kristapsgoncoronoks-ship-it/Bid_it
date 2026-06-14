"""
Master workbook (build_master.build) regression tests for the FX-analysis sheet,
the explicit per-line FX rate on Transactions, the Python-computed cross-tab pivots,
and the Transactions Excel Table.

These run against the demo data: consolidate.py is executed once to (re)write the
period-stamped pickle, then build() is called and the resulting .xlsx is opened with
openpyxl and asserted. The generated workbook is cleaned up; we never assert against a
populated ecb_rates.db (it is gitignored / network-seeded) — for the ECB-dependent
markup we seed a throwaway ecb_rates DB so the deviation column is deterministic.
"""
import collections
import os
import subprocess
import sys

import openpyxl
import pytest

WORKDIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, WORKDIR)

import money  # noqa: E402

PERIOD = "2026-05"


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    """Consolidate the demo period, point ecb_rates at a throwaway DB seeded with a
    known PLN rate, build the master workbook, and yield (workbook, rows). The output
    .xlsx is removed on teardown; the real ecb_rates.DB is restored."""
    proc = subprocess.run([sys.executable, os.path.join(WORKDIR, "consolidate.py")],
                          cwd=WORKDIR, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stdout + proc.stderr

    import consolidate
    import ecb_rates
    import build_master

    rows = consolidate.load_rows(PERIOD)

    # Deterministic ECB cache: a single seeded PLN rate so the FX-analysis deviation /
    # markup for the BP/PLN row is computed regardless of the live (gitignored) cache.
    # The seed stays active THROUGH the yielded tests so any canonical recompute
    # (supplier_fx.analysis_from_rows) uses the same rate the workbook was built with;
    # the real ecb_rates.DB is restored in teardown.
    seed_db = tmp_path_factory.mktemp("ecb") / "ecb_rates.db"
    orig_db = ecb_rates.DB
    ecb_rates.DB = str(seed_db)
    try:
        ecb_rates.store([("2026-05-31", "PLN", 4.31)], source="test seed")
        out_path = build_master.build(PERIOD)
        wb = openpyxl.load_workbook(out_path)
        yield wb, rows
    finally:
        ecb_rates.DB = orig_db
        try:
            os.remove(out_path)
        except OSError:
            pass


def test_expected_sheets_present(workbook):
    wb, _ = workbook
    for name in ("Transactions", "FX analysis", "Pivot Supplier x Country",
                 "Pivot Entity x Country", "Pivot Product x Country"):
        assert name in wb.sheetnames, f"missing sheet {name}: {wb.sheetnames}"


def test_transactions_fx_rate_column_consistent(workbook):
    """The Transactions sheet carries an explicit FX-rate column (S), populated, and
    net_local / rate == net_eur for non-EUR lines; EUR lines carry rate 1.0."""
    wb, _ = workbook
    ws = wb["Transactions"]
    assert ws.cell(row=1, column=19).value == "FX rate (local/EUR)"
    checked = 0
    for rr in range(2, ws.max_row + 1):
        ccy = ws.cell(row=rr, column=11).value
        net_local = ws.cell(row=rr, column=12).value
        net_eur = ws.cell(row=rr, column=15).value
        fx = ws.cell(row=rr, column=19).value
        if not ccy or ccy == "EUR":
            assert fx == 1.0, f"EUR line {rr} fx={fx!r}"
            continue
        if net_eur and fx:
            # net_local / rate ~= net_eur (rate = foreign units per 1 EUR)
            assert abs(net_local / fx - net_eur) < 0.01, (rr, net_local, fx, net_eur)
            checked += 1
    assert checked > 0, "expected at least one non-EUR line to verify"


def test_fx_analysis_known_supplier_row(workbook):
    """FX-analysis sheet exists with a deviation/markup column and a known supplier row
    (BP/PLN): implied rate = net_local/net_eur, and with the seeded ECB rate the
    deviation % and EUR markup are populated and self-consistent."""
    wb, _ = workbook
    ws = wb["FX analysis"]
    hdr = [c.value for c in ws[4]]
    assert "Deviation %" in hdr and "EUR gain/loss (markup)" in hdr, hdr
    ix = {h: i for i, h in enumerate(hdr)}

    bp = None
    for rr in range(5, ws.max_row + 1):
        if ws.cell(row=rr, column=1).value == "BP" and ws.cell(row=rr, column=2).value == "PLN":
            bp = [c.value for c in ws[rr]]
            break
    assert bp is not None, "BP/PLN row not found in FX analysis"

    net_local = bp[ix["Net local"]]
    net_eur = bp[ix["Net EUR"]]
    implied = bp[ix["Implied rate"]]
    ecb = bp[ix["ECB rate"]]
    dev = bp[ix["Deviation %"]]
    markup = bp[ix["EUR gain/loss (markup)"]]

    assert abs(implied - net_local / net_eur) < 1e-4, (implied, net_local, net_eur)
    assert abs(ecb - 4.31) < 1e-6, ecb
    # deviation = (implied - ecb) / ecb * 100
    assert abs(dev - (implied - ecb) / ecb * 100) < 1e-3, (dev, implied, ecb)
    # eur markup = net_eur - net_local / ecb
    assert abs(markup - money.f2(net_eur - net_local / ecb)) < 0.01, (markup,)


def test_pivot_supplier_country_cell_matches_raw(workbook):
    """A known cell in Pivot Supplier x Country (BP, Poland) matches the litres and
    Net EUR summed from the raw rows."""
    wb, rows = workbook
    raw_litres = collections.defaultdict(float)
    raw_net = collections.defaultdict(float)
    for r_ in rows:
        raw_litres[(r_[1], r_[2])] += r_[9] or 0.0
        raw_net[(r_[1], r_[2])] += r_[14] or 0.0

    ws = wb["Pivot Supplier x Country"]
    hdr = [c.value for c in ws[5]]
    assert "Poland" in hdr, hdr
    pol = hdr.index("Poland") + 1

    # litres block: first BP row under header row 5
    bp_litres = None
    for rr in range(6, ws.max_row + 1):
        if ws.cell(row=rr, column=1).value == "BP":
            bp_litres = ws.cell(row=rr, column=pol).value
            break
    assert bp_litres == round(raw_litres[("BP", "Poland")], 0), bp_litres

    # Net EUR block: BP row under the "Net EUR" metric header
    bp_net = None
    for rr in range(1, ws.max_row + 1):
        if ws.cell(row=rr, column=1).value == "Net EUR":
            for k in range(rr + 2, ws.max_row + 1):
                if ws.cell(row=k, column=1).value == "BP":
                    bp_net = ws.cell(row=k, column=pol).value
                    break
            break
    assert bp_net == money.f2(raw_net[("BP", "Poland")]), bp_net


def test_executive_summary_sheet_present_and_first(workbook):
    """The Executive summary is the FIRST sheet and carries the period + NET basis."""
    wb, _ = workbook
    assert wb.sheetnames[0] == "Executive summary", wb.sheetnames
    ws = wb["Executive summary"]
    assert "EXECUTIVE SUMMARY" in str(ws["A1"].value)
    assert "net EUR/L" in str(ws["A2"].value) and PERIOD in str(ws["A2"].value)


def _kpi_value(ws, label):
    """Find the KPI value cell sitting directly ABOVE its label cell."""
    for row in ws.iter_rows():
        for c in row:
            if c.value == label:
                return ws.cell(row=c.row - 1, column=c.column).value
    return None


def test_executive_summary_figures_preserved(workbook):
    """PRESERVATION: the headline KPIs on the Executive summary EQUAL the figures summed
    straight from the raw rows — proving the KPI sheet only re-displays, never recomputes
    differently. Net spend and avoidable overpay are the load-bearing checks."""
    wb, rows = workbook
    ws = wb["Executive summary"]

    # Net spend (EUR) == money.f2 of sum of net_eur over ALL rows (col 14).
    raw_net = money.f2(money.fsum(r_[14] or 0.0 for r_ in rows))
    assert _kpi_value(ws, "Net spend (EUR)") == raw_net

    # Reclaimable VAT (EUR) == money.f2 of sum of vat_eur (col 15).
    raw_vat = money.f2(money.fsum(r_[15] or 0.0 for r_ in rows))
    assert _kpi_value(ws, "Reclaimable VAT (EUR)") == raw_vat

    # Avoidable overpay (EUR): recompute the SAME head-to-head logic and confirm equality.
    hh = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0, 0.0]))
    for r_ in rows:
        if r_[8] == "Diesel":
            hh[(r_[4], r_[2])][r_[1]][0] += r_[9]
            hh[(r_[4], r_[2])][r_[1]][1] += r_[16]
    overpay = 0.0
    for _k, bysup in hh.items():
        if len(bysup) >= 2:
            prices = {s: v[1] / v[0] for s, v in bysup.items()}
            cheap = min(prices.values())
            overpay += sum(v[0] * (prices[s] - cheap) for s, v in bysup.items())
    assert _kpi_value(ws, "Avoidable overpay (EUR)") == money.f2(overpay)


def test_executive_summary_and_benchmark_have_charts(workbook):
    """At least one chart on the Executive summary (overpay by country) and one on the
    Diesel benchmark sheet (effective €/L by supplier)."""
    wb, _ = workbook
    assert len(wb["Executive summary"]._charts) >= 1
    assert len(wb["Diesel benchmark"]._charts) >= 1


def test_fx_markup_cell_unchanged(workbook):
    """A known EUR cell on an existing sheet (BP/PLN markup on FX analysis) is byte-for-byte
    the canonical supplier_fx figure — formatting/highlighting did not perturb the number."""
    wb, rows = workbook
    import supplier_fx
    canon = {(d["supplier"], d["currency"]): d for d in supplier_fx.analysis_from_rows(rows)}
    bp = canon[("BP", "PLN")]
    ws = wb["FX analysis"]
    hdr = [c.value for c in ws[4]]
    ix = {h: i for i, h in enumerate(hdr)}
    for rr in range(5, ws.max_row + 1):
        if ws.cell(row=rr, column=1).value == "BP" and ws.cell(row=rr, column=2).value == "PLN":
            cell_markup = ws[rr][ix["EUR gain/loss (markup)"]].value
            assert cell_markup == money.f2(bp["eur_diff"]), (cell_markup, bp["eur_diff"])
            return
    pytest.fail("BP/PLN row not found")


def test_transactions_registered_as_excel_table(workbook):
    """The Transactions data range is registered as a proper Excel Table so the user
    can drop their own native PivotTables / filters on clean tabular data."""
    wb, _ = workbook
    ws = wb["Transactions"]
    assert "TransactionsTbl" in ws.tables, list(ws.tables)
    ref = ws.tables["TransactionsTbl"].ref
    assert ref.startswith("A1:S"), ref


def test_build_reads_stored_fx_rate(monkeypatch):
    """build SURFACES the PERSISTED transactions.fx_rate on the Transactions sheet: a
    seeded stored rate that DIFFERS from the net_local/net_eur recompute shows up in
    column S, proving the sheet reads storage rather than always re-deriving."""
    import os
    import sqlite3
    import consolidate
    import history
    import dataproduct
    import build_master

    period = "2099-09"
    import tempfile
    tmp = tempfile.mkdtemp()
    pkl = os.path.join(tmp, "consolidated_rows.pkl")
    hist_db = os.path.join(tmp, "fuel_history.db")
    out = os.path.join(tmp, f"Fleet_Fuel_Master_{period}.xlsx")

    fixture = [
        ["ENT", "BP", "PL", "CARP", "2026-05-10", "08:00", "Stat PL", "Diesel",
         "Diesel", 500.0, "PLN", 4270.0, 982.10, 5252.10, 1000.0, 230.0, 1000.0, ""],
    ]
    consolidate._dump_pickle(fixture, period, path=pkl)

    real = consolidate.load_rows
    monkeypatch.setattr(consolidate, "load_rows",
                        lambda p, path=pkl: real(p, path=path))
    monkeypatch.setattr(history, "DB", hist_db, raising=True)
    monkeypatch.setattr(build_master, "WORKDIR", tmp, raising=True)
    # dataproduct.connect("fuel_history") must hit the temp DB build_master reads from
    monkeypatch.setattr(dataproduct, "connect",
                        lambda which="fuel_history", path=None: sqlite3.connect(hist_db),
                        raising=True)

    history.load(period)

    # Overwrite the STORED rate with a sentinel that is NOT net_local/net_eur (4.27),
    # so reading-vs-recompute is distinguishable on the sheet.
    con = sqlite3.connect(hist_db)
    con.execute("UPDATE transactions SET fx_rate=? WHERE period=?", (9.99, period))
    con.commit(); con.close()

    out_path = build_master.build(period)
    try:
        wb = openpyxl.load_workbook(out_path)
        ws = wb["Transactions"]
        assert ws.cell(row=2, column=19).value == 9.99, "build did not surface stored fx_rate"
    finally:
        try:
            os.remove(out_path)
        except OSError:
            pass


def test_build_fx_falls_back_when_unstored(monkeypatch):
    """When NO stored fx_rate is available (empty / unloaded period), build falls back to
    the canonical net_local/net_eur recompute — the figure is identical by construction."""
    import build_master
    monkeypatch.setattr(build_master, "_stored_fx_map", lambda period: {})
    import consolidate
    period = "2026-05"
    rows = consolidate.load_rows(period)
    out_path = build_master.build(period)
    try:
        wb = openpyxl.load_workbook(out_path)
        ws = wb["Transactions"]
        checked = 0
        for rr in range(2, ws.max_row + 1):
            ccy = ws.cell(row=rr, column=11).value
            net_local = ws.cell(row=rr, column=12).value
            net_eur = ws.cell(row=rr, column=15).value
            fx = ws.cell(row=rr, column=19).value
            if ccy and ccy != "EUR" and net_eur and fx:
                assert abs(fx - net_local / net_eur) < 1e-9, (rr, fx, net_local, net_eur)
                checked += 1
        assert checked > 0
    finally:
        try:
            os.remove(out_path)
        except OSError:
            pass
