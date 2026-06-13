"""Tests for the polished Excel reports (summary + styled exports)."""
import os


def test_summary_workbook_structure(tmp_path):
    import reports
    from openpyxl import load_workbook
    path = reports.summary_workbook(path=str(tmp_path / "sum.xlsx"))
    assert os.path.exists(path)
    wb = load_workbook(path)
    # all the expected sheets are present
    for sheet in ("Summary", "By supplier", "By country", "By entity (VAT)", "Trend", "Savings"):
        assert sheet in wb.sheetnames
    ws = wb["Summary"]
    assert "Fleet Fuel Report" in str(ws["A1"].value)
    assert "NET EUR/L" in str(ws["A2"].value)  # basis stated on the first sheet
    assert len(ws._charts) >= 1               # a chart was added
    # KPI value cells are numeric
    assert isinstance(ws.cell(4, 3).value, (int, float))
    # By supplier carries two charts (net spend + effective €/L); savings + entity each one.
    assert len(wb["By supplier"]._charts) >= 2
    assert len(wb["By entity (VAT)"]._charts) >= 1
    assert len(wb["Savings"]._charts) >= 1


def test_summary_kpi_figure_preserved(tmp_path):
    """PRESERVATION: the Net-spend KPI on the Summary sheet EQUALS the sum of the per-supplier
    Net EUR rows on the 'By supplier' sheet — the polish/charts did not change the figure."""
    import reports
    from openpyxl import load_workbook
    path = reports.summary_workbook(path=str(tmp_path / "sum.xlsx"))
    wb = load_workbook(path)
    sup = wb["By supplier"]
    # sum the Net EUR column (col 3) over the data rows (header=1, last row is TOTAL).
    body_net = 0.0
    for rr in range(2, sup.max_row):  # exclude TOTAL row (has a formula, not a value)
        v = sup.cell(rr, 3).value
        if isinstance(v, (int, float)):
            body_net += v
    # the Summary "Net spend (EUR)" KPI sits at row 4, col 3 (INK card).
    kpi_net = wb["Summary"].cell(4, 3).value
    assert abs(kpi_net - round(body_net, 2)) < 0.01, (kpi_net, body_net)


def test_export_summary_route(client):
    r = client.get("/export/summary")
    assert r.status_code == 200
    assert r.get_data()[:2] == b"PK"          # valid xlsx (zip)


def test_export_compare_styled(client):
    r = client.get("/export/compare?period=2026-05")
    assert r.status_code == 200 and r.get_data()[:2] == b"PK"


def test_fees_statement_workbook(tmp_path):
    import reports
    from openpyxl import load_workbook
    rows = [
        # paid to US: fee deducted, net remitted = 1000-130 = 870
        dict(entity="Acme", country="DE", period="2026-Q1", vat_eur=1000.0, status="paid",
             paid_amount=1000.0, age_days="", fee_eur=130.0, fee_pct=8.0, fee_min=130.0,
             fee_billed_date="2026-04-01", payout_to="us", fee_invoice_no=None),
        # paid to CUSTOMER: fee invoiced (receivable)
        dict(entity="Acme", country="PL", period="2026-Q1", vat_eur=2000.0, status="paid",
             paid_amount=2000.0, age_days="", fee_eur=160.0, fee_pct=8.0, fee_min=130.0,
             fee_billed_date="2026-04-02", payout_to="customer", fee_invoice_no="F2026-0001"),
        # not yet billed -> excluded from the statement
        dict(entity="Beta", country="DE", period="2026-Q1", vat_eur=500.0, status="submitted",
             paid_amount=None, age_days=10, fee_eur=None, fee_pct=8.0, fee_min=130.0,
             fee_billed_date=None, payout_to=None, fee_invoice_no=None),
    ]
    path = reports.fees_statement_workbook(rows, "2026", path=str(tmp_path / "fees.xlsx"))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Fees by customer", "Fee detail"]
    ws = wb["Fees by customer"]
    # one aggregated customer row (Acme); Beta excluded (unbilled)
    body = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] == "Acme"]
    assert len(body) == 1
    _, claims, vat, fee, recv, net = body[0]
    assert claims == 2 and vat == 1000 + 2000 and fee == 130 + 160
    assert recv == 160 and net == 870
    # detail sheet only has the two billed lines
    assert wb["Fee detail"].max_row - 1 == 2


def test_export_fees_route(client):
    r = client.get("/export/fees?year=2026")
    assert r.status_code == 200 and r.get_data()[:2] == b"PK"
