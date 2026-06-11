import pickle, collections
from supplier_specs import SPECS
from month_config import PERIOD, PAYMENTS, OPEN_ITEMS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import os
WORKDIR = os.path.dirname(os.path.abspath(__file__))
ROWS = pickle.load(open(f"{WORKDIR}/consolidated_rows.pkl","rb"))

wb = Workbook()
fillH = PatternFill("solid", start_color="1F3864")
fillI = PatternFill("solid", start_color="DCE6F1")
bw = Font(bold=True, color="FFFFFF", name="Arial", size=9)
b10 = Font(bold=True, name="Arial", size=10)
norm = Font(name="Arial", size=9)
it8 = Font(italic=True, name="Arial", size=8)
def head(ws, row):
    for c in ws[row]:
        if c.value: c.font = bw; c.fill = fillH; c.alignment = Alignment(horizontal="center", wrap_text=True)

# ============ 1. RUNBOOK ============
ws = wb.active; ws.title = "Runbook"
ws["A1"] = "FLEET FUEL MONTHLY CLOSE - RUNBOOK"; ws["A1"].font = Font(bold=True, size=13, name="Arial")
ws["A2"] = "Period: May 2026 (template - reuse monthly)"; ws["A2"].font = it8
steps = [
("1. Collect","Download invoice PDF + transaction CSV/XLS export from each supplier portal (see Supplier specs, col 'Portal export'). Target: machine-readable data, PDF kept as audit artifact."),
("2. Extract","One workbook per supplier in canonical layout (Transactions + control sheets). If CSV obtained: import directly. If PDF only: extract per the supplier spec sheet."),
("3. Self-control","Run the control formulas (built into each supplier workbook): line math, card/vehicle subtotals, invoice grand totals, VAT recomputation. Tolerances per Supplier specs."),
("4. Exceptions only","Review ONLY: CHECK flags, cash-at-pump, zero-discount stations, unexplained rebate lines, product anomalies (e.g. HVO premium), missing vehicles."),
("5. Consolidate","Run consolidate.py - maps all supplier workbooks into the Transactions sheet of this master. Update FX inputs first."),
("6. Review views","Pivot: supplier x country net EUR/L. Entity & VAT view: payables + reclaimable VAT per registration. Station scorecard: routing actions."),
("7. Actions","Send routing guidance to dispatchers (worst stations), file VAT refund data per registration, log supplier negotiation evidence, schedule payments per calendar."),
("8. History","Run history.py - loads the period into fuel_history.db (SQLite, idempotent) and regenerates Fleet_Fuel_History_Report.xlsx with month-over-month trends. The DB is the permanent archive; the report is the exportable Excel."),
]
r = 4
for s, d in steps:
    ws[f"A{r}"] = s; ws[f"A{r}"].font = b10
    ws[f"B{r}"] = d; ws[f"B{r}"].font = norm; ws[f"B{r}"].alignment = Alignment(wrap_text=True)
    r += 1
ws[f"A{r+1}"] = "Onboarding a NEW supplier (one-time training, ~30 min):"; ws[f"A{r+1}"].font = b10
onb = ["1. Get first invoice PDF + data source: portal CSV/XLS export, XML e-invoice, or API access (preferred).",
       "2. File sources: build/import the supplier workbook. XML/API sources: declare 'source' in the spec",
       "   (record path + field map for XML; URL + auth env-var for API) - see ingest.py header for examples.",
       "3. In supplier_specs.py: copy _TEMPLATE, fill metadata (entity, currency, quirks, controls, portal, terms)",
       "   and write the row_map (~5-10 lines mapping workbook columns to canonical fields).",
       "4. Fill 'expected' with headline figures FROM THE INVOICE (line count, gross total) - the training target.",
       "5. Add filename to FILES in month_config.py; run consolidate.py. PASS = trained: the supplier appears",
       "   automatically in Diesel benchmark, Entity & VAT view, Station scorecard and the Supplier specs sheet.",
       "   FAIL = engine shows doc vs calc per metric; fix the row_map and rerun. Engine/renderer never change."]
for i, o in enumerate(onb):
    ws[f"B{r+2+i}"] = o; ws[f"B{r+2+i}"].font = norm
r = r + 3 + len(onb)
ws[f"A{r}"] = "Open items this month:"; ws[f"A{r}"].font = b10
for i, o in enumerate(OPEN_ITEMS):
    ws[f"B{r+1+i}"] = "- " + o; ws[f"B{r+1+i}"].font = norm
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 110

# ============ 2. SUPPLIER SPECS ============
ws = wb.create_sheet("Supplier specs")
ws.append(["Supplier","Entity","Country/Scope","Invoice structure & quirks","Control rules & tolerances","Portal export","Payment terms"])
head(ws, 1)
specs = [(("%s" % k), v["entity"][0] + " (" + v["entity"][1] + ")", v["scope"],
          v["quirks"], v["controls"], v["portal"], v["terms"]) for k, v in SPECS.items()]
r = 2
for s in specs:
    ws.append(list(s)); r += 1
for row in ws.iter_rows(min_row=2, max_row=r-1):
    for c in row: c.font = norm; c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEFG",[18,22,14,50,46,22,22]):
    ws.column_dimensions[col].width = w

# ============ 3. FX & INPUTS ============
ws = wb.create_sheet("FX & inputs")
ws.append(["Input","Value","Note"]); head(ws,1)
ws.append(["PLN per EUR (BP conversion)", 4.27, "Indicative May 2026 - update monthly from ECB average; blue = input"])
ws.append(["SEK per EUR (DKV)", "per line", "DKV converts per transaction date; EUR taken from invoice lines"])
ws.append(["Q8 rebate allocation", "per litre by country", "Assumption: Port One rebate spread over ALL litres per country - resplit in Q8 adjusted workbook (blue cells)"])
ws["B2"].fill = fillI
for row in ws.iter_rows(min_row=2, max_row=4):
    for c in row: c.font = norm
for col, w in zip("ABC",[30,16,70]): ws.column_dimensions[col].width = w

# ============ 4. TRANSACTIONS ============
ws = wb.create_sheet("Transactions")
hdr = ["Entity","Supplier","Country","Vehicle/Card","Date","Time","Station","Product (doc)",
       "Product group","Qty (L/pc)","Currency","Net local","VAT local","Gross local",
       "Net EUR","VAT EUR","Net EUR effective","Note"]
ws.append(hdr); head(ws,1)
r = 2
for row in ROWS:
    ws.append(row); r += 1
last = r-1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = norm
for col in "JLMNOPQ":
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=ord(col)-64, max_col=ord(col)-64):
        row[0].number_format = "#,##0.00"
for col, w in zip("ABCDEFGHIJKLMNOPQR",[24,9,10,14,11,6,22,13,11,9,5,10,9,10,10,9,10,22]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:R{last}"

# ============ 5. PIVOT: SUPPLIER x COUNTRY (DIESEL) ============
ws = wb.create_sheet("Diesel benchmark", 1)
ws["A1"] = "DIESEL NET EUR/L BENCHMARK - MAY 2026 (excl. VAT, after on-invoice discounts; 'effective' adds Q8 rebate layer)"
ws["A1"].font = Font(bold=True, size=11, name="Arial")
ws.append([])
ws.append(["Supplier","Country","Litres","Net EUR (doc)","EUR/L (doc)","Net EUR (effective)","EUR/L (effective)"])
head(ws,3)
combos = sorted({(r_[1], r_[2]) for r_ in ROWS if r_[8]=="Diesel"})
r = 4
for sup, ctry in combos:
    b = (f'SUMPRODUCT((Transactions!$B$2:$B${last+0}="{sup}")*(Transactions!$C$2:$C${last}="{ctry}")*'
         f'(Transactions!$I$2:$I${last}="Diesel")*Transactions!')
    ws.append([sup, ctry, f"={b}$J$2:$J${last})", f"=ROUND({b}$O$2:$O${last}),2)",
               f"=D{r}/C{r}", f"=ROUND({b}$Q$2:$Q${last}),2)", f"=F{r}/C{r}"])
    r += 1
tr = r
ws.append(["ALL","", f"=SUM(C4:C{r-1})", f"=SUM(D4:D{r-1})", f"=D{tr}/C{tr}", f"=SUM(F4:F{r-1})", f"=F{tr}/C{tr}"])
r += 1
for row in ws.iter_rows(min_row=4, max_row=r-1):
    for c in row: c.font = norm
    row[2].number_format = "#,##0"; row[3].number_format = "#,##0"; row[5].number_format = "#,##0"
    row[4].number_format = "0.0000"; row[6].number_format = "0.0000"
for c in ws[tr]: c.font = b10
ws.append([]); r += 1
ws.append(["Other product groups","","Litres/pcs","Net EUR","EUR/unit"]); 
for c in ws[r]: 
    if c.value: c.font = b10
r += 1
for pg in ("AdBlue","HVO","Parking","Service/Other","Promo adj"):
    b = f'SUMPRODUCT((Transactions!$I$2:$I${last}="{pg}")*Transactions!'
    ws.append([pg,"",f"={b}$J$2:$J${last})", f"=ROUND({b}$O$2:$O${last}),2)",
               f"=IF(C{r}>0,D{r}/C{r},\"\")"])
    for c in ws[r]: c.font = norm
    ws[f"C{r}"].number_format = "#,##0.00"; ws[f"D{r}"].number_format = "#,##0.00"; ws[f"E{r}"].number_format = "0.0000"
    r += 1
for col, w in zip("ABCDEFG",[16,12,11,13,11,15,13]):
    ws.column_dimensions[col].width = w


# ============ 5b. SUPPLIER COMPARISON ============
ws = wb.create_sheet("Supplier comparison", 2)
ws["A1"] = "SUPPLIER COMPARISON - by location, date, product"
ws["A1"].font = Font(bold=True, size=11, name="Arial")
sups = sorted({r_[1] for r_ in ROWS})

# --- A) Interactive query (edit blue cells) ---
ws["A3"] = "A) INTERACTIVE QUERY - edit the blue cells, all supplier rows recompute"; ws["A3"].font = b10
ws.append([]); ws.append(["Product group","Country (or ALL)","Date from","Date to"])
head(ws, 5)
ws.append(["Diesel","ALL","2026-05-01","2026-05-31"])
for c in ws[6]: c.fill = fillI; c.font = b10
ws.append([])
ws.append(["Supplier","Litres/qty","Net EUR","EUR/L (doc)","EUR/L (effective)","Share of qty"])
head(ws, 8)
r = 9
for s in sups:
    b = (f'SUMPRODUCT((Transactions!$B$2:$B${last}="{s}")*'
         f'(Transactions!$I$2:$I${last}=$A$6)*'
         f'(((Transactions!$C$2:$C${last}=$B$6)+($B$6="ALL"))>0)*'
         f'(Transactions!$E$2:$E${last}>=$C$6)*(Transactions!$E$2:$E${last}<=$D$6)*Transactions!')
    ws.append([s, f"={b}$J$2:$J${last})", f"=ROUND({b}$O$2:$O${last}),2)",
               f"=IF(B{r}>0,C{r}/B{r},\"\")",
               f"=IF(B{r}>0,ROUND({b}$Q$2:$Q${last})/B{r},4),\"\")",
               f"=IF($B${9+len(sups)}>0,B{r}/$B${9+len(sups)},\"\")"])
    r += 1
ws.append(["TOTAL", f"=SUM(B9:B{r-1})", f"=SUM(C9:C{r-1})",
           f"=IF(B{r}>0,C{r}/B{r},\"\")", "", ""])
tq = r; r += 1
for row in ws.iter_rows(min_row=9, max_row=tq):
    for c in row: c.font = norm
    row[1].number_format = "#,##0.00"; row[2].number_format = "#,##0.00"
    row[3].number_format = "0.0000"; row[4].number_format = "0.0000"; row[5].number_format = "0.0%"
for c in ws[tq]: c.font = b10

# --- B) Head-to-head: same day, same country, Diesel ---
r += 2
ws[f"A{r}"] = "B) HEAD-TO-HEAD - days where 2+ suppliers bought DIESEL in the same country (apples-to-apples)"
ws[f"A{r}"].font = b10; r += 1
ws.append(["Date","Country","Supplier prices EUR/L (effective)","Cheapest","Spread EUR/L","Litres that day","Overpay vs cheapest EUR"])
head(ws, r); r += 1
g = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0, 0.0]))
for r_ in ROWS:
    if r_[8] == "Diesel":
        g[(r_[4], r_[2])][r_[1]][0] += r_[9]; g[(r_[4], r_[2])][r_[1]][1] += r_[16]
h2h = []
for (d, c), bysup in g.items():
    if len(bysup) >= 2:
        prices = {s: v[1]/v[0] for s, v in bysup.items()}
        cheap = min(prices, key=prices.get)
        litres = sum(v[0] for v in bysup.values())
        overpay = sum(v[0]*(prices[s]-prices[cheap]) for s, v in bysup.items())
        h2h.append((d, c, " | ".join(f"{s} {prices[s]:.4f}" for s in sorted(prices)),
                    cheap, max(prices.values())-min(prices.values()), litres, overpay))
h2h.sort()
tot_over = 0.0
for row_ in h2h:
    ws.append([row_[0], row_[1], row_[2], row_[3], round(row_[4],4), round(row_[5],0), round(row_[6],2)])
    tot_over += row_[6]
    for c in ws[r]: c.font = norm
    ws[f"E{r}"].number_format = "0.0000"; ws[f"F{r}"].number_format = "#,##0"; ws[f"G{r}"].number_format = "#,##0"
    r += 1
ws.append(["", "", "", "", "", "TOTAL overpay vs cheapest-available:", round(tot_over,2)])
for c in ws[r]: c.font = b10
ws[f"G{r}"].number_format = "#,##0"
r += 1

# --- C) Daily diesel matrix ---
r += 2
ws[f"A{r}"] = "C) DAILY DIESEL PRICE MATRIX - effective net EUR/L by supplier by date (blank = no fueling)"
ws[f"A{r}"].font = b10; r += 1
ws.append(["Date"] + sups + ["Day litres"]); head(ws, r); mstart = r + 1; r += 1
dates = sorted({r_[4] for r_ in ROWS if r_[8] == "Diesel"})
for d in dates:
    rowvals = [d]
    for s in sups:
        b = (f'SUMPRODUCT((Transactions!$B$2:$B${last}="{s}")*(Transactions!$E$2:$E${last}=$A{r})*'
             f'(Transactions!$I$2:$I${last}="Diesel")*Transactions!')
        col = chr(66 + sups.index(s))
        rowvals.append(f'=IFERROR(ROUND({b}$Q$2:$Q${last})/{b}$J$2:$J${last}),4),"")')
    b = (f'SUMPRODUCT((Transactions!$E$2:$E${last}=$A{r})*(Transactions!$I$2:$I${last}="Diesel")*Transactions!')
    rowvals.append(f"={b}$J$2:$J${last})")
    ws.append(rowvals)
    for c in ws[r]: c.font = norm
    for i in range(2, 2+len(sups)): ws.cell(row=r, column=i).number_format = "0.0000"
    ws.cell(row=r, column=2+len(sups)).number_format = "#,##0"
    r += 1
ws[f"A{r+1}"] = ("Effective = after Port One rebate for Q8. Head-to-head compares country-level prices; station-level "
                 "detail is in Station scorecard / Transactions filters. New suppliers appear automatically.")
ws[f"A{r+1}"].font = it8
for col, w in zip("ABCDEFG", [11,13,38,10,12,13,22]):
    ws.column_dimensions[col].width = w
for i in range(len(sups)+2):
    ws.column_dimensions[chr(65+i)].width = max(ws.column_dimensions[chr(65+i)].width or 0, 10)

# ============ 6. ENTITY & VAT VIEW ============
ws = wb.create_sheet("Entity & VAT view", 2)
ws["A1"] = "PER-ENTITY TOTALS, RECLAIMABLE VAT AND PAYMENT CALENDAR - MAY 2026"
ws["A1"].font = Font(bold=True, size=11, name="Arial")
ws.append([])
ws.append(["Entity","Supplier","Country","Net EUR","VAT EUR (reclaimable)","Gross EUR"])
head(ws,3)
combos2 = sorted({(r_[0], r_[1], r_[2]) for r_ in ROWS})
r = 4
for ent, sup, ctry in combos2:
    b = (f'SUMPRODUCT((Transactions!$A$2:$A${last}="{ent}")*(Transactions!$B$2:$B${last}="{sup}")*'
         f'(Transactions!$C$2:$C${last}="{ctry}")*Transactions!')
    ws.append([ent, sup, ctry, f"=ROUND({b}$O$2:$O${last}),2)", f"=ROUND({b}$P$2:$P${last}),2)",
               f"=D{r}+E{r}"])
    r += 1
tr = r
ws.append(["TOTAL","","",f"=SUM(D4:D{r-1})",f"=SUM(E4:E{r-1})",f"=SUM(F4:F{r-1})"]); r += 1
for row in ws.iter_rows(min_row=4, max_row=r-1):
    for c in row: c.font = norm
    for c in row[3:6]: c.number_format = "#,##0.00"
for c in ws[tr]: c.font = b10
r += 1
ws.append([]); r += 1
ws.append(["PAYMENT CALENDAR (from invoices; verify on payment run)"])
ws[f"A{r}"].font = b10; r += 1
ws.append(["Due date","Entity","Supplier","Amount","Currency","Note"]); head(ws, r); r += 1
cal = PAYMENTS
for c_ in cal:
    ws.append(list(c_))
    for c in ws[r]: c.font = norm
    ws[f"D{r}"].number_format = "#,##0.00"
    r += 1
for col, w in zip("ABCDEF",[28,16,14,13,20,40]):
    ws.column_dimensions[col].width = w

# ============ 7. STATION SCORECARD ============
ws = wb.create_sheet("Station scorecard", 3)
ws["A1"] = "DIESEL STATION SCORECARD (>=300 L) - routing guidance"; ws["A1"].font = Font(bold=True, size=11, name="Arial")
agg = collections.defaultdict(lambda: [0.0,0.0,0.0])
for r_ in ROWS:
    if r_[8]=="Diesel":
        k = (r_[1], r_[2], r_[6]); agg[k][0]+=r_[9]; agg[k][1]+=r_[14]; agg[k][2]+=r_[16]
rows = [(s,c,st,v[0],v[1],v[1]/v[0],v[2]/v[0]) for (s,c,st),v in agg.items() if v[0]>=300]
rows.sort(key=lambda x: x[5])
ws.append([])
ws.append(["Supplier","Country","Station","Litres","Net EUR","EUR/L (doc)","EUR/L (eff)","Action hint"])
head(ws,3)
r = 4
for s,c,st,L,E,pl,ple in rows:
    hint = ""
    if pl <= 1.40: hint = "PREFER"
    elif pl >= 1.62: hint = "AVOID / renegotiate"
    ws.append([s,c,st,round(L,0),round(E,2),round(pl,4),round(ple,4),hint])
    for cc in ws[r]: cc.font = norm
    ws[f"D{r}"].number_format = "#,##0"; ws[f"E{r}"].number_format = "#,##0"
    ws[f"F{r}"].number_format = "0.0000"; ws[f"G{r}"].number_format = "0.0000"
    if hint == "PREFER": ws[f"H{r}"].font = Font(bold=True, color="1B7340", name="Arial", size=9)
    if hint.startswith("AVOID"): ws[f"H{r}"].font = Font(bold=True, color="C8102E", name="Arial", size=9)
    r += 1
ws[f"A{r+1}"] = "Sorted cheapest first by doc net EUR/L. 'eff' includes Port One rebate for Q8 lines. Static snapshot - regenerate via consolidate.py."
ws[f"A{r+1}"].font = it8
for col, w in zip("ABCDEFGH",[9,10,30,9,10,11,11,20]):
    ws.column_dimensions[col].width = w

wb.save(f"{WORKDIR}/Fleet_Fuel_Master_{PERIOD}.xlsx")
print("master saved;", last-1, "txn rows;", len(combos), "diesel combos;", len(rows), "scorecard stations")
