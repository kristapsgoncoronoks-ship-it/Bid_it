"""
HISTORY LAYER - SQLite database for cross-month comparison + exportable Excel report.

Monthly use (after consolidate.py PASSes):
    python3 history.py          # loads current period into fuel_history.db (idempotent:
                                # re-running replaces that period), then regenerates
                                # Fleet_Fuel_History_Report.xlsx from the full database.

Database: fuel_history.db
    transactions  - every canonical line, all periods (full fidelity, ~690 rows/month)
    v_supplier_month / v_entity_month / v_station_month - reporting views

Ad-hoc queries (sqlite3 fuel_history.db or any SQL tool):
    SELECT period, supplier, ROUND(SUM(net_eur_eff)/SUM(qty),4) AS eur_l
    FROM transactions WHERE product_group='Diesel' GROUP BY period, supplier;
"""
import sqlite3, collections
import db_tuning
import db_migrate
import money
import consolidate
import month_config
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

import os
WORKDIR = os.path.dirname(os.path.abspath(__file__))
DB = f"{WORKDIR}/fuel_history.db"
FIELDS = ["entity","supplier","country","vehicle","date","time","station","product",
          "product_group","qty","currency","net_local","vat_local","gross_local",
          "net_eur","vat_eur","net_eur_eff","note"]

# Migrations for fuel_history.db, run once each via db_migrate's versioning. APPEND new
# DDL at the END — positions are stable. The base `transactions` schema is created by the
# CREATE TABLE IF NOT EXISTS in load(); these alter it forward for already-existing DBs.
_MIGR = [
    # finding #4 (FX provenance): freeze the APPLIED local->EUR rate per line at
    # consolidation so a historical claim's EUR is traceable to a stored rate even if FX
    # sources change later. Convention = foreign units per 1 EUR (ECB, = net_local/net_eur).
    "ALTER TABLE transactions ADD COLUMN fx_rate REAL",
]


def fx_rate(net_local, net_eur, currency=None):
    """APPLIED FX rate for a line, in ECB convention (foreign units per 1 EUR), i.e.
    exactly the rate that produced net_eur: `net_local / net_eur`.

    Returns None when there is no EUR basis to divide by (net_eur 0/None) or no local
    amount (net_local None) — we store NULL rather than fabricate or divide by zero.
    EUR-native lines fall out naturally at 1.0 (net_local == net_eur). The currency arg
    is advisory only (not required for the math) so callers can pass it for clarity.

    Decimal is used for the division so the stored REAL is exact-ish; net_local/net_eur
    are NOT touched — this is a purely additional, derived value.
    """
    if net_local is None or not net_eur:
        return None
    return float(money.D(net_local) / money.D(net_eur))


def load(period=None):
    """Load `period` (default month_config.PERIOD) into fuel_history.db, then
    regenerate Fleet_Fuel_History_Report.xlsx from the full database.

    Idempotent / RESTARTABLE: the load is DELETE-by-period + INSERT, so re-running
    REPLACES that period's rows (never duplicates). Reads the consolidate pickle via
    consolidate.load_rows(period), which ASSERTS the pickle was written for `period`
    (reliability finding #3) before any DB write.

    IMPORTING this module is SIDE-EFFECT-FREE — the whole load/report runs only on this
    call (or via main()/__main__), so a failure leaves a clean restart point instead of
    a half-built Workbook aborted at module import (reliability finding #7). Returns the
    output Excel path.
    """
    period = period or month_config.PERIOD
    ROWS = consolidate.load_rows(period)

    # ---------------- 1. LOAD ----------------
    # This is the canonical engine WRITER of fuel_history.db. Apply the same WAL +
    # busy_timeout tuning the app/consumer connections use (pricing_intelligence,
    # vat_refund.analytics_connect) so every handle on this file agrees on the
    # journal mode — mixed WAL/rollback risks reader/writer contention on close.
    # (build_master.py writes only consolidated_rows.pkl and consolidate.py writes
    # no DB — neither holds a product-DB handle, so neither needs tuning.)
    con = sqlite3.connect(DB)
    db_tuning.tune(con)   # WAL + busy_timeout; idempotent, no-ops on :memory:/Postgres
    con.executescript("""
CREATE TABLE IF NOT EXISTS transactions (
    period TEXT NOT NULL,
    entity TEXT, supplier TEXT, country TEXT, vehicle TEXT,
    date TEXT, time TEXT, station TEXT, product TEXT, product_group TEXT,
    qty REAL, currency TEXT, net_local REAL, vat_local REAL, gross_local REAL,
    net_eur REAL, vat_eur REAL, net_eur_eff REAL, note TEXT
);
CREATE INDEX IF NOT EXISTS ix_t_psc ON transactions(period, supplier, country);
CREATE INDEX IF NOT EXISTS ix_t_pg  ON transactions(product_group, period);
CREATE VIEW IF NOT EXISTS v_supplier_month AS
    SELECT period, supplier, country, product_group,
           ROUND(SUM(qty),2) litres, ROUND(SUM(net_eur),2) net_eur,
           ROUND(SUM(vat_eur),2) vat_eur, ROUND(SUM(net_eur_eff),2) net_eur_eff,
           ROUND(SUM(net_eur)/NULLIF(SUM(qty),0),4) eur_l_doc,
           ROUND(SUM(net_eur_eff)/NULLIF(SUM(qty),0),4) eur_l_eff
    FROM transactions GROUP BY period, supplier, country, product_group;
CREATE VIEW IF NOT EXISTS v_entity_month AS
    SELECT period, entity, country,
           ROUND(SUM(net_eur),2) net_eur, ROUND(SUM(vat_eur),2) vat_eur,
           ROUND(SUM(net_eur)+SUM(vat_eur),2) gross_eur
    FROM transactions GROUP BY period, entity, country;
CREATE VIEW IF NOT EXISTS v_station_month AS
    SELECT period, supplier, country, station,
           ROUND(SUM(qty),2) litres,
           ROUND(SUM(net_eur_eff)/NULLIF(SUM(qty),0),4) eur_l_eff
    FROM transactions WHERE product_group='Diesel'
    GROUP BY period, supplier, country, station;
""")
    # Add the derived fx_rate column (and any later additive columns) to the base
    # transactions table. db_migrate is versioned, so each ALTER runs EXACTLY ONCE per DB
    # (idempotent on re-run); a fresh DB gets it here too, just after the CREATE.
    db_migrate.apply(con, "history", _MIGR)

    con.execute("DELETE FROM transactions WHERE period=?", (period,))
    # Persist the APPLIED FX rate per line alongside the figures (additive — net_*/vat_*
    # are written unchanged). net_local idx 11, net_eur idx 14, currency idx 10 in FIELDS.
    con.executemany(
        f"INSERT INTO transactions (period,{','.join(FIELDS)},fx_rate) "
        f"VALUES ({','.join('?'*19)},?)",
        [[period]+list(r)+[fx_rate(r[11], r[14], r[10])] for r in ROWS])
    con.commit()
    periods = [p[0] for p in con.execute("SELECT DISTINCT period FROM transactions ORDER BY period")]
    n = con.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    print(f"DB loaded: period {period} replaced; total {n} rows across periods {periods}")

    # ---------------- 2. REPORT ----------------
    wb = Workbook()
    fillH = PatternFill("solid", start_color="2F5233")
    bw = Font(bold=True, color="FFFFFF", name="Arial", size=9)
    b10 = Font(bold=True, name="Arial", size=10)
    norm = Font(name="Arial", size=9)
    it8 = Font(italic=True, name="Arial", size=8)
    def head(ws, row):
        for c in ws[row]:
            if c.value: c.font = bw; c.fill = fillH; c.alignment = Alignment(horizontal="center", wrap_text=True)
    def style(ws, r0, r1, numcols):
        for row in ws.iter_rows(min_row=r0, max_row=r1):
            for c in row: c.font = norm
            for i in numcols: row[i].number_format = "#,##0.00" if i < 90 else "0.0000"

    # --- Sheet 1: Diesel price trend (period x supplier) ---
    ws = wb.active; ws.title = "Diesel trend"
    ws["A1"] = "DIESEL EFFECTIVE NET EUR/L - TREND BY SUPPLIER AND MONTH"; ws["A1"].font = Font(bold=True, size=11, name="Arial")
    sups = [s[0] for s in con.execute("SELECT DISTINCT supplier FROM transactions ORDER BY supplier")]
    ws.append([]); ws.append(["Period"] + sups + ["FLEET eur/L", "FLEET litres", "MoM fleet %"])
    head(ws, 3)
    r = 4
    prev = None
    for p in periods:
        row = [p]
        for s in sups:
            v = con.execute("""SELECT SUM(net_eur_eff)/NULLIF(SUM(qty),0) FROM transactions
                               WHERE period=? AND supplier=? AND product_group='Diesel'""",(p,s)).fetchone()[0]
            row.append(round(v,4) if v else "")
        f = con.execute("""SELECT SUM(net_eur_eff)/NULLIF(SUM(qty),0), SUM(qty) FROM transactions
                           WHERE period=? AND product_group='Diesel'""",(p,)).fetchone()
        row += [round(f[0],4), round(f[1],0), (round((f[0]/prev-1)*100,2) if prev else "")]
        prev = f[0]
        ws.append(row)
        for c in ws[r]: c.font = norm
        for i in range(1, len(sups)+2): ws.cell(row=r, column=i+1).number_format = "0.0000"
        ws.cell(row=r, column=len(sups)+3).number_format = "#,##0"
        r += 1
    ws[f"A{r+1}"] = "MoM columns populate automatically as more periods are loaded. Effective = incl. rebate layers (Q8/Port One)."
    ws[f"A{r+1}"].font = it8
    if len(periods) >= 2:
        chart = LineChart(); chart.title = "Fleet diesel EUR/L trend"; chart.height = 7; chart.width = 16
        last_row = 3 + len(periods)
        data = Reference(ws, min_col=len(sups)+2, min_row=3, max_row=last_row)
        cats = Reference(ws, min_col=1, min_row=4, max_row=last_row)
        chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
        ws.add_chart(chart, f"A{r+4}")
    for i in range(len(sups)+4): ws.column_dimensions[chr(65+i)].width = 11

    # --- Sheet 2: Supplier x country x product history ---
    ws = wb.create_sheet("Supplier-country history")
    ws.append(["Period","Supplier","Country","Product group","Litres/qty","Net EUR","VAT EUR","EUR/L doc","EUR/L eff"])
    head(ws, 1)
    r = 2
    for row in con.execute("""SELECT period, supplier, country, product_group, litres, net_eur, vat_eur,
                              eur_l_doc, eur_l_eff FROM v_supplier_month ORDER BY period, supplier, country, product_group"""):
        ws.append(list(row)); r += 1
    style(ws, 2, r-1, [])
    for rr in ws.iter_rows(min_row=2, max_row=r-1):
        for i in (4,5,6): rr[i].number_format = "#,##0.00"
        for i in (7,8): rr[i].number_format = "0.0000"
    for col, w in zip("ABCDEFGHI",[10,10,10,13,11,11,10,10,10]): ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:I{r-1}"; ws.freeze_panes = "A2"

    # --- Sheet 3: Entity & VAT history ---
    ws = wb.create_sheet("Entity-VAT history")
    ws.append(["Period","Entity","Country","Net EUR","VAT EUR (reclaimable)","Gross EUR"]); head(ws, 1)
    r = 2
    for row in con.execute("SELECT * FROM v_entity_month ORDER BY period, entity, country"):
        ws.append(list(row)); r += 1
    for rr in ws.iter_rows(min_row=2, max_row=r-1):
        for c in rr: c.font = norm
        for i in (3,4,5): rr[i].number_format = "#,##0.00"
    for col, w in zip("ABCDEF",[10,26,10,12,16,12]): ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:F{r-1}"; ws.freeze_panes = "A2"

    # --- Sheet 4: Station history (diesel, >=300 L/month) ---
    ws = wb.create_sheet("Station history")
    ws.append(["Period","Supplier","Country","Station","Litres","EUR/L eff"]); head(ws, 1)
    r = 2
    for row in con.execute("SELECT * FROM v_station_month WHERE litres>=300 ORDER BY period, eur_l_eff"):
        ws.append(list(row)); r += 1
    for rr in ws.iter_rows(min_row=2, max_row=r-1):
        for c in rr: c.font = norm
        rr[4].number_format = "#,##0"; rr[5].number_format = "0.0000"
    for col, w in zip("ABCDEF",[10,10,10,30,10,10]): ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:F{r-1}"; ws.freeze_panes = "A2"

    # --- Sheet 5: How to query ---
    ws = wb.create_sheet("DB guide")
    ws["A1"] = "FUEL HISTORY DATABASE - fuel_history.db (SQLite)"; ws["A1"].font = Font(bold=True, size=11, name="Arial")
    guide = [
    ("What it is","Single-file SQL database holding every canonical transaction line for every loaded month. Open with any SQLite tool (DB Browser for SQLite, DBeaver, Excel via ODBC, Python)."),
    ("Monthly load","After consolidate.py PASSes: python3 history.py - replaces the current period and regenerates this report. Safe to re-run."),
    ("Tables/views","transactions (line level); v_supplier_month, v_entity_month, v_station_month (reporting views this report is built from)."),
    ("Example: supplier trend","SELECT period, supplier, ROUND(SUM(net_eur_eff)/SUM(qty),4) FROM transactions WHERE product_group='Diesel' GROUP BY period, supplier;"),
    ("Example: station drift","SELECT period, eur_l_eff FROM v_station_month WHERE station LIKE '%Poperinge%' ORDER BY period;"),
    ("Example: vehicle audit","SELECT * FROM transactions WHERE vehicle LIKE '%MIJ641%' ORDER BY date;"),
    ("Example: VAT claim pack","SELECT entity, country, period, SUM(vat_eur) FROM transactions GROUP BY entity, country, period;"),
    ]
    r = 3
    for k, v in guide:
        ws[f"A{r}"] = k; ws[f"A{r}"].font = b10
        ws[f"B{r}"] = v; ws[f"B{r}"].font = norm; ws[f"B{r}"].alignment = Alignment(wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 110

    out_path = f"{WORKDIR}/Fleet_Fuel_History_Report.xlsx"
    wb.save(out_path)
    print("history report saved")
    con.close()
    return out_path


def main():
    load()


if __name__ == "__main__":
    main()
