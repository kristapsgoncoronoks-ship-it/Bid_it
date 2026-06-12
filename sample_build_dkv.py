import os, sys
WORKDIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, WORKDIR)
from sample_dkv_data import T, CT, A, B
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "Transactions"
hdr = ["Invoice","Vehicle","Date","Time","Brand","City","Product","Qty (L/pc)",
       "Gross SEK/unit","Net SEK/unit","Base net SEK","Discount net SEK","Service fee net SEK",
       "Total net SEK","VAT 25% SEK","Total gross SEK","Payable EUR","Eff. net SEK/L after disc",
       "Chk1: base=qty x net","Chk2: net=base+disc+fee","Chk3: VAT=25%","Chk4: gross=net+VAT"]
ws.append(hdr)
fill = PatternFill("solid", start_color="C8102E")
bw = Font(bold=True, color="FFFFFF", name="Arial", size=9)
norm = Font(name="Arial", size=9)
for c in ws[1]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(horizontal="center", wrap_text=True)

r = 2
for inv, veh, d, t, br, city, prod, qty, gpu, npu, base, disc, fee, net, vat, gross, eur in T:
    ws.append([inv, veh, d, t, br, city, prod, qty, gpu, npu, base, disc,
        fee if fee else None, net, vat, gross, eur,
        f'=IF(G{r}="parking","",(K{r}+L{r})/H{r})',
        f'=IF(ABS(ROUND(H{r}*J{r},2)-K{r})<=0.02,"OK","CHECK")',
        f'=IF(ABS(ROUND(K{r}+L{r}+IF(M{r}="",0,M{r}),2)-N{r})<=0.011,"OK","CHECK")',
        f'=IF(ABS(ROUND(N{r}*0.25,2)-O{r})<=0.011,"OK","CHECK")',
        f'=IF(ABS(N{r}+O{r}-P{r})<=0.011,"OK","CHECK")'])
    r += 1
last = r - 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = norm
for col, f in (("H","#,##0.000"),("I","0.0000"),("J","0.0000"),("K","#,##0.00"),("L","#,##0.00"),
               ("M","#,##0.00"),("N","#,##0.00"),("O","#,##0.00"),("P","#,##0.00"),
               ("Q","#,##0.00"),("R","0.0000")):
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=ord(col)-64, max_col=ord(col)-64):
        row[0].number_format = f
for col, w in zip("ABCDEFGHIJKLMNOPQRSTUV",[12,7,10,6,9,13,9,9,9,9,10,10,9,10,9,10,9,10,10,11,9,10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:V{last}"

# Card totals check
ws2 = wb.create_sheet("Card totals check")
ws2.append(["Invoice","Vehicle","Stated: qty","base","disc","fee","net","VAT","gross","EUR",
            "Calc net","Calc gross","Calc EUR","Check"])
for c in ws2[1]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(wrap_text=True)
r2 = 2
for (inv, veh), (q, base, disc, fee, net, vat, gross, eur) in CT.items():
    b = (f'SUMPRODUCT((Transactions!$A$2:$A${last}=A{r2})*'
         f'(Transactions!$B$2:$B${last}=B{r2})*Transactions!')
    ws2.append([inv, veh, q, base, disc, fee, net, vat, gross, eur,
        f"=ROUND({b}$N$2:$N${last}),2)",
        f"=ROUND({b}$P$2:$P${last}),2)",
        f"=ROUND({b}$Q$2:$Q${last}),2)",
        f'=IF(AND(ABS(K{r2}-G{r2})<=0.02,ABS(L{r2}-I{r2})<=0.02,ABS(M{r2}-J{r2})<=0.02),"OK","CHECK")'])
    r2 += 1
for row in ws2.iter_rows(min_row=2, max_row=r2-1):
    for c in row: c.font = norm
    for c in row[2:13]: c.number_format = "#,##0.00"
for col, w in zip("ABCDEFGHIJKLMN",[12,7,10,10,9,7,10,9,10,9,10,10,9,7]):
    ws2.column_dimensions[col].width = w

# Invoice totals check
ws3 = wb.create_sheet("Invoice totals check")
ws3.append(["Check item","Invoice states","Calculated","Diff","Check"])
for c in ws3[1]:
    c.font = bw; c.fill = fill
def S(inv, prod, col):
    return (f'=ROUND(SUMPRODUCT((Transactions!$A$2:$A${last}="{inv}")*'
            f'(Transactions!$G$2:$G${last}="{prod}")*Transactions!${col}$2:${col}${last}),2)')
def SI(inv, col):
    return f'=ROUND(SUMIF(Transactions!$A$2:$A${last},"{inv}",Transactions!${col}$2:${col}${last}),2)'
items = [
 (A+" DIESEL litres",10407.490,S(A,"DIESEL","H")),(A+" DIESEL net SEK",166997.74,S(A,"DIESEL","N")),
 (A+" ADBLUE litres",354.080,S(A,"ADBLUE","H")),(A+" ADBLUE net SEK",2461.33,S(A,"ADBLUE","N")),
 (A+" parking net SEK",1812.54,S(A,"parking","N")),
 (A+" TOTAL net SEK",171271.61,SI(A,"N")),(A+" TOTAL VAT SEK",42817.90,SI(A,"O")),
 (A+" TOTAL gross SEK",214089.51,SI(A,"P")),(A+" TOTAL payable EUR",19779.26,SI(A,"Q")),
 (B+" DIESEL litres",6924.390,S(B,"DIESEL","H")),(B+" DIESEL net SEK",109108.85,S(B,"DIESEL","N")),
 (B+" ADBLUE litres",400.230,S(B,"ADBLUE","H")),(B+" ADBLUE net SEK",2678.35,S(B,"ADBLUE","N")),
 (B+" HVO 100 litres",251.000,S(B,"HVO 100","H")),(B+" HVO 100 net SEK",5732.34,S(B,"HVO 100","N")),
 (B+" parking net SEK",1144.99,S(B,"parking","N")),
 (B+" TOTAL net SEK",118664.53,SI(B,"N")),(B+" TOTAL VAT SEK",29666.13,SI(B,"O")),
 (B+" TOTAL gross SEK",148330.66,SI(B,"P")),(B+" TOTAL payable EUR",13757.93,SI(B,"Q")),
 ("BOTH invoices gross SEK",362420.17,f"=ROUND(SUM(Transactions!$P$2:$P${last}),2)"),
 ("BOTH invoices payable EUR",33537.19,f"=ROUND(SUM(Transactions!$Q$2:$Q${last}),2)"),
 ("BOTH Swedish VAT SEK (reclaimable)",72484.03,f"=ROUND(SUM(Transactions!$O$2:$O${last}),2)"),
]
r3 = 2
for name, doc, calc in items:
    ws3.append([name, doc, calc, f"=C{r3}-B{r3}", f'=IF(ABS(D{r3})<=0.02,"OK","CHECK")'])
    r3 += 1
for row in ws3.iter_rows(min_row=2, max_row=r3-1):
    for c in row: c.font = norm
    for c in row[1:4]: c.number_format = "#,##0.00"
for col, w in zip("ABCDE",[40,14,14,9,8]):
    ws3.column_dimensions[col].width = w

# Price summary
ws4 = wb.create_sheet("Price summary", 0)
ws4.append(["DKV SWEDEN MAY 2026 - JUPITER PLUS AS - NET PRICE OVERVIEW (SEK excl. 25% VAT, after DKV discount)"])
ws4["A1"].font = Font(bold=True, name="Arial", size=11)
ws4.append(["Invoices 26/651689595/011 (1-15 May, due per DKV terms) + 26/652169828/011 (16-31 May); payable in EUR"])
ws4["A2"].font = Font(italic=True, name="Arial", size=8)
ws4.append([])
ws4.append(["Diesel by station (city)","Litres","Net SEK after disc","Eff. net SEK/L","approx EUR/L @10.80","Discount SEK/L"])
for c in ws4[4]:
    c.font = bw; c.fill = fill; c.alignment = Alignment(wrap_text=True, horizontal="center")
cities = ["TANUMSHEDE","FALKENBERG","TRELLEBORG","HELSINGBORG","BASTAD","MARKARYD","ARLOV",
          "LANDSKRONA","GOTEBORG","HALMSTAD","LILLA EDET","ANGELHOLM","ORKELLJUNGA","KVANUM",
          "MALMO","ODESHOG","GRANNA","VAXJO","HISINGS BACKA"]
r4 = 5
for city in cities:
    b = (f'SUMPRODUCT((Transactions!$F$2:$F${last}=A{r4})*'
         f'(Transactions!$G$2:$G${last}="DIESEL")*Transactions!')
    ws4.append([city, f"={b}$H$2:$H${last})", f"=ROUND({b}$N$2:$N${last}),2)",
        f"=C{r4}/B{r4}", f"=D{r4}/10.80",
        f"=-ROUND({b}$L$2:$L${last}),2)/B{r4}"])
    r4 += 1
tr = r4
ws4.append(["TOTAL DIESEL", f"=SUM(B5:B{r4-1})", f"=SUM(C5:C{r4-1})", f"=C{tr}/B{tr}", f"=D{tr}/10.80",
            f'=-SUMPRODUCT((Transactions!$G$2:$G${last}="DIESEL")*Transactions!$L$2:$L${last})/B{tr}'])
r4 += 1
ws4.append([])
r4 += 1
ws4.append(["By half-month (diesel)","Litres","Net SEK","Eff. net SEK/L","approx EUR/L"])
for c in ws4[r4]:
    if c.value: c.font = Font(bold=True, name="Arial", size=10)
r4 += 1
for inv, lbl in ((A,"01-15 May"),(B,"16-31 May")):
    b = (f'SUMPRODUCT((Transactions!$A$2:$A${last}="{inv}")*'
         f'(Transactions!$G$2:$G${last}="DIESEL")*Transactions!')
    ws4.append([lbl, f"={b}$H$2:$H${last})", f"=ROUND({b}$N$2:$N${last}),2)", f"=C{r4}/B{r4}", f"=D{r4}/10.80"])
    r4 += 1
ws4.append([])
r4 += 1
ws4.append(["Other products","Qty","Net SEK","Eff. net SEK/unit","approx EUR/unit"])
for c in ws4[r4]:
    if c.value: c.font = Font(bold=True, name="Arial", size=10)
r4 += 1
for prod, lbl in (("ADBLUE","AdBlue (bulk), L"),("HVO 100","HVO 100 renewable diesel, L"),("parking","Parking services, pc")):
    b = f'SUMPRODUCT((Transactions!$G$2:$G${last}="{prod}")*Transactions!'
    ws4.append([lbl, f"={b}$H$2:$H${last})", f"=ROUND({b}$N$2:$N${last}),2)", f"=C{r4}/B{r4}", f"=D{r4}/10.80"])
    r4 += 1
for row in ws4.iter_rows(min_row=5, max_row=r4-1):
    for c in row:
        if c.font is None or not c.font.bold: c.font = norm
    if row[1].value: row[1].number_format = "#,##0.00"
    if row[2].value: row[2].number_format = "#,##0.00"
    for c in row[3:6]: c.number_format = "0.0000"
for c in ws4[tr]: c.font = Font(bold=True, name="Arial", size=10)
ws4[f"A{r4+1}"] = ("EUR/L shown at indicative 10.80 SEK/EUR (DKV converts per transaction date; implied invoice rates "
                   "10.82 and 10.78). Swedish VAT 25% is charged on top and reclaimable via EE registration. Pump net "
                   "prices fell 17.60 -> 16.856 -> 17.16 SEK/L over the month; DKV discount averages ~1.3 SEK/L net.")
ws4[f"A{r4+1}"].font = Font(italic=True, name="Arial", size=8)
for col, w in zip("ABCDEF",[28,11,15,13,14,13]):
    ws4.column_dimensions[col].width = w

wb.save(os.path.join(WORKDIR, "DKV_SE_May2026_transactions.xlsx"))
print("saved", last-1)
