from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = load_workbook("/home/claude/work/Q8_payment_summary_DE00752298_full_transactions.xlsx")
ws = wb["Transactions"]
last = ws.max_row
fill = PatternFill("solid", start_color="1F4E78")
bold_w = Font(bold=True, color="FFFFFF", name="Arial", size=10)
norm = Font(name="Arial", size=10)
blue = Font(name="Arial", size=10, color="0000FF")

# --- Sheet: Rebate allocation (per country) ---
wsr = wb.create_sheet("Rebate allocation")
wsr.append(["Country","Port One rebate EUR (input)","Litres (from transactions)",
            "Rebate EUR per litre","Q8 net EUR","Adjusted net EUR","Eff. avg net price EUR/L"])
for c in wsr[1]:
    c.font = bold_w; c.fill = fill; c.alignment = Alignment(horizontal="center", wrap_text=True)
rebates = [("Belgium",8027.61),("Germany",1133.14),("Denmark",1609.10),("Spain",1096.15),
           ("France",398.70),("Poland",15.00),("Austria",81.12),("Luxembourg",7.01)]
r = 2
for ctry, reb in rebates:
    wsr.append([ctry, reb,
        f'=SUMIF(Transactions!$F$2:$F${last},A{r},Transactions!$L$2:$L${last})',
        f'=IF(C{r}>0,B{r}/C{r},0)',
        f'=SUMIF(Transactions!$F$2:$F${last},A{r},Transactions!$M$2:$M${last})',
        f'=E{r}-B{r}', f'=IF(C{r}>0,F{r}/C{r},"n/a")'])
    wsr[f"B{r}"].font = blue
    r += 1
tr = r
wsr.append(["TOTAL", f"=SUM(B2:B{r-1})", f"=SUM(C2:C{r-1})", "",
            f"=SUM(E2:E{r-1})", f"=SUM(F2:F{r-1})", f"=F{tr}/C{tr}"])
for c in wsr[tr]: c.font = Font(bold=True, name="Arial", size=10)
for row in wsr.iter_rows(min_row=2, max_row=tr):
    for c in row:
        if c.font != blue and c.row != tr: c.font = norm
    row[1].number_format = "#,##0.00"; row[2].number_format = "#,##0.00"
    row[3].number_format = "0.0000"; row[4].number_format = "#,##0.00"
    row[5].number_format = "#,##0.00"; row[6].number_format = "0.0000"
wsr["B2"].font = blue
for ctry_row in range(2, tr): wsr[f"B{ctry_row}"].font = blue
wsr[f"A{tr+2}"] = ("Rebate inputs (blue) from Port One invoice EE2605310167. Allocation assumption: each country's "
                   "rebate is spread evenly over ALL litres bought in that country (diesel + AdBlue). "
                   "Luxembourg shows EUR 7.01 rebate but zero transactions in this summary - likely relates to a "
                   "prior period or a fee adjustment; left unallocated. Ask Port One to confirm the per-litre "
                   "rebate split (diesel vs AdBlue) if exact per-product pricing is needed.")
wsr[f"A{tr+2}"].font = Font(italic=True, name="Arial", size=9)
for col, w in zip("ABCDEFG",[13,22,21,15,13,15,17]):
    wsr.column_dimensions[col].width = w

# --- Extend Transactions sheet with adjusted pricing columns ---
ws["Q1"] = "Rebate EUR/L (country)"; ws["R1"] = "Adjusted net price EUR/L"
ws["S1"] = "Adjusted net amount EUR"; ws["T1"] = "Discount %"
for col in ("Q1","R1","S1","T1"):
    ws[col].font = bold_w; ws[col].fill = fill
    ws[col].alignment = Alignment(horizontal="center", wrap_text=True)
for rr in range(2, last+1):
    ws[f"Q{rr}"] = f"=IFERROR(VLOOKUP(F{rr},'Rebate allocation'!$A$2:$D$9,4,FALSE),0)"
    ws[f"R{rr}"] = f"=K{rr}-Q{rr}"
    ws[f"S{rr}"] = f"=ROUND(R{rr}*L{rr},2)"
    ws[f"T{rr}"] = f"=Q{rr}/K{rr}"
    for col, f in (("Q","0.0000"),("R","0.0000"),("S","#,##0.00"),("T","0.0%")):
        ws[f"{col}{rr}"].number_format = f
        ws[f"{col}{rr}"].font = norm
ws.auto_filter.ref = f"A1:T{last}"
for col, w in zip("QRST",[12,13,14,9]):
    ws.column_dimensions[col].width = w

# --- Sheet: Effective cost summary ---
wse = wb.create_sheet("Effective cost", 0)
rows = [
 ("EFFECTIVE FUEL COST - PERIOD 16-30 MAY 2026 (Summary DE00752298 + Port One EE2605310167)", None, None),
 ("", None, None),
 ("Item","EUR","Note"),
 ("Q8 net amount (all transactions)", f"=SUM(Transactions!$M$2:$M${last})", "56,057.99 per Transaction Details"),
 ("VAT charged by Q8 (mixed rates)", 11169.04, "Reclaimable via VAT refund (invoices marked for VAT reclaim)"),
 ("Q8 gross (payment summary)", "=B4+B5", "67,227.03"),
 ("Port One rebates", "=-'Rebate allocation'!B10", "8 countries, per invoice EE2605310167"),
 ("Port One invoice payable (due 14.06.2026)", "=B6+B7", "54,859.20 - note due date is 14.06, one day before Q8 summary due date"),
 ("", None, None),
 ("Cash cost now", "=B8", ""),
 ("Less: VAT reclaimable", "=-B5", "DE 19%, BE 21%, FR 20%, DK 25%, PL 23%, ES 21%/10%, AT 20%"),
 ("True net fuel cost after VAT recovery", "=B10+B11", ""),
 ("Total litres", f"=SUM(Transactions!$L$2:$L${last})", "32,523.14 L"),
 ("Effective blended cost EUR/L (after rebates & VAT recovery)", "=B12/B13", ""),
 ("Q8 list net price avg EUR/L (before rebates)", "=B4/B13", ""),
 ("Average saving from rebates EUR/L", "=B15-B14", ""),
]
r = 1
for a, b, c in rows:
    wse.append([a, b, c]); r += 1
wse["A1"].font = Font(bold=True, name="Arial", size=11)
for cell in ("A3","B3","C3"):
    wse[cell].font = bold_w; wse[cell].fill = fill
for rr in range(4, 17):
    wse[f"A{rr}"].font = norm; wse[f"C{rr}"].font = Font(italic=True, name="Arial", size=9)
    wse[f"B{rr}"].font = norm
    wse[f"B{rr}"].number_format = "#,##0.00" if rr not in (14,15,16) else "0.0000"
wse["B5"].font = blue
for rr in (8,12,14):
    wse[f"A{rr}"].font = Font(bold=True, name="Arial", size=10)
    wse[f"B{rr}"].font = Font(bold=True, name="Arial", size=10)
wse.column_dimensions["A"].width = 50
wse.column_dimensions["B"].width = 14
wse.column_dimensions["C"].width = 60

wb.save("/home/claude/work/Q8_DE00752298_adjusted_pricing.xlsx")
print("ok")
