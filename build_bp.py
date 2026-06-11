import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

WORKDIR = os.path.dirname(os.path.abspath(__file__))

# (lp, card, date, reg, location, product, category, qty, unit_net, gross, vat%, vat, net)
T = [
(1,"000080","27/05/26","BE9919","SUWALKI","ON ACT.","Diesel",300.00,5.59,1812.13,8,134.23,1677.90),
(2,"000122","15/05/26","GH9909","SUWALKI","ADBLUE","AdBlue",20.02,2.16,53.11,23,9.93,43.18),
(3,"000122","15/05/26","GH9909","SUWALKI","ON ACT.","Diesel",256.63,5.82,1612.52,8,119.45,1493.07),
(4,"000122","26/05/26","GH9909","SUWALKI","ADBLUE","AdBlue",20.02,2.16,53.11,23,9.93,43.18),
(5,"000122","28/05/26","GH9909","SUWALKI","ADBLUE","AdBlue",40.07,2.16,106.31,23,19.88,86.43),
(6,"000122","28/05/26","GH9909","SUWALKI","ON ACT.","Diesel",110.12,5.40,642.10,8,47.56,594.54),
(7,"000130","29/05/26","GR9","SUWALKI","ON ACT.","Diesel",400.04,5.38,2323.09,8,172.08,2151.01),
(8,"000171","23/05/26","KP8998","SUWALKI","ON ACT.","Diesel",800.06,5.85,5055.64,8,374.49,4681.15),
(9,"000171","28/05/26","KP8998","SUWALKI","ON ACT.","Diesel",568.72,5.40,3316.16,8,245.64,3070.52),
(10,"000189","20/05/26","KS9339","RADZYMIN CIEMNE","ON ACT.","Diesel",550.00,5.92,3514.10,8,260.30,3253.80),
(11,"000247","23/05/26","MV5097","SUWALKI","ON ACT.","Diesel",723.73,5.85,4573.30,8,338.76,4234.54),
(12,"000247","23/05/26","MV5097","SUWALKI","ON ACT.","Diesel",366.56,5.85,2316.32,8,171.58,2144.74),
(13,"000254","16/05/26","NI3993","SUWALKI","ON ACT.","Diesel",620.00,5.82,3895.73,8,288.57,3607.16),
(14,"000270","31/05/26","NI9449","SUWALKI","ON ACT.","Diesel",770.00,5.38,4471.51,8,331.22,4140.29),
(15,"000288","22/05/26","NJ959","SUWALKI","ON ACT.","Diesel",670.01,5.85,4233.85,8,313.62,3920.23),
(16,"000296","16/05/26","NL1991","AWSA A2 - PPO L","Toll Service","Toll",1.00,130.08,160.00,23,29.92,130.08),
(17,"000296","16/05/26","NL1991","AWSA A2 - PPO L","Oplata ORS","Toll fee",1.00,3.25,4.00,23,0.75,3.25),
(18,"000296","16/05/26","NL1991","AWSA A2 - PPO N","Toll Service","Toll",1.00,130.08,160.00,23,29.92,130.08),
(19,"000296","16/05/26","NL1991","AWSA A2 - PPO N","Oplata ORS","Toll fee",1.00,3.25,4.00,23,0.75,3.25),
(20,"000296","17/05/26","NL1991","AWSA A2 - SPO T","Toll Service","Toll",1.00,156.91,193.00,23,36.09,156.91),
(21,"000296","17/05/26","NL1991","AWSA A2 - SPO T","Oplata ORS","Toll fee",1.00,3.92,4.82,23,0.90,3.92),
(22,"000296","17/05/26","NL1991","AWSA A2 - PPO T","Toll Service","Toll",1.00,5.69,7.00,23,1.31,5.69),
(23,"000296","17/05/26","NL1991","AWSA A2 - PPO T","Oplata ORS","Toll fee",1.00,0.14,0.17,23,0.03,0.14),
(24,"000296","20/05/26","NL1991","AWSA A2 - PPO G","Toll Service","Toll",1.00,156.91,193.00,23,36.09,156.91),
(25,"000296","20/05/26","NL1991","AWSA A2 - PPO G","Oplata ORS","Toll fee",1.00,3.92,4.82,23,0.90,3.92),
(26,"000296","20/05/26","NL1991","AWSA A2 - SPO T","Toll Service","Toll",1.00,5.69,7.00,23,1.31,5.69),
(27,"000296","20/05/26","NL1991","AWSA A2 - SPO T","Oplata ORS","Toll fee",1.00,0.14,0.17,23,0.03,0.14),
(28,"000296","21/05/26","NL1991","AWSA A2 - PPO L","Toll Service","Toll",1.00,130.08,160.00,23,29.92,130.08),
(29,"000296","21/05/26","NL1991","AWSA A2 - PPO L","Oplata ORS","Toll fee",1.00,3.25,4.00,23,0.75,3.25),
(30,"000296","21/05/26","NL1991","AWSA A2 - PPO N","Toll Service","Toll",1.00,130.08,160.00,23,29.92,130.08),
(31,"000296","21/05/26","NL1991","AWSA A2 - PPO N","Oplata ORS","Toll fee",1.00,3.25,4.00,23,0.75,3.25),
(32,"000312","23/05/26","NO4994","SUWALKI","ADBLUE","AdBlue",43.45,2.16,115.28,23,21.56,93.72),
(33,"000312","23/05/26","NO4994","SUWALKI","ON ACT.","Diesel",290.80,5.85,1837.59,8,136.12,1701.47),
(34,"000312","23/05/26","NO4994","SUWALKI","ON ACT.","Diesel",144.79,5.85,914.93,8,67.77,847.16),
(35,"000346","18/05/26","RA99","SUWALKI","ON ACT.","Diesel",350.00,5.85,2210.17,8,163.72,2046.45),
(36,"000346","23/05/26","RA99","SUWALKI","ON ACT.","Diesel",150.00,5.85,947.86,8,70.21,877.65),
(37,"000346","28/05/26","RA99","SUWALKI","ON ACT.","Diesel",490.00,5.40,2857.15,8,211.64,2645.51),
(38,"000361","28/05/26","NI8998","SUWALKI","ON ACT.","Diesel",565.00,5.40,3294.47,8,244.03,3050.44),
(39,"000379","17/05/26","NR6247","SUWALKI","ON ACT.","Diesel",512.50,5.82,3220.27,8,238.54,2981.73),
(40,"000403","17/05/26","HC994","SUWALKI","ON ACT.","Diesel",770.01,5.82,4838.31,8,358.39,4479.92),
(41,"000411","26/05/26","HZ995","SUWALKI","ON ACT.","Diesel",250.03,5.75,1552.68,8,115.01,1437.67),
(42,"000445","25/05/26","OMUSS","SUWALKI","ADBLUE","AdBlue",20.00,2.16,53.06,23,9.92,43.14),
(43,"000445","31/05/26","OMUSS","SUWALKI","ON ACT.","Diesel",485.00,5.38,2816.48,8,208.63,2607.85),
(44,"000445","31/05/26","OMUSS","SUWALKI","ON ACT.","Diesel",435.00,5.38,2526.12,8,187.12,2339.00),
(45,"000460","21/05/26","LB929","SUWALKI","ON ACT.","Diesel",664.86,5.91,4240.80,8,314.13,3926.67),
(46,"000494","30/05/26","ON989","SUWALKI","ON ACT.","Diesel",300.53,5.38,1745.23,8,129.28,1615.95),
]
card_totals = {"000080":1812.13,"000122":2467.15,"000130":2323.09,"000171":8371.80,
"000189":3514.10,"000247":6889.62,"000254":3895.73,"000270":4471.51,"000288":4233.85,
"000296":1065.98,"000312":2867.80,"000346":6015.18,"000361":3294.47,"000379":3220.27,
"000403":4838.31,"000411":1552.68,"000445":5395.66,"000460":4240.80,"000494":1745.23}

wb = Workbook()
ws = wb.active
ws.title = "Transactions"
hdr = ["Lp","Card","Date","Vehicle reg","Location","Product","Category","Qty (L/pcs)",
       "Unit net PLN (doc)","Gross PLN","VAT %","VAT PLN","Net PLN",
       "Check1: Net+VAT=Gross","Check2: implied unit vs doc","Check3: VAT = Net x rate"]
ws.append(hdr)
fill = PatternFill("solid", start_color="00553E")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    c.fill = fill; c.alignment = Alignment(horizontal="center", wrap_text=True)
norm = Font(name="Arial", size=10)
r = 2
for row in T:
    ws.append(list(row) + [
        f'=IF(ABS(M{r}+L{r}-J{r})<=0.01,"OK","CHECK")',
        f'=IF(ABS(M{r}/H{r}-I{r})<=0.005+0.01/H{r},"OK","CHECK")',
        f'=IF(ABS(L{r}-ROUND(M{r}*K{r}/100,2))<=0.02,"OK","CHECK")'])
    r += 1
last = r - 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = norm
for col, f in (("H","#,##0.00"),("I","0.00"),("J","#,##0.00"),("L","#,##0.00"),("M","#,##0.00")):
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=ord(col)-64, max_col=ord(col)-64):
        row[0].number_format = f
for col, w in zip("ABCDEFGHIJKLMNOP",[5,9,10,11,17,12,9,11,11,11,7,10,11,11,12,11]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:P{last}"

# Card totals
ws2 = wb.create_sheet("Card totals check")
ws2.append(["Card","Stated gross PLN","Calc gross","Diff","Check"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10); c.fill = fill
r2 = 2
for card, tot in card_totals.items():
    ws2.append([card, tot, f'=ROUND(SUMIF(Transactions!$B$2:$B${last},A{r2},Transactions!$J$2:$J${last}),2)',
                f"=C{r2}-B{r2}", f'=IF(ABS(D{r2})<=0.01,"OK","CHECK")'])
    r2 += 1
ws2.append(["TOTAL", f"=SUM(B2:B{r2-1})", f"=SUM(C2:C{r2-1})", f"=C{r2}-B{r2}",
            f'=IF(ABS(D{r2})<=0.01,"OK","CHECK")'])
last2 = r2
for row in ws2.iter_rows(min_row=2, max_row=last2):
    for c in row: c.font = norm
    for c in row[1:4]: c.number_format = "#,##0.00"
for c in ws2[last2]: c.font = Font(bold=True, name="Arial", size=10)
for col, w in zip("ABCDE",[9,17,14,10,8]):
    ws2.column_dimensions[col].width = w

# Invoice summary checks
ws3 = wb.create_sheet("Invoice totals check")
ws3.append(["Check item","Invoice states","Calculated","Diff","Check"])
for c in ws3[1]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10); c.fill = fill
items = [
 ("Diesel (OLEJ NAPEDOWY) litres", 11544.39,
  f'=ROUND(SUMIF(Transactions!$G$2:$G${last},"Diesel",Transactions!$H$2:$H${last}),2)', 0.01),
 ("Diesel gross PLN", 70768.51,
  f'=ROUND(SUMIF(Transactions!$G$2:$G${last},"Diesel",Transactions!$J$2:$J${last}),2)', 0.01),
 ("Diesel net PLN", 65526.42,
  f'=ROUND(SUMIF(Transactions!$G$2:$G${last},"Diesel",Transactions!$M$2:$M${last}),2)', 0.01),
 ("AdBlue (AKCESORIA SAM.) litres", 143.56,
  f'=ROUND(SUMIF(Transactions!$G$2:$G${last},"AdBlue",Transactions!$H$2:$H${last}),2)', 0.01),
 ("AdBlue gross PLN", 380.87,
  f'=ROUND(SUMIF(Transactions!$G$2:$G${last},"AdBlue",Transactions!$J$2:$J${last}),2)', 0.01),
 ("Tolls (INNE PRODUKTY) item count", 16.00,
  f'=SUMPRODUCT((LEFT(Transactions!$G$2:$G${last},4)="Toll")*Transactions!$H$2:$H${last})', 0.01),
 ("Tolls gross PLN", 1065.98,
  f'=ROUND(SUMPRODUCT((LEFT(Transactions!$G$2:$G${last},4)="Toll")*Transactions!$J$2:$J${last}),2)', 0.01),
 ("VAT 8% base (net) PLN", 65526.42,
  f'=ROUND(SUMIF(Transactions!$K$2:$K${last},8,Transactions!$M$2:$M${last}),2)', 0.01),
 ("VAT 8% amount PLN", 5242.09,
  f'=ROUND(SUMIF(Transactions!$K$2:$K${last},8,Transactions!$L$2:$L${last}),2)', 0.02),
 ("VAT 23% base (net) PLN", 1176.29,
  f'=ROUND(SUMIF(Transactions!$K$2:$K${last},23,Transactions!$M$2:$M${last}),2)', 0.01),
 ("VAT 23% amount PLN", 270.56,
  f'=ROUND(SUMIF(Transactions!$K$2:$K${last},23,Transactions!$L$2:$L${last}),2)', 0.02),
 ("TOTAL gross PLN (amount payable)", 72215.36,
  f'=ROUND(SUM(Transactions!$J$2:$J${last}),2)', 0.01),
 ("TOTAL VAT PLN", 5512.65, f'=ROUND(SUM(Transactions!$L$2:$L${last}),2)', 0.02),
 ("TOTAL net PLN", 66702.71, f'=ROUND(SUM(Transactions!$M$2:$M${last}),2)', 0.01),
]
r3 = 2
for name, doc, calc, tol in items:
    ws3.append([name, doc, calc, f"=C{r3}-B{r3}", f'=IF(ABS(D{r3})<={tol},"OK","CHECK")'])
    r3 += 1
for row in ws3.iter_rows(min_row=2, max_row=r3-1):
    for c in row: c.font = norm
    for c in row[1:4]: c.number_format = "#,##0.00"
for col, w in zip("ABCDE",[33,15,15,10,8]):
    ws3.column_dimensions[col].width = w

# Price analysis
ws4 = wb.create_sheet("Price analysis")
ws4.append(["Diesel net unit price by date (PLN/L, rebated per invoice note)","","",""])
ws4.append(["Date","Litres","Net PLN","Avg net PLN/L"])
for c in ws4[2]:
    c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10); c.fill = fill
dates = sorted({t[2] for t in T if t[6]=="Diesel"}, key=lambda d: d.split("/")[0])
r4 = 3
for d in dates:
    ws4.append([d,
      f'=SUMPRODUCT((Transactions!$C$2:$C${last}=A{r4})*(Transactions!$G$2:$G${last}="Diesel")*Transactions!$H$2:$H${last})',
      f'=SUMPRODUCT((Transactions!$C$2:$C${last}=A{r4})*(Transactions!$G$2:$G${last}="Diesel")*Transactions!$M$2:$M${last})',
      f"=C{r4}/B{r4}"])
    r4 += 1
ws4.append(["TOTAL", f"=SUM(B3:B{r4-1})", f"=SUM(C3:C{r4-1})", f"=C{r4}/B{r4}"])
for row in ws4.iter_rows(min_row=3, max_row=r4):
    for c in row: c.font = norm
    row[1].number_format = "#,##0.00"; row[2].number_format = "#,##0.00"; row[3].number_format = "0.0000"
for c in ws4[r4]: c.font = Font(bold=True, name="Arial", size=10)
ws4["A1"].font = Font(bold=True, name="Arial", size=10)
for col, w in zip("ABCD",[12,12,13,13]):
    ws4.column_dimensions[col].width = w

wb.save(os.path.join(WORKDIR, "BP_PL_0261167596_transactions.xlsx"))
print("lines:", last-1)
