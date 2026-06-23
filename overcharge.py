"""overcharge.py — SUPPLIER OVERCHARGE CLAIM-BACK: turn the contract-audit €-overcharges into
an actionable supplier claim.

`contract_audit.audit()` DETECTS overcharges (short discounts / over-ceiling prices) per
(supplier, country, station, product, period). This module gives each supplier×period an
evidence packet and a claim-back LIFECYCLE so the recovered cash is tracked toward the
north-star metric "€ overcharges recovered":

    detected → packaged → claimed → recovered (€ credited) / rejected / written_off

App-owned store (`overcharge.db`, gitignored), audited writes, tenant-stamped. READ-ONLY over
the analytics via `contract_audit`; never mutates a transaction or a VAT figure. Never raises
into the caller -> safe defaults.
"""
import io
import os
import sqlite3

import applog
import audit
import db_migrate
import money
import tenancy

WORKDIR = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("OVERCHARGE_DB", f"{WORKDIR}/overcharge.db")
log = applog.get("overcharge")

# Claim-back lifecycle. OPEN_STATUSES are still being chased; RECOVERED is the only one that
# books cash; REJECTED / WRITTEN_OFF close it out with no recovery.
STATUSES = ("detected", "packaged", "claimed", "recovered", "rejected", "written_off")
OPEN_STATUSES = ("detected", "packaged", "claimed")
STATUS_LABELS = {
    "detected": "Detected", "packaged": "Evidence packaged", "claimed": "Claimed to supplier",
    "recovered": "Recovered", "rejected": "Rejected by supplier", "written_off": "Written off"}

_MIGRATIONS = [
    """CREATE TABLE IF NOT EXISTS overcharge_claims (
        id            INTEGER PRIMARY KEY AUTOINCREMENT,
        supplier      TEXT NOT NULL,
        period        TEXT NOT NULL,
        detected_eur  REAL,                         -- € overcharge snapshot at open
        status        TEXT NOT NULL DEFAULT 'detected',
        recovered_eur REAL,
        note          TEXT,
        opened_at     TEXT DEFAULT CURRENT_TIMESTAMP,
        opened_by     TEXT,
        updated_at    TEXT,
        updated_by    TEXT,
        tenant_id     TEXT NOT NULL DEFAULT 'default',
        UNIQUE (tenant_id, supplier, period))""",
]
_READY = set()


def connect():
    con = sqlite3.connect(DB, timeout=30)
    con.row_factory = sqlite3.Row
    if DB == ":memory:" or DB not in _READY:
        db_migrate.apply(con, "overcharge", _MIGRATIONS)
        if DB != ":memory:":
            _READY.add(DB)
    return con


def _claim_row(con, supplier, period):
    frag, tp = tenancy.scope_clause()
    return con.execute("SELECT * FROM overcharge_claims WHERE supplier=? AND period=?" + frag,
                       [supplier, period, *tp]).fetchone()


def ensure_claim(supplier, period, detected_eur=None, actor=None):
    """Make sure a claim row exists for (supplier, period); snapshot the detected € on create.
    Idempotent; never raises."""
    sup, per = (supplier or "").strip(), (period or "").strip()
    if not sup or not per:
        return
    try:
        con = connect()
        if actor:
            audit.set_actor(con, actor)
        try:
            if _claim_row(con, sup, per):
                return
            con.execute(
                "INSERT INTO overcharge_claims (supplier, period, detected_eur, status, "
                "opened_by, tenant_id) VALUES (?,?,?, 'detected', ?, ?)",
                (sup, per, float(money.f2(detected_eur or 0)), actor or "",
                 tenancy.write_tenant()))
            con.commit()
        finally:
            con.close()
    except Exception as e:
        log.warning("ensure_claim(%s,%s) failed: %s", supplier, period, e)


def set_status(supplier, period, status, actor=None, detected_eur=None):
    """Advance a supplier×period claim to `status`. Returns (True,"") or (False,error)."""
    if status not in STATUSES:
        return False, f"unknown status {status!r}"
    sup, per = (supplier or "").strip(), (period or "").strip()
    if not sup or not per:
        return False, "supplier and period are required"
    ensure_claim(sup, per, detected_eur, actor)
    try:
        con = connect()
        if actor:
            audit.set_actor(con, actor)
        try:
            frag, tp = tenancy.scope_clause()
            con.execute("UPDATE overcharge_claims SET status=?, updated_at=datetime('now'), "
                        "updated_by=? WHERE supplier=? AND period=?" + frag,
                        [status, actor or "", sup, per, *tp])
            con.commit()
            return True, ""
        finally:
            con.close()
    except Exception as e:
        log.warning("set_status(%s,%s,%s) failed: %s", supplier, period, status, e)
        return False, f"could not update the claim ({str(e)[:60]})"


def record_recovery(supplier, period, eur, actor=None, detected_eur=None):
    """Book recovered cash on a claim (status -> recovered). Returns (True,"") or (False,error)."""
    try:
        amt = float(money.f2(float(eur)))
    except (TypeError, ValueError):
        return False, "the recovered amount must be a number"
    if amt < 0:
        return False, "the recovered amount must be ≥ 0"
    sup, per = (supplier or "").strip(), (period or "").strip()
    if not sup or not per:
        return False, "supplier and period are required"
    ensure_claim(sup, per, detected_eur, actor)
    try:
        con = connect()
        if actor:
            audit.set_actor(con, actor)
        try:
            frag, tp = tenancy.scope_clause()
            con.execute("UPDATE overcharge_claims SET status='recovered', recovered_eur=?, "
                        "updated_at=datetime('now'), updated_by=? WHERE supplier=? AND period=?"
                        + frag, [amt, actor or "", sup, per, *tp])
            con.commit()
            return True, ""
        finally:
            con.close()
    except Exception as e:
        log.warning("record_recovery(%s,%s) failed: %s", supplier, period, e)
        return False, f"could not record the recovery ({str(e)[:60]})"


def _claims_map(period=None):
    """{(supplier, period): row} of stored claims. Never raises -> {}."""
    try:
        con = connect()
        try:
            frag, tp = tenancy.scope_clause()
            sql = "SELECT * FROM overcharge_claims WHERE 1=1" + frag
            params = list(tp)
            if period:
                sql += " AND period=?"
                params.append(period)
            return {(r["supplier"], r["period"]): r for r in con.execute(sql, params)}
        finally:
            con.close()
    except Exception as e:
        log.warning("claims_map failed: %s", e)
        return {}


def overview(period=None):
    """Per (supplier, period): the DETECTED overcharge (aggregated from contract_audit) joined
    with the claim-back status + recovered €. Returns a list sorted by detected € desc; each
    item carries `lines` (the underlying audit flags) for the evidence packet. Never raises."""
    try:
        import contract_audit
        flags, _summ = contract_audit.audit(period)
        agg = {}
        for f in flags:
            k = (f["supplier"], f["period"])
            a = agg.setdefault(k, {"supplier": f["supplier"], "period": f["period"],
                                   "detected_eur": 0.0, "flags": 0, "lines": []})
            a["detected_eur"] = float(money.f2(a["detected_eur"] + f["recover_eur"]))
            a["flags"] += 1
            a["lines"].append(f)
        claims = _claims_map(period)
        # include claims that no longer have a live detection (e.g. recovered) so they show.
        for (sup, per), c in claims.items():
            if (sup, per) not in agg:
                agg[(sup, per)] = {"supplier": sup, "period": per,
                                   "detected_eur": float(c["detected_eur"] or 0),
                                   "flags": 0, "lines": []}
        out = []
        for k, a in agg.items():
            c = claims.get(k)
            a["status"] = c["status"] if c else "detected"
            a["recovered_eur"] = float(c["recovered_eur"] or 0) if c else 0.0
            out.append(a)
        out.sort(key=lambda x: x["detected_eur"], reverse=True)
        return out
    except Exception as e:
        log.warning("overview(%s) failed: %s", period, e)
        return []


def recovered_total(period=None):
    """Total € overcharges RECOVERED (booked) — the north-star figure. Never raises -> 0.0."""
    try:
        con = connect()
        try:
            frag, tp = tenancy.scope_clause()
            sql = ("SELECT COALESCE(SUM(recovered_eur),0) FROM overcharge_claims "
                   "WHERE status='recovered'" + frag)
            params = list(tp)
            if period:
                sql += " AND period=?"
                params.append(period)
            return float(con.execute(sql, params).fetchone()[0] or 0)
        finally:
            con.close()
    except Exception as e:
        log.warning("recovered_total(%s) failed: %s", period, e)
        return 0.0


def evidence_workbook(supplier, period):
    """An Excel EVIDENCE PACKET for one supplier×period: every overcharge line (country,
    station, product, litres, contracted vs actual price, € recoverable) + a total — the
    document a fleet sends the supplier to claim the money back. Returns xlsx bytes, or None
    when there's nothing to package / on error. Never raises."""
    sup, per = (supplier or "").strip(), (period or "").strip()
    try:
        import contract_audit
        flags, _ = contract_audit.audit(per)
        lines = [f for f in flags if f["supplier"] == sup]
        if not lines:
            return None
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Overcharge claim"
        bold = Font(bold=True)
        ws.append([f"Supplier overcharge claim — {sup} — {per}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Basis: NET EUR/L, VAT excluded, rebates applied. Prices contracted vs "
                   "actually invoiced; recoverable = (gap) × litres."])
        ws.append([])
        hdr = ["Country", "Station", "Product", "Litres", "Issue",
               "Contracted EUR/L", "Actual EUR/L", "Recoverable EUR", "Note"]
        ws.append(hdr)
        for c in ws[ws.max_row]:
            c.font = bold
        total = 0.0
        for ln in sorted(lines, key=lambda x: x["recover_eur"], reverse=True):
            ws.append([ln["country"], ln["station"], ln["product"], ln["litres"],
                       ln["issue"], ln["expected"], ln["actual"], ln["recover_eur"],
                       ln["note"]])
            total += ln["recover_eur"]
        ws.append([])
        ws.append(["", "", "", "", "", "", "Total recoverable EUR", float(money.f2(total))])
        for c in ws[ws.max_row]:
            c.font = bold
        for col, w in zip("ABCDEFGHI", (14, 26, 12, 10, 14, 16, 14, 16, 30)):
            ws.column_dimensions[col].width = w
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as e:
        log.warning("evidence_workbook(%s,%s) failed: %s", supplier, period, e)
        return None
